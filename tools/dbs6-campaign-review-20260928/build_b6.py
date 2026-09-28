"""B6 — campaign-by-campaign review of run 20260928-535a4eef (docx + xlsx register)."""
import json, os, re
from collections import Counter, defaultdict
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

S = os.path.dirname(os.path.abspath(__file__))
A = json.load(open(f'{S}/a.json'))
RV = json.load(open(f'{S}/review.json'))
ROWS = json.load(open(f'{S}/rows.json'))
SALES = json.load(open(f'{S}/child_sales.json'))
INV = {r['sku']: r for r in json.load(open(f'{S}/inventory_children.json'))}
OUT = os.environ.get('OUT', '/home/user/PPC-Audit-/docs/DBS6_Campaign_Review_20260928-535a4eef.docx')
NAVY = RGBColor(0x1F, 0x3A, 0x5F); GREY = RGBColor(0x55, 0x55, 0x55)
BE = A['brief']['sections']['margin']['be_acos']

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
for lvl, size in ((1, 15), (2, 12.5), (3, 11)):
    h = doc.styles[f'Heading {lvl}']; h.font.name = 'Calibri'; h.font.size = Pt(size); h.font.color.rgb = NAVY
sec = doc.sections[0]
for m in ('left_margin', 'right_margin'):
    setattr(sec, m, Cm(1.8))
sec.top_margin = sec.bottom_margin = Cm(1.6)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), fill); tcPr.append(s)


def runs(p, text, bold=False, size=None, color=None):
    for part in re.split(r'(\*\*[^*]+\*\*)', text):
        if not part:
            continue
        b = bold
        if part.startswith('**') and part.endswith('**'):
            part, b = part[2:-2], True
        r = p.add_run(part); r.bold = b
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
    return p


def para(t='', **k):
    p = doc.add_paragraph(); runs(p, t, **k); p.paragraph_format.space_after = Pt(4); return p


def lead(l, t):
    p = doc.add_paragraph(); r = p.add_run(l + ' '); r.bold = True; r.font.color.rgb = NAVY; runs(p, t); p.paragraph_format.space_after = Pt(4); return p


def bullet(t):
    p = doc.add_paragraph(style='List Bullet'); runs(p, t); p.paragraph_format.space_after = Pt(2); return p


def table(headers, rows, widths=None, size=8):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ''
        r = c.paragraphs[0].add_run(str(h)); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shade(c, '1F3A5F')
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''; runs(cells[i].paragraphs[0], '' if v is None else str(v), size=size)
            if ri % 2: shade(cells[i], 'F2F5F9')
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def pct(x, d=0):
    return '—' if x is None else f'{x * 100:.{d}f}%'


def usd(x, d=2):
    return '—' if x is None else f'${x:,.{d}f}'


def f1(x, d=1):
    return '—' if x is None else f'{x:.{d}f}'


def short(n):
    return re.sub(r'^DBS6-SP-', '', n)


sec_no = [0]; part = ['A']


def section(t):
    sec_no[0] += 1; doc.add_heading(f'{part[0]}.{sec_no[0]} {t}', 1)


RK = [r for r in RV if r.get('ceil_tos') is not None]
PREFERRED_B6 = {'KING': 'BAMBOO-KING-WHITE-6PCS', 'CALIFKING': 'BAMBOO-CALIFKING-WHITE-6PCS', 'QUEEN': 'BAMBOO-QUEEN-WHITE-6PCS', 'FULL': 'BAMBOO-FULL-WHITE-6PCS', 'TWIN': 'BAMBOO-TWIN-SAGE GREEN-6PCS'}
NR = [r for r in RV if r.get('ceil_tos') is None]
iss = lambda r, k: [i for i in r['issues'] if i.startswith(k)]

KWR = {}
import glob
for f in glob.glob(f'{S}/b46/b6_keywords_30d_p*.json'):
    for r in json.load(open(f))['rows']:
        KWR[r['keyword']] = r
DAILY = {p['date']: p for p in json.load(open(f'{S}/b46/b6_daily.json'))['points']}


def davg(a, b):
    ds = [k for k in DAILY if a <= k <= b]
    sp = sum(DAILY[k]['spend'] for k in ds); sa = sum(DAILY[k]['sales'] for k in ds)
    return sp / len(ds), (sp / sa if sa else None)


DEP = [json.loads(l) for l in open(f'{S}/decisions_recent.jsonl')]
DEP = [x for x in DEP if x.get('status') == 'deployed' and (x.get('decided') or '') >= '2026-09-22']
OUTC = A['run']['outcome']

# ================================================================ title
t = doc.add_paragraph(); runs(t, 'B6 — campaign-by-campaign review of the 28 September run', bold=True, size=17, color=NAVY)
para('Decolure Bamboo Sheets 6 Piece · run 20260928-535a4eef · reviewed against “Bids and Placements on Exact Campaigns” (25 September revision) · data to 27 September', size=9, color=GREY)
para(f"Every live campaign the run touches — {len(RV)} of them, including {sum(1 for r in RV if r.get('not_in_export'))} “FIX” campaigns the run writes to but its campaign list leaves out — was re-read on its own numbers: objective, advertised SKU, clicks, impressions, CTR, CVR, spend, CPC and orders over 90, 45, 30, 14 and 7 days; top-of-search clicks and share, product-page share, top-of-search impression share on its keyword, the required clicks for its target rank, the 7-day median rank against target; and the bid, modifier, budget, placement and SKU the run proposes. B6 is judged on its own data only — its own stock, sales, conversion, ranks and deal calendar; nothing is carried over from B4. Where the run, or the 25 September corrections, got a decision wrong, it is flagged and corrected here rather than carried forward.")
lead('B6’s situation.', 'No deal is running and none ran in September: the last Best Deal was 31 July → 13 August, a Lightning Deal ran on 16 August, and the next dated event is a one-day Lightning Deal on 26 October. Conversion rates here are read on the 45 days from 14 August, after the July deal, so the deal does not inflate the ceilings. Contribution is each child’s at what it actually sold for in the last 30 days, not at list price — B6 children sell well below list (Queen Creme lists at $109.99 and sold at $75.67). Stock is read on Available units only (reserved-inventory guideline, 28 September). Every size’s White has 7+ days of Available stock except Twin.')

doc.add_heading('Part A — What the run (and the earlier corrections) got wrong', 1)

# ---------------------------------------------------------------- A.1 pricing
section('Ranking prices — no velocity window, so maintenance')
above = [r for r in RK if r.get('tos_price') and r['tos_price'] > r['ceil_tos'] * 1.005]
run_up = [r for r in RK if r.get('run_price') and r.get('tos_price') and r['run_price'] > r['tos_price'] + 0.01]
run_up_above = [r for r in run_up if r['run_price'] > r['ceil_tos']]
a1, c1 = davg('2026-09-15', '2026-09-21'); a2, c2 = davg('2026-09-22', '2026-09-25'); a3, c3 = davg('2026-09-26', '2026-09-27')
dk = Counter(x.get('field') for x in DEP)
lead('The rule.', 'Top of search may sit above its ceiling (contribution × the campaign’s blended top-of-search conversion) only on a funded push — sized, dated, ceilinged, predicted and funded — and a push needs all four conditions: a velocity window, CTR above market, CVR above market, and stock to hold the position. Without a velocity window a ranking campaign is in maintenance: top of search at its ceiling, no premium, the 70/20 mix. A price above the ceiling comes down to it — straight there on a row with no top-of-search clicks, in steps of at most 50% on a row that has them.')
lead('What the 25 September corrections put live — wrong for B6.', f"Their daily top-of-search push (step by the delivery gap, bounded by a loss stop of 3× contribution per order) has been deploying since 22 September: {dk.get('placement_multiplier', 0)} top-of-search modifier raises, {dk.get('budget', 0)} budget raises (most doubled) and {dk.get('bid', 0)} base-bid raises. B6 has no deal, so that push has no velocity window behind it. Daily spend went from ${a1:,.0f} (15–21 Sep, ACoS {pct(c1)}) to ${a2:,.0f} (22–25 Sep, {pct(c2)}) to ${a3:,.0f} on 26–27 September (ACoS {pct(c3)} on sales still attributing — it will settle lower, but not to break-even {pct(BE)}). Stop the daily push; every row it raised is re-priced below.")
lead('What the run does.', f"It extends the push: the run’s own outcome takes spend from ${OUTC['spend_day_before']:,.0f} to ${OUTC['spend_day_after']:,.0f} a day, ACoS from {pct(OUTC['acos_before'])} to {pct(OUTC['acos_after'])} (break-even {pct(BE)}) and TACoS from {pct(OUTC['tacos_before'], 1)} to {pct(OUTC['tacos_after'], 1)}, on {OUTC['tos_clicks_wk_before']:.0f} → {OUTC['tos_clicks_wk_after']:.0f} top-of-search clicks a week. Of {len(RK)} ranking campaigns read, {len(above)} already sit above their ceiling and the run raises top of search on {len(run_up)} — {len(run_up_above)} of them to a price above the ceiling.")
ex = sorted(run_up, key=lambda r: -r['d30']['spend'])[:6]
table(['Campaign', 'ACoS 30d', 'Ceiling (rate basis)', 'TOS price now', 'Run writes', 'Corrected (this write → end point)'],
      [[short(r['campaign'])[:52], pct(r['d30']['acos'] / 100 if r['d30']['acos'] is not None else None), f"{usd(r['ceil_tos'])} ({pct(r['tos_rate'], 1)}, {r['tos_src'][:22]})", usd(r['tos_price']), usd(r.get('run_price')), f"{usd(r.get('price_to'))} → {usd(r.get('descent_final') or r.get('price_to'))}" + (' (collapsing: frozen)' if 'COLLAPS' in r['situation'] else '')] for r in ex],
      widths=[5.6, 1.4, 4, 1.8, 1.8, 3.6], size=7.5)
lead('Correction.', 'Every ranking campaign at its ceiling, no premium, until a push is funded and dated. Where top of search today is below the ceiling and the mix is right, it climbs toward the ceiling at most 30% a write; where product pages take over 20% of clicks, the base comes down and the modifier is re-solved so top of search holds. Budgets the push doubled go back to what the ceiling-priced requirement needs. The one real velocity window on the calendar is the 26 October Lightning Deal — a single day, too short to carry a ranking push on its own.')

# ---------------------------------------------------------------- A.2 push candidates
section('Push candidates — B6’s own terms that pass three of the four conditions')
PC = [r for r in RK if r['conds']['ctr'] and r['conds']['cvr'] and r['conds']['stock'] and r.get('rank_now') and r.get('rank_tgt') and r['rank_now'] > r['rank_tgt']]
lead('What they are.', f"{len(PC)} ranking campaigns beat the market on top-of-search click-through and conversion, have 7+ days of Available stock, and sit short of their target rank. Only the velocity window is missing. None is funded in this review — whether B6 funds a push, and on which terms, is the operator’s decision. What each would need is below; most must fix their placement mix first, because a push on a campaign leaking to product pages pays for the wrong placement.")
rows = []
for r in sorted(PC, key=lambda r: -((KWR.get(r['kw']) or {}).get('impressions') or 0)):
    k = KWR.get(r['kw']) or {}
    blk = 'rank collapsing — never scale' if 'COLLAPS' in r['situation'] else ('mix first — ' + r['situation'].split(' — ', 1)[1]) if ('LEAK' in r['situation'] or 'NOT A RANKING' in r['situation']) else 'ready if funded'
    rows.append([r['kw'], (r.get('child') or '').replace('BAMBOO-', '').replace('-6PCS', ''), f"{k.get('impressions') or 0:,}", f"{r['rank_30']} → {r['rank_now']} · {r['rank_tgt']}", f"{f1(r['tos_cvr'])}% vs {r['mcvr']}%", f"{f1(r.get('req_day'))} / {f1(r['deliv_14'])}", f"{usd(r['ceil_tos'])} / {usd(r.get('tos_price'))}", blk])
table(['Term', 'SKU', 'Impr. 30d', 'Rank 08-28 → now · target', 'TOS CVR vs market', 'Req / del TOS/day', 'Ceiling / TOS now', 'Status'], rows, widths=[3.4, 2.2, 1.4, 2.2, 1.8, 1.6, 2, 3.4], size=7)
fl = next((r for r in PC if r['kw'] == 'bamboo sheets'), None)
if fl:
    lead('The flagship, “bamboo sheets” (King White).', f"The biggest term B6 has ({(KWR['bamboo sheets'].get('searchVolume') or 0):,} searches a month by the library). Rank {fl['rank_30']} → {fl['rank_now']} against a target of {fl['rank_tgt']}; it needs {f1(fl['req_day'])} top-of-search clicks a day and takes {f1(fl['deliv_14'])} — at {f1(fl['is30'])}% impression share, so price is what limits it. At the ceiling {usd(fl['ceil_tos'])} the requirement costs about ${fl['req_day'] * fl['ceil_tos']:,.0f} a day; any premium above that is money lost on every order. It is the natural head of a funded B6 push — but the size of that commitment is the operator’s call, not something to run as a daily step.")

# ---------------------------------------------------------------- A.3 SKU
section('Ranking SKU — White where it has the stock')
lead('The rule.', 'Ranking campaigns advertise White. Where White cannot carry the traffic on Available stock, the size’s highest-selling variation with healthy Available stock carries it until White lands. Colour-named terms advertise their own colour. A swap never changes the size.')
S30 = {r['sku']: r for r in SALES['d30']}
RUNPICK = Counter()
for x in A['decisions']:
    if x['kind'] == 'child' and x['verdict'] == 'change':
        RUNPICK[(x['current_at_audit'], x['suggested'])] += 1
size_rows = []
for sz, nm in (('KING', 'King'), ('QUEEN', 'Queen'), ('CALIFKING', 'Cal King'), ('FULL', 'Full'), ('TWIN', 'Twin')):
    w = f'BAMBOO-{sz}-WHITE-6PCS'; wi = INV.get(w, {}); ws = S30.get(w, {})
    top = sorted([r for r in SALES['d30'] if r['sku'].startswith(f'BAMBOO-{sz}-')], key=lambda r: -r['units'])
    rank_w = next((i + 1 for i, r in enumerate(top) if r['sku'] == w), None)
    runp = ', '.join(f"{b.replace('BAMBOO-', '').replace('-6PCS', '').title()} ({n})" for (a, b), n in RUNPICK.items() if a == w) or '—'
    v = wi.get('velocity_day') or 0; b = wi.get('reserved_breakdown') or {}
    pw = PREFERRED_B6[sz]; pi = INV.get(pw, {}); ps = S30.get(pw, {})
    size_rows.append([nm, f"{ws.get('units', 0)} units (#{rank_w or '—'} in size, {pct(ws['organic_units'] / ws['units']) if ws.get('units') else '—'} organic); {wi.get('fba_available')} Available = {(wi.get('fba_available') or 0) / v:.0f} days; reserved {b.get('fc_transfers', 0)} FC transfer / {b.get('fc_processing', 0)} processing / {b.get('customer_orders', 0)} orders; {wi.get('inbound')} inbound" if v else f"{ws.get('units', 0)} units; {wi.get('fba_available')} Available; reserved {b.get('fc_transfers', 0)} FC transfer; {wi.get('inbound')} inbound; no sales",
                      runp, pw.replace('BAMBOO-', '').replace('-6PCS', '').split('-', 1)[1].title() + (' — keep' if pw == w else f" — {ps.get('units')} units, {pi.get('fba_available')} Available = {(pi.get('fba_available') or 0) / (pi.get('velocity_day') or 1):.0f} days")])
table(['Size', 'White: 30-day units · stock', 'The run moves White to', 'Ranking SKU (corrected)'], size_rows, widths=[1.6, 8.2, 3, 5.2], size=7.5)
para('Units and organic share are Sellerboard’s 30-day figures (08-29 → 09-27); its PPC split credits only same-SKU ad sales, so organic is an upper bound. Stock is Sellerboard on 28 September; days = Available ÷ weighted velocity. Next Queen White arrival 30 October (1,002 units across five shipments).', size=8, color=GREY)
qw = INV.get('BAMBOO-QUEEN-WHITE-6PCS', {})
lead('Queen White → Olive — wrong for B6.', f"The run moves {RUNPICK.get(('BAMBOO-QUEEN-WHITE-6PCS', 'BAMBOO-QUEEN-OLIVE-6PCS'), 0)} campaigns off Queen White because, at the pace of the push it proposes, White would run out on day 29, three days before the 30 October arrival. That is the push creating its own stockout. At today’s pace Queen White has {qw.get('fba_available')} Available — {qw['fba_available'] / qw['velocity_day']:.0f} days against 32 to the arrival — and it is the size’s best seller. With the push stopped (A.1), White carries Queen ranking, and Olive (the size’s seventh seller) is not the backup the rule would pick anyway. Keep Queen on White; if a funded push is approved later, size it so White’s Available cover stays above the arrival date.")
bad = [r for r in RV if any('different size' in i for i in r['issues'])]
lead('Size changes — plain errors.', 'The run moves ' + '; '.join(f"“{short(r['campaign'])[:40]}” {next(i for i in r['issues'] if 'different size' in i).split('(')[1].split(')')[0]}" for r in bad) + '. A swap is a stock fix inside a size; none of these has a stock reason (the current children have 43 and 47 days Available), and the Cal King term goes to Cal King White.')
lead('Twin.', 'Twin White has 0 Available (36 in FC transfer) and no sales in 30 days. Twin ranking stays on Sage Green, which the Twin campaigns already advertise: it ties for the size’s best seller and has the most Available stock of the three (54, 142 days). Back to White once the transfers show as Available and it sells.')
sku_iss = [r for r in RV if iss(r, 'SKU: ranking') or iss(r, 'SKU: the term')]
lead('Correction.', f"{len(sku_iss)} ranking campaigns advertise a child other than the size’s ranking SKU (the register names each). Re-point them by hand (the loader skips every swap). Auto, Broad and Phrase campaigns should advertise the size’s LTSF SKU, not White; {sum(1 for r in RV if any('LTSF' in i for i in r['issues']))} of them advertise White today and are flagged — the SKU is named once the LTSF list is in (Sellerboard gives no inventory age).")

# ---------------------------------------------------------------- A.4 stock
section('Stock — PPC runs on Available units only')
lead('The rule (28 September guideline).', 'Ads can only sell Available units. Under 5 days of Available with reserved mostly FC transfer → throttle; under 5 days with reserved mostly customer orders → stockout risk; 7+ days → normal. Customer orders never count; FC transfer and FC processing count for replenishment only.')
SR = [r for r in RV if (r.get('situation') or '').startswith(('STOCK', 'LOW STOCK'))]
low = sorted([(k, v) for k, v in INV.items() if v.get('velocity_day') and (v.get('fba_available') or 0) / v['velocity_day'] < 7 and S30.get(k, {}).get('units', 0) >= 5], key=lambda kv: -S30[kv[0]]['units'])
lead('What the data shows.', (f"{len(SR)} advertised campaigns sit on a child under 7 days." if SR else 'No advertised B6 child is under 7 days of Available stock: every ranking SKU above clears the line, and no campaign needs throttling today.') + ' Children that sell but are short, to keep out of any new ad group or swap: ' + '; '.join(f"{k.replace('BAMBOO-', '').replace('-6PCS', '').title()} {v['fba_available']} Available = {v['fba_available'] / v['velocity_day']:.0f} days ({(v.get('reserved_breakdown') or {}).get('fc_transfers', 0)} in FC transfer)" for k, v in low[:6]) + '.')

# ---------------------------------------------------------------- A.5 objectives
section('Campaign objectives')
ob = Counter()
for r in RV:
    for i in iss(r, 'OBJECTIVE'):
        ob[re.sub(r'\d+', '#', i.split(':', 1)[1].strip())[:80]] += 1
lead('The rule.', 'The objective follows the targeting: Exact = Ranking, Auto/Broad/Phrase = Discovery (clearance keeps Liquidation), product targeting = Conversions, the brand term = Defensive — one objective per campaign.')
table(['Finding', 'Campaigns'], [[k, v] for k, v in ob.most_common()], widths=[14, 2.5], size=8)
lead('What it means.', 'The Exact campaigns tagged Conversions are ranking campaigns under the standard — priced on the ranking ceiling, managed to the 70/20 mix, read against a rank target where one exists. The brand-term campaigns tagged Ranking or Discovery are Defensive: they hold the brand shelf and are not climbed.')

# ---------------------------------------------------------------- A.6 traffic
section('Required clicks and whether the traffic exists')
tr = [r for r in RK if iss(r, 'TRAFFIC')]
withreq = [r for r in RK if r.get('req_day')]
lead('The rule.', 'Required top-of-search clicks a day = the units a day the market sells at the target rank ÷ the campaign’s top-of-search conversion, split across the campaigns on the same term. It sets volume and budget, never price. A shortfall is priced only when the budget is intact and impression share is low; where winning every auction at today’s volume still falls short, the requirement is above what the auction holds.')
lead('What the data shows.', f"{len(withreq)} ranking campaigns carry a term with a target rank and units. {sum(1 for r in withreq if r['deliv_14'] >= 0.7 * r['req_day'])} deliver 70% or more of their requirement; {len(tr)} have a requirement above what the whole auction holds at today’s volume.")
ex = sorted(tr, key=lambda r: -(r['d30']['spend'] or 0))[:6]
table(['Campaign', 'Target rank · now', 'Required TOS/day', 'Delivered (14d)', 'TOS impr. share', 'Most the auction holds'],
      [[short(r['campaign'])[:50], f"{r.get('rank_tgt') or '—'} · {r.get('rank_now') or '—'}", f1(r.get('req_day')), f1(r.get('deliv_14')), f"{f1(r.get('is30'))}%", f1(r.get('reach_day'))] for r in ex], widths=[6.5, 2.4, 2.2, 2.2, 2.2, 2.6], size=7.5)
lead('Correction.', 'Grade these on impression share and rank, not on a click count the auction cannot supply; re-base their requirement. They are not reasons to raise price.')

# ---------------------------------------------------------------- A.7 placement
section('Placement — where the clicks land')
lk = [r for r in RK if 'LEAK' in (r.get('situation') or '') or 'NOT A RANKING' in (r.get('situation') or '')]
lead('The rule.', 'Ranking campaigns aim for 70–90% of clicks at top of search and product pages at or under 20%. Over 20% → the base comes down toward what product pages afford (≤50% a write, ≤25% where product pages carry orders) and the modifier is re-solved so top of search holds; under 30% at top of search that fix comes before any price or budget move.')
lead('What the data shows.', f"This is B6’s largest problem. {sum(1 for r in lk if 'NOT A RANKING' in r['situation'])} ranking campaigns take under 30% of their clicks at top of search and {sum(1 for r in lk if 'LEAK' in r['situation'])} more lose over 20% to product pages — {len(lk)} of {len(RK)}. The run’s push raises top of search on several of them, which buys nothing while the base keeps buying product pages.")
ex = sorted(lk, key=lambda r: -(r['d30']['spend'] or 0))[:8]
table(['Campaign', 'Clicks 30d', 'TOS / PP share', 'Base now → corrected', 'TOS price now → corrected', 'The run'],
      [[short(r['campaign'])[:50], r['d30']['clicks'], f"{pct(r['d30']['tos_sh'])} / {pct(r['d30']['pp_sh'])}", f"{usd(r.get('base'))} → {usd(r.get('base_to'))}", f"{usd(r.get('tos_price'))} → {usd(r.get('price_to'))}", '; '.join(r['run'][:2]) or 'holds'] for r in ex], widths=[6, 1.4, 2.2, 3, 3, 3.6], size=7.5)

# ---------------------------------------------------------------- A.8 keywords
section('Keywords and search terms — one owner per ranking term')
st_r = [r for r in RV if iss(r, 'SEARCH TERM') and r.get('ceil_tos') is not None]
st_b = [r for r in RV if iss(r, 'SEARCH TERM') and r['match'] in ('Auto', 'Broad', 'Phrase')]
kd = [r for r in RV if iss(r, 'KEYWORD')]
ne = [r for r in RV if r.get('not_in_export')]
lead('The rule.', 'A ranking term earns rank credit only through the exact campaign that owns it. Where a Broad, Phrase or Auto campaign also buys the term, it is negative-exact there. Two enabled exact campaigns on one term split the read; one keeps it.')
lead('What the data shows.', f"{len(st_r)} ranking terms are also bought by broad or auto campaigns, and {len(st_b)} broad/auto/phrase campaigns buy exact ranking terms — the size-level Broad campaigns take up to 15 clicks a month each on “bamboo sheets” alone. {len(kd)} campaigns share an exact term with another enabled exact campaign.")
table(['Campaign', 'What it buys (30 days)'], [[short(r['campaign'])[:55], iss(r, 'SEARCH TERM')[0].split(' — ', 1)[-1][:150]] for r in sorted(st_b, key=lambda r: -(r['d30']['spend'] or 0))], widths=[6.5, 11], size=7.5)
lead('The FIX campaigns the export leaves out.', f"The run writes {sum(1 for x in A['decisions'] if x['verdict'] == 'change' and x.get('campaign_id') in {r['cid'] for r in ne})} changes to {len(ne)} “FIX” exact campaigns that are not in its campaign list (${sum(r['d30']['spend'] for r in ne):,.0f} spend in 30 days). They are reviewed in the register from Amazon’s placement report: {sum(1 for r in ne if 'LEAK' in (r.get('situation') or '') or 'NOT A RANKING' in (r.get('situation') or ''))} leak to product pages and {sum(1 for r in ne if (r.get('situation') or '').startswith('NO CLICKS'))} take no clicks. Confirm each one’s status, child and targets before any of those changes loads.")
lead('Correction.', 'Negative-exact every ranking term in the Broad, Phrase and Auto campaigns that buy it (the register lists each); keep one exact owner per term — the campaign on the ranking SKU with the clicks — and pause the duplicate.')

# ---------------------------------------------------------------- A.9 other types
section('Other campaign types — inside their own ceilings')
wrong = [r for r in NR if any(i.startswith('RUN:') or i.startswith('MISSING') for i in r['issues'])]
lead('The rule.', f'Conversions, Discovery, Defensive and Liquidation campaigns are priced inside their ceiling: ACoS at or under break-even ({pct(BE, 1)}), read on 30 days against 90. Over the ceiling the base comes down; inside it nothing moves; under 15 clicks nothing moves.')
lead('What the data shows.', f"{len(NR)} non-ranking campaigns: {sum(1 for r in NR if (r.get('situation') or '').startswith('INSIDE'))} inside their ceiling, {sum(1 for r in NR if (r.get('situation') or '').startswith('OVER'))} over it, {sum(1 for r in NR if (r.get('situation') or '').startswith('THIN'))} too thin to read. {len(wrong)} run decisions go against the rule.")
if wrong:
    table(['Campaign', 'Objective', 'Clicks / orders 30d', 'ACoS 30d (90d)', 'The run', 'Corrected'],
          [[short(r['campaign'])[:48], r['objective'], f"{r['d30']['clicks']} / {r['d30']['orders']}", f"{pct(r.get('acos30'))} ({pct(r.get('acos90'))})", '; '.join(r['run'][:2]) or 'holds', (r.get('action') or [''])[0][:70]] for r in sorted(wrong, key=lambda r: -(r['d30']['spend'] or 0))[:10]],
          widths=[5.4, 2, 2.2, 2.4, 3.4, 4.8], size=7.5)

# ---------------------------------------------------------------- A.10 rows not to load
section('Rows in the change file that should not load')
CS = {c['campaign_id']: c for c in A['campaigns']}
ch = [x for x in A['decisions'] if x['verdict'] == 'change']
pz = [x['seq'] for x in ch if x['kind'] in ('tos', 'bid', 'ros', 'pp', 'budget') and CS.get(x['campaign_id'], {}).get('status') == 'PAUSED']
ms = [x for x in ch if x['kind'] in ('tos', 'bid', 'budget', 'ros', 'pp', 'state') and x.get('campaign_id') and x['campaign_id'] not in CS]
wh = [r['seq'] for r in ROWS if r['verdict'] == 'change' and r['withheld']]
nch = sum(1 for x in ch if x['kind'] == 'child')
bullet(f"**Every top-of-search, bid and budget raise on a ranking campaign** — {len(run_up)} top-of-search raises in all (A.1). Load the corrected prices from the register instead.")
bullet(f"**{len(wh)} rows the engine itself withholds** (their text says “change → HOLD”) but export as changes: seq {', '.join(map(str, wh))}.")
bullet(f"**{len(pz)} writes on campaigns paused at the audit**: seq {', '.join(map(str, pz))}.")
bullet(f"**{len(ms)} rows on the {len({x['campaign_id'] for x in ms})} FIX campaigns the export does not carry** (A.8) — hold until each campaign is confirmed; the cuts among them can stand.")
bullet(f"**The {RUNPICK.get(('BAMBOO-QUEEN-WHITE-6PCS', 'BAMBOO-QUEEN-OLIVE-6PCS'), 0)} Queen White → Olive swaps and the 3 size-changing swaps** (A.3). All {nch} swaps skip at load anyway (no product-ad roster); the ones that move a campaign onto White are right and are done by hand.")
bullet('**Engine check, 1 failure**: a video probe placed beside a Sponsored Brands probe on the same group (R14) — hold that build.')

# ================================================================ Part B — register
new = doc.add_section(); new.orientation = WD_ORIENT.LANDSCAPE
new.page_width, new.page_height = sec.page_height, sec.page_width
for m in ('left_margin', 'right_margin'):
    setattr(new, m, Cm(1.2))
part[0] = 'B'; sec_no[0] = 0
doc.add_heading('Part B — Every campaign, one row each', 1)
para('Sorted by objective, then 30-day spend. Clicks/CTR/CVR/ACoS are 30 days (08-29 → 09-27); TOS/PP are the share of clicks at top of search and on product pages; Req/Del are required vs delivered top-of-search clicks a day (14 days); IS is top-of-search impression share on the campaign’s keyword (30 days). Price = the top-of-search price (base × (1 + modifier)); “now → corrected” is this write. The last column is the corrected decision; “!” marks what the run or the earlier corrections got wrong.', size=8, color=GREY)
order = {'Ranking': 0, 'Conversions': 1, 'Discovery': 2, 'Defensive': 3, 'Liquidation': 4}


def verdict(r):
    bits = []
    if r.get('situation'):
        bits.append(r['situation'])
    if r.get('ceil_tos') is not None and r.get('price_to') is not None:
        bits.append(f"TOS {usd(r.get('tos_price'))} → {usd(r['price_to'])}" + (f" (then {usd(r['descent_final'])})" if r.get('descent_final') and abs(r['descent_final'] - r['price_to']) > 0.01 else '') + (f"; base {usd(r.get('base'))} → {usd(r.get('base_to'))}" if r.get('base_to') and r.get('base') and abs(r['base_to'] - r['base']) > 0.004 else '') + (f"; mod {r['mod_to']}%" if r.get('mod_to') is not None else ''))
    elif r.get('action'):
        bits.append(r['action'][0])
    for fx in r.get('fixes', []):
        bits.append(fx)
    for i in r['issues']:
        if i.startswith(('RUN', 'MISSING', 'OBJECTIVE', 'SKU', 'TRAFFIC', 'SEARCH TERM', 'KEYWORD', 'ROS')):
            bits.append('! ' + i)
    return '. '.join(bits)[:520]


rows = []
for r in sorted(RV, key=lambda r: (order.get(r['objective'], 9), -(r['d30']['spend'] or 0))):
    d = r['d30']
    rows.append([short(r['campaign'])[:60], f"{r['objective'][:4]}→{(r.get('obj_expected') or r['objective'])[:4]}" if r.get('obj_expected') and r.get('obj_expected') != r['objective'] else r['objective'][:4],
                 (r.get('child') or '—').replace('BAMBOO-', ''), f"{d['clicks']}/{f1(d['ctr'], 2)}%/{f1(d['cvr'])}%", usd(d['spend'], 0), pct(d['acos'] / 100 if d['acos'] is not None else None),
                 f"{pct(d['tos_sh'])}/{pct(d['pp_sh'])}", f"{r.get('rank_now') or '—'}/{r.get('rank_tgt') or '—'}", f"{f1(r.get('req_day'))}/{f1(r.get('deliv_14'))}",
                 f1(r.get('is30')), '; '.join(r['run'][:3]) or 'holds', verdict(r)])
table(['Campaign', 'Obj', 'SKU', 'Clk/CTR/CVR', 'Spend', 'ACoS', 'TOS/PP', 'Rank/tgt', 'Req/Del', 'IS', 'The run', 'Corrected decision'], rows,
      widths=[5.2, 1.3, 2.4, 2.2, 1.2, 1.1, 1.5, 1.3, 1.4, 0.9, 3.2, 8.5], size=6)

doc.save(OUT); print('saved', OUT)

# ================================================================ xlsx register
import openpyxl
from openpyxl.styles import Font, PatternFill
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Every campaign'
hdr = ['Campaign', 'Campaign ID', 'Status', 'Objective', 'Objective (standard)', 'Match', 'Advertised SKU', 'Ranking SKU for size', 'Focus', 'Budget', 'Budget use 7d',
       'Impr 30d', 'Clicks 30d', 'CTR 30d', 'Orders 30d', 'CVR 30d', 'Spend 30d', 'CPC 30d', 'ACoS 30d', 'Clicks 90d', 'Orders 90d',
       'TOS clicks 30d', 'TOS share 30d', 'PP share 30d', 'TOS clicks pre-deal/day', 'TOS clicks/day 14d',
       'Keyword', 'Rank 30d ago', 'Rank now', 'Target rank', 'Units/day at target', 'Required TOS clicks/day', 'Most the auction holds/day', 'TOS impr share 30d',
       'Contribution', 'TOS rate', 'TOS rate basis', 'PP rate', 'TOS ceiling', 'PP ceiling (base)', 'Base now', 'TOS mod now', 'TOS price now', 'Run TOS price',
       'Corrected TOS price (this write)', 'End point', 'Corrected base', 'Corrected modifier',
       'Situation', 'Action', 'The run', 'Issues', 'Fixes']
ws.append(hdr)
for r in sorted(RV, key=lambda r: (order.get(r['objective'], 9), -(r['d30']['spend'] or 0))):
    d = r['d30']
    ws.append([r['campaign'], r['cid'], r['status'], r['objective'], r.get('obj_expected'), r['match'], r.get('child'), r.get('preferred_child'), r['focus'], r.get('budget'), r.get('util7'),
               d['impr'], d['clicks'], d['ctr'], d['orders'], d['cvr'], d['spend'], d['cpc'], d['acos'], r['d90']['clicks'], r['d90']['orders'],
               d['tos_c'], d['tos_sh'], d['pp_sh'], r.get('deliv_pre'), r.get('deliv_14'),
               r.get('kw'), r.get('rank_30'), r.get('rank_now'), r.get('rank_tgt'), r.get('units_day'), r.get('req_day'), r.get('reach_day'), r.get('is30'),
               r.get('contrib'), r.get('tos_rate'), r.get('tos_src'), r.get('pp_rate'), r.get('ceil_tos'), r.get('ceil_pp'), r.get('base'), r.get('tos_mod'), r.get('tos_price'), r.get('run_price'),
               r.get('price_to'), r.get('descent_final'), r.get('base_to'), r.get('mod_to'),
               r.get('situation'), ' | '.join(r.get('action') or []), ' | '.join(r['run']), ' | '.join(r['issues']), ' | '.join(r.get('fixes') or [])])
for cell in ws[1]:
    cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='1F3A5F')
ws.freeze_panes = 'B2'; ws.auto_filter.ref = ws.dimensions; ws.column_dimensions['A'].width = 60
w2 = wb.create_sheet('Child sales & stock')
w2.append(['SKU', 'Units 30d', 'PPC units 30d (Sellerboard)', 'Organic units 30d (Sellerboard)', 'Units pre-deal', 'FBA available', 'Reserved', 'Inbound', 'Velocity/day', 'Days of stock'])
pre = {r['sku']: r for r in SALES['pre']}
for r in sorted(SALES['d30'], key=lambda r: (r['sku'].split('-')[1], -r['units'])):
    i = INV.get(r['sku'], {})
    w2.append([r['sku'], r['units'], r['ppc_units'], r['organic_units'], (pre.get(r['sku']) or {}).get('units'), i.get('fba_available'), i.get('reserved'), i.get('inbound'), i.get('velocity_day'), i.get('days_of_stock')])
for cell in w2[1]:
    cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='1F3A5F')
XL = OUT.replace('.docx', '.xlsx'); wb.save(XL); print('saved', XL)
