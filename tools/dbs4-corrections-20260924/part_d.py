# Part D — B4: why rank has not moved. Executed inside build_cat_doc.py's namespace (doc, para, lead, bullet, table, section, part, sec_no, CATS, A).
import statistics as _st
import glob as _glob
import json as _json
import diag as _dg
from docx.shared import Cm as _Cm

part[0] = 'D'; sec_no[0] = 0
doc.add_page_break()
doc.add_heading('Part D — B4: why rank has not moved, and what to do about it', 1)
para('Erik’s question: across the implementations and tuning so far, rank traction has not come. Is that because we are not driving enough traffic to the top of search (a **traffic gap**), or because we are there and not meeting the target CTR and CVR (a **performance gap**)? This part answers it for B4 on the data to 24 September — for the focus this run sets (Bamboo|California King and Bamboo|Cooling) and for the Bamboo, Queen and King terms that moved out of focus on the operator’s ruling of 24 September — and adds the performance-vs-tuning history, competitors, the campaign types that should also be running, and the metrics the app should be reading.')

# ---------------------------------------------------------------- data
PLAN = {k['keyword']: k for k in A['brief']['sections']['targets']['keywords']}
KW4 = _dg.kw_table('b4', plan=PLAN)
GR4 = _dg.group_table('b4')
INF = [c for c in CATS if c['focus_in']]
_iss = [c['tos_is'] for c in INF if c.get('tos_is') is not None]
_need = sum(c['need_day'] or 0 for c in INF); _deal = sum(c['deal_tos_day'] or 0 for c in INF)
_nd = [(c['deal_tos_day'] or 0) / c['need_day'] for c in INF if c.get('need_day')]
FK = {c['kw'] for c in INF if c.get('kw')}


def f1(x, d=1, suf=''):
    return '—' if x is None else f'{x:.{d}f}{suf}'


def rk(r):
    v = list(r['ranks'].values())
    return ' → '.join('—' if x is None else str(x) for x in (v[0], v[2], v[3]))


def verdict_short(v):
    return v.split(' — ')[0].replace('NOT PRESENT at top of search', 'NOT PRESENT').title()


def group_is(prod):
    ti = _dg.targets_is(prod); k30 = _dg.kw_rows(prod, '30d')
    acc = {}
    for k, t in ti.items():
        r = k30.get(k)
        if not r or t['is_'] is None: continue
        w = r.get('tos_impr') or 0
        a = acc.setdefault(r['syntax'], [0, 0]); a[0] += t['is_'] * w; a[1] += w
    return {g: (a[0] / a[1] if a[1] else None) for g, a in acc.items()}


GIS4 = group_is('b4')
_up = _n = 0
for _f in _glob.glob(f'{_dg.S}/b4_targets_30d_*.json'):
    _cid = _f.rsplit('_', 1)[1][:-5]
    try:
        _a = _json.load(open(_f)); _b = _json.load(open(f'{_dg.S}/b4_targets_deal_{_cid}.json'))
    except Exception:
        continue
    _bd = {r['target']: r.get('tosImpressionShare') for r in _b['rows']}
    for r in _a['rows']:
        x, y = r.get('tosImpressionShare'), _bd.get(r['target'])
        if x is not None and y is not None:
            _n += 1; _up += y > x
G = {g['group']: g for g in GR4}

# ---------------------------------------------------------------- D1 answer
section('The answer in one page')
lead('Short answer.', 'Still overwhelmingly a **traffic gap**, and on the new focus a volume gap as much as a price gap. Where B4 reaches the top of search it beats the app’s CTR bar (1.1× market CTR) and CVR bar (3× market CVR) on every group with a readable sample except the broad Cooling head term. But it wins only a sliver of the top-of-search auctions, and on the California King terms the whole top-of-search pool is small.')
bullet(f'**We are rarely at the top.** Amazon’s top-of-search impression share on the in-focus exact targets has a median of **{_st.median(_iss):.1f}%** ({sum(i < 10 for i in _iss)} of {len(_iss)} under 10%). On “bamboo cooling sheets” — the biggest in-focus term — it is 2.6% at a $7.61 effective top-of-search bid; on “california king bamboo sheets” 0.3% at $3.84; the best is “king sheets bamboo cooling” at 21.9%.')
bullet(f'**The volume is under a tenth of the plan.** The run’s click requirements for the {len(INF)} in-focus campaigns add to {_need:.0f} top-of-search clicks a day; over the deal (15–24 Sep) they got **{_deal:.1f} a day ({_deal / _need * 100:.0f}%)**. The median campaign is at {_st.median(_nd) * 100:.0f}% of its requirement and none is at it. In-focus spend is about $67 a day of the product’s $1,142 in the deal: the push this run declares is a small share of what B4 spends.')
bullet('**When we are there, we perform.** At the top of search over 30 days: Bamboo|Cooling CTR 5.2% vs a 2.1% bar and CVR 19.1% vs 7.5%; “bamboo cooling sheets” 3.5% / 19.5% on 54 clicks. Bamboo|California King: CTR high, CVR 10.1% vs an 8.8% bar — but on only 35 top-of-search clicks in 30 days across the whole group, so the call rests on a thin sample.')
bullet('**Cal King: the ranks are already good; the pool is small.** On the 7-day median “cal king bamboo sheet set” is at 7, “california king bamboo sheets” 9, “bamboo california king sheets” 11, “bamboo sheets california king” 12 — the group moved 28 → 18 since June. Every Cal King term took 2–7 top-of-search clicks in 30 days. The push can buy the top slot on them, but the requirement (40 clicks a week per campaign in the run) is above what the auction holds on most (A-section reachability check), and 7-day rank is the grade, not clicks.')
bullet('**Cooling: the push is working where it was given price — with one exception.** Bamboo|Cooling moved 36 → 26 on the 7-day median in the deal; “bamboo cooling sheets king size” 28 → 11 and “king size sheets set bamboo cooling” 43 → 18. The head term “bamboo cooling sheets” is better over 30 days (45 → 29) but slipped in the deal (24 → 29) while its impression share rose from 2.6% to 14.3% — its ad sends shoppers to Queen White, which is running short, so the re-point comes before more price. Top-of-search impression share rose on {up} of {n} targets with both reads after the 15 September raise.'.format(up=_up, n=_n))
bullet('**The real performance gap is still narrow:** “cooling sheets”, the broad head term (CVR 7.5% vs an 8.7% bar, ACoS 75%, rank 55 → 60) — out of focus this run, and it should stay that way until the offer answers the query. The size-qualified cooling groups (Cooling|King, Cooling|Queen) meet both bars and are traffic gaps.')
bullet('**The tuning worked against rank until the deal.** From July to the week before the deal the rounds took ad clicks from 401 to 235 a day and orders from 46 to 28; top-of-search clicks fell to 92 a day while product pages carried 125. The deal round bought volume back (342 clicks, 148 at top of search a day) at a price: CVR 11.9% → 10.1%, CPC $2.92 → $3.33, ACoS 45% against a 33.8% break-even.')
lead('What changes.', 'Buy the traffic where it performs: step the top-of-search price every day of the deal by the delivery gap, inside the loss stop and the envelope (Part A), and read impression share with it — share rising with price is the proof the price works. On Cal King, grade on share and 7-day rank, not on a click count the auction cannot supply. Hold price and fix the offer only on the terms with a real CTR/CVR shortfall (“cooling sheets”). The out-of-focus hero terms hold (Part B): their rank loss on Queen coincides with the Queen White hero running short, which price does not fix.')

# ---------------------------------------------------------------- D2 performance vs tuning
section('Overall performance against the tuning so far')
lead('What happened.', 'Four blocks from the daily product series, and the organic group rank (7-day median) at the start of each. The deal is the Best Deal (15–28 September); its first 10 days are in, with attribution settled to 18 September.')
rows = []
pts = _json.load(open(f'{_dg.B}/b4_daily.json'))['points']
for a, b, wl in (('2026-06-24', '2026-07-23', '24 Jun – 23 Jul'), ('2026-07-24', '2026-08-23', '24 Jul – 23 Aug'), ('2026-08-24', '2026-09-14', '24 Aug – 14 Sep'), ('2026-09-15', '2026-09-24', 'Deal 15 – 24 Sep')):
    r = [x for x in pts if a <= x['date'] <= b]; n = len(r)
    s = lambda k: sum((x.get(k) or 0) for x in r)
    c, o, sp, sa = s('clicks'), s('orders'), s('spend'), s('sales')
    tac = [x['tacos'] for x in r if x.get('tacos')]
    rows.append(['B4' if a == '2026-06-24' else '', wl, f'{c / n:.0f}', f'{o / n:.1f}', f'${sp / n:,.0f}', f'${sa / n:,.0f}', f'{sp / sa * 100:.1f}%', f'{o / c * 100:.1f}%', f'${sp / c:.2f}', f'{sum(tac) / len(tac):.1f}%' if tac else '—'])
table(['Product', 'Window', 'Ad clicks/day', 'Ad orders/day', 'Spend/day', 'Ad sales/day', 'ACoS', 'CVR', 'CPC', 'TACoS (mean of days)'], rows, widths=[1.2, 2.8, 1.6, 1.6, 1.6, 1.7, 1.3, 1.2, 1.2, 1.8], size=7.5)
rows = []
for g in GR4:
    if g['group'].startswith('Size') or g['group'] == 'Organic & Natural' or (g['spend'] or 0) < 20: continue
    v = list(g['ranks'].values())
    rows.append([g['group'] + (' (in focus)' if g['group'] in ('Bamboo|California King', 'Bamboo|Cooling') else ''), *['—' if x is None else x for x in v], f"${g['spend']:,.0f}"])
table(['Syntax group', 'Rank 24 Jun', 'Rank 15 Aug', 'Rank 14 Sep', 'Rank 24 Sep', 'Spend 30d'], rows, widths=[5, 2, 2, 2, 2, 2], size=8)
p = f'{_dg.S}/chart_b4.png'
if os.path.exists(p):
    doc.add_picture(p, width=_Cm(16.5)); para('B4 — top of search and product-page clicks per day (all campaigns), the spend-weighted effective top-of-search bid, and the rank of the two in-focus groups and Bamboo|Queen. Dashed lines: the change rounds of 24 Aug, 10 Sep, 15/16 Sep and 22 Sep; shaded: the deal to 24 Sep.', size=8.5, color=GREY)
lead('B4.', 'Through August and early September every round took volume out (ad clicks 401 → 328 → 235 a day, orders 46 → 38 → 28) while the placement mix improved — the mix moved because product pages were cut harder than top of search. Rank drifted on the big groups: Bamboo 31 → 35, Bamboo|Queen 32 → 44. The deal round (15/16 Sep) bought volume back: top-of-search clicks 92 → 148 a day. The groups given top-of-search volume moved — Bamboo|King 31 → 24, Bamboo|Cooling 36 → 26, Cooling|King 73 → 53, Cooling|Queen 79 → 58 — while Bamboo (35 → 33) and Bamboo|Queen (44 → 43) held and Bamboo|California King sat at 18–19. The deal also lowered product CVR (11.9% → 10.1%) and lifted ACoS to 45%; part is attribution lag (settled to 18 September).')
lead('The lesson for the app.', 'A round that improves ACoS or the mix but lowers top-of-search clicks on an in-focus keyword is a rank loss, not a win — B4’s August rounds did exactly that, and the deal round, which bought volume, is the one that moved rank. Grade every round on top-of-search clicks delivered against the requirement, and on rank, before ACoS.')

# ---------------------------------------------------------------- D3 method
section('How traffic gap and performance gap are told apart')
lead('Performance bars.', 'The app’s own: target CTR = 1.1 × the market CTR of the keyword’s syntax group (Search Query Performance, 30 days — the deal window’s market figures are blank); target CVR = 3 × market CVR. CTR is judged only with 300+ top-of-search impressions, CVR only with 15+ top-of-search clicks — below that the call is “not enough sample”, never a verdict.')
lead('Traffic bars.', 'Three reads, any of which marks a traffic gap: Amazon’s measured top-of-search impression share on the target under 10%; fewer than 15 top-of-search clicks in 30 days; delivered top-of-search clicks under 70% of the requirement.')
lead('Verdicts.', '**Traffic gap** — meets the bars when shown, not shown enough → buy the traffic. **Performance gap** — shown enough, under the CTR or CVR bar → fix the offer before buying more. **Both** — fix the offer while probing. **Not present** — under 100 top-of-search impressions in 30 days → eligibility/price first. **On track** — meets both, impression share ≥ 10% → scale while rank moves.')
lead('Caveat.', 'Amazon reports placement per campaign and performance per keyword, never both; the keyword-level top-of-search figures distribute each keyword’s traffic by its campaigns’ measured split (Command Center’s estimate). Impression share and effective bid are measured per target. Where a keyword runs in one exact campaign — true of most ranking keywords — the split is that campaign’s own.')

# ---------------------------------------------------------------- D4 syntax
section('Syntax groups')
rows = []
for g in sorted(GR4, key=lambda r: -(r['spend'] or 0)):
    if not g['clicks']: continue
    is_ = GIS4.get(g['group'])
    v = _dg.verdict(g['tos_impr'], g['tos_clicks'], g['tos_ctr'], g['tos_cvr'], g['mctr'], g['mcvr'], is_=is_)
    rows.append([g['group'], f"${g['spend']:,.0f} / {g['clicks']:,}", f"{g['tos_clicks']:.0f} ({f1(g['tos_share'], 0)}%)", f"{f1(g['tos_ctr'])} / {f1(v['tctr'])}", f"{f1(g['tos_cvr'])} / {f1(v['tcvr'])}", f1(is_, 1, '%') if is_ is not None else '—',
                 f"{f1(g['tos_cpc'], 2)}", '—' if g['group'].startswith('Size') else ' → '.join('—' if x is None else str(x) for x in g['ranks'].values()), v['verdict']])
doc.add_heading('B4 — 30 days to 24 Sep', 3)
table(['Group', 'Spend / clicks 30d', 'TOS clicks (share)', 'TOS CTR / bar', 'TOS CVR / bar', 'TOS impr. share', 'TOS CPC', 'Rank 24 Jun → 15 Aug → 14 Sep → 24 Sep', 'Verdict'], rows, widths=[2.4, 1.5, 1.8, 1.5, 1.5, 1.3, 1.1, 2.2, 3.4], size=7)
para('Size groups (“Size: King” etc.) are cross-cutting: a keyword is in its primary group and in its size group, so the size rows re-count traffic already shown above. Impression share for a group is the impression-weighted mean of its member targets’ measured share.', size=8.5, color=GREY)
lead('Reading it.', 'Every Bamboo group and every size-qualified Cooling group meets its CTR and CVR bars at top of search at low single-digit impression share — a traffic gap. Cooling (the broad group) is the exception on CVR, carried by the head term “cooling sheets”. The two in-focus groups: Bamboo|Cooling is a clean traffic gap on a real sample (92 top-of-search clicks, CVR 19.1% vs 7.5%); Bamboo|California King meets both bars on a thin sample (35 clicks) and is small in absolute terms — $298 in 30 days across the group.')

# ---------------------------------------------------------------- D5 keywords
section('The main ranking keywords — one call each')
lead('What the columns say.', 'Req/wk = the run’s weekly click requirement for the campaign carrying the keyword (the brief carries no per-keyword requirement for the new focus terms; for the out-of-focus hero terms it is the brief’s). TOS/wk = top-of-search clicks a week over 30 days, and in the deal. IS = Amazon’s top-of-search impression share on the exact target, at the effective top-of-search bid shown. Rank = our tracked rank (7-day median) 24 Jun → 14 Sep → 24 Sep.')
_req = {c['kw']: c['plan_wk'] for c in INF if c.get('kw')}
for title_, pick in (('In focus — Bamboo|California King and Bamboo|Cooling', lambda r: r['kw'] in FK), ('Out of focus — the hero terms (held; Part B)', lambda r: r['kw'] not in FK and (r['clicks'] or 0) >= 150)):
    rows = []
    for r in [r for r in KW4 if pick(r)][:16]:
        pl = PLAN.get(r['kw'])
        req = _req.get(r['kw']) or (pl or {}).get('ppc_clicks_target')
        wk = r['tos_clicks'] / 30 * 7
        rows.append([r['kw'][:34], f'{req:.0f}' if req else '—', f"{wk:.0f} / {r['deal_tos_day'] * 7:.0f}",
                     f"{f1(r['tos_ctr'])}/{f1(r['tctr'])}", f"{f1(r['tos_cvr'])}/{f1(r['tcvr'])}", f1(r['is_'], 1, '%') if r['is_'] is not None else '—', f"${r['eff']:.2f}" if r['eff'] else '—',
                     rk(r), verdict_short(r['verdict'])])
    doc.add_heading(title_, 3)
    table(['Keyword', 'Req/wk', 'TOS/wk 30d / deal', 'CTR / bar', 'CVR / bar', 'IS', 'Eff. TOS bid', 'Rank', 'Verdict'], rows, widths=[3.8, 1.1, 1.7, 1.5, 1.6, 1.1, 1.3, 2.2, 2.1], size=7)
lead('In focus — what to do, by keyword.', '')
for kw, txt in [
    ('bamboo cooling sheets', 'Traffic gap, the one in-focus term with a real sample: CTR 3.5% vs 2.1%, CVR 19.5% vs 7.5% on 54 top-of-search clicks; 2.6% impression share at $7.61. Rank 45 → 29 over 30 days, but 24 → 29 inside the deal while its share rose from 2.6% to 14.3%; ACoS 77% at a $7.61 effective bid. It needs 132 clicks a week and got 31 in the deal. The step is +30% to its loss stop ($9.73); its modifier is at 718%, so the 900% wall is reached and the base carries the rest. Its campaign advertises Queen White, which runs short before stock lands — re-point the ad to Queen Olive in the same write (operator 2026-09-25) and read the rank the day after, before the next step.'),
    ('bamboo cooling sheets king size / king size sheets set bamboo cooling', 'Traffic gaps that are moving: 28 → 11 and 43 → 18 on the 7-day median. Keep stepping while share rises; they are the evidence the Cooling push buys rank.'),
    ('cal king bamboo sheet set / california king bamboo sheets / bamboo california king sheets', 'Already at 7–11 on the 7-day median with 2–4 top-of-search clicks in 30 days. The target (1–5) is a few positions away, and the push can buy it: grade on impression share and rank. “bamboo california king sheets” has had no top-of-search order on 15 clicks in 90 days — it takes no step until it converts.'),
    ('bamboo sheets cal king', 'The one in-focus term that lost ground: 23 → 37 on the 7-day median. Push and check the cause the same day — the Cal King White hero has 49 units and 120 inbound with no date.'),
    ('king sheets bamboo cooling / bamboo cal king sheets', 'The only in-focus targets near 20% impression share (21.9% and 19.9%). Neither is rank-tracked — add both to the tracker, or the push cannot be graded.'),
]:
    bullet(f'**{kw}** — {txt}')
lead('Out of focus — the hero terms.', 'They hold flat this run (Part B). For the record: “bamboo sheets” is a traffic gap (CTR 3.5% vs 1.9%, CVR 12.7% vs 8.3%, 3.1% impression share at $7.42, rank 38 → 32 in the deal); “bamboo sheets queen size” and “bamboo sheets queen” lost 11–17 positions (36 → 47, 54 → 71) while Queen White ran short, and are frozen; “bamboo sheets king size” is climbing (26 → 22) on a 0.8% share (0.3% in the deal). When the hero sizes come back into focus, the push starts there.')

# ---------------------------------------------------------------- D6 why IS is low at the bid
section('Why a $5–8 effective bid wins only 1–5% of top-of-search')
lead('What we know.', f'Impression share at the top is low on almost every target. The deal gives the read on whether price moves it: after the 15 September raise, share rose on {_up} of {_n} B4 targets with both a 30-day and a deal-window read (“bamboo cooling sheets” 2.6% → 14.3%; “bamboo sheets” 3.1% → 5.1%). Share responds to price; where it did not move, the price was not raised or sits far below the slot. Four causes, each with its test:')
bullet('**Price below the top-slot winners.** The engine’s per-campaign market bound (“clears at $X … bound is $Y”) is built from what we paid, not what wins the slot. A 3% share at $7.42 says the winning price is above it. Test: step the price and read impression share daily; if share climbs with price, it was price — and if share climbs and rank does not (“bamboo cooling sheets” in the deal), the constraint is elsewhere.')
bullet('**Budget cut-off.** Share is counted over the whole day. “Bamboo Cooling Sheets” used 108% of its budget in the last 7 days; a campaign out of budget by evening loses the evening auctions. Test: hours out of budget per campaign; raise budget with the price step where 80%+ is used.')
bullet('**The hero the ad serves.** On Queen the rank loss runs with Queen White’s stock (79 units, next arrival 27 October); on Full and Twin the cooling and bamboo terms fell 30–80 positions in the week Full White and Twin White sold out. A price push cannot fix an ad that sends shoppers to a child that cannot ship.')
bullet('**Eligibility and relevance.** Low share on a high bid can be relevance — but relevance shows as low CTR and CVR, and ours are above the bars. It remains the check on rows still at zero after two steps (A-section daily read).')
lead('Correction for the app.', 'Read top-of-search impression share on every ranking target, every run, beside the effective bid; grade every push on whether share rose. The engine has the field (tosImpressionShare) and does not use it.')

# ---------------------------------------------------------------- D7 competitors
section('Competitors')
riv = _json.load(open(f'{_dg.B}/b6_context.json'))['competitors']['rivals']
rows = []
for r in sorted(riv, key=lambda r: -(r.get('estWeeklySales') or 0))[:12]:
    rows.append([r['brand'][:20], r['tier'], f"${r['price']:.2f}" + (f" (list ${r['listPrice']:.2f})" if r.get('listPrice') else ''), r['rating'], f"{r['reviews']:,}", r['estWeeklySales'], f"{r['bsr']:,}"])
table(['Brand', 'Tier', 'Price', 'Rating', 'Reviews', 'Est. units/wk', 'BSR'], rows, widths=[3.4, 2.2, 3.4, 1.3, 1.8, 1.8, 1.8], size=8)
_cs = A['brief']['sections']['competitors']['summary']
bullet(f"**Prices around us.** The run’s brief watches {_cs['competitorCount']} rivals ({_cs['withHistory']} with history, 24 Aug → 23 Sep): {_cs['priceRaisers']} raised price, {_cs['priceCutters']} cut, {_cs['onDiscountNow']} are on a strike-through discount now. B4’s deal price sits under most of the reference rivals — the reason CVR at the top is above the bar and this is the window to buy volume.")
bullet('**A new low-price entrant.** Shilucheng cut from $69.98 to $31.99–33.90 (15% off) and went from BSR 556,629 to 9,850 with 22,150 new reviews in a month — a price point half of ours. Watch it on the Cooling terms first; it is not on the reference list for any keyword yet.')
bullet('**Who holds the top of the terms.** On the hero terms the top three ASINs take 34–40% of clicks but 6–10% of conversions (ASINsight, 17 Sep): clicks concentrate at the top, purchases do not — the category is won on visibility, and our above-bar CVR at the top means the clicks we buy there convert.')
lead('Correction for the app.', 'Carry each rival’s current price, discount and review count beside the keyword it is the reference for; re-weight the push toward terms where the price gap is widest during a deal; flag rivals without an ASINsight export (their traffic is unmeasured, not zero); add coupon/lightning-deal detection (the brief notes it is not scraped).')

# ---------------------------------------------------------------- D8 more campaign types
section('Campaign types that should also be running')
lead('What the run proposes.', '8 builds: 2 Sponsored Products exact tail rows on the Cooling focus (loadable) and 6 the loader cannot write — Sponsored Brands and Sponsored Brands Video on “bamboo sheets california king” and “bamboo cooling sheets”, Sponsored Display own-page defence, and a product-targeting conquest — written as PAUSED rows for a person to build, with nothing that makes that happen. The stored review held three of them because they feature heroes the run moves campaigns off.')
rows = [
    ['Sponsored Brands headline — exact on the focus terms', 'Builds 2383 (“bamboo sheets california king”, $12.41/day), 2399 (“bamboo cooling sheets”, $4.07)', 'A second top-of-page surface above the SP top slot, in a separate auction. Our above-bar CTR at the top says the creative will be clicked.', 'Build both this week; size budgets to a readable click count (~50 clicks a week each); read new-to-brand and branded-search lift, not only ACoS.'],
    ['Sponsored Brands Video — exact', 'Builds 2375 (Cal King), 2391 (“bamboo cooling sheets”)', 'Video sits mid-search where we have no presence; cooling and deep-pocket fit are what video shows.', 'Build as a 100-click probe each.'],
    ['Sponsored Display — own-page defence', 'Build 2407 (our own ASINs)', 'Product pages still take 140 clicks a day in the deal — rivals’ ads on our pages take shoppers away.', 'Build; keep separate from ranking campaigns so the ranking mix reads clean.'],
    ['SP product targeting — conquest', 'Build 2415 (rival ASINs), on Queen White', 'Most reference rivals price above our deal price; a cheaper, well-rated offer on their page converts. Queen White cannot carry new traffic.', 'Build on King White (2,430 inbound) or Queen Olive, not Queen White; own campaign, own ceiling (contribution × its CVR).'],
    ['SP exact tail on Cooling', 'Builds 2450, 2462 ($12.11/day each)', 'Coverage for converting cooling variants.', 'Load; sized to 15 clicks by the next read.'],
]
table(['Type', 'In the run', 'Why, from the data', 'What to do'], rows, widths=[3.4, 3.8, 5.4, 4.4], size=7.5)
lead('Correction for the app.', 'Non-SP builds need a way to load (or a task created for a person with a due date). Budgets on new formats should be sized to a readable click count in the read window.')

# ---------------------------------------------------------------- D9 metrics
section('Metrics the app should be reading')
rows = [
    ['Top-of-search impression share, per target', 'Measured by Amazon; in Command Center; not read by the engine', f'In-focus median {_st.median(_iss):.1f}%; rose on {_up} of {_n} targets after the 15 Sep raise', 'The direct measure of “are we at the top”; grade every push on it.'],
    ['Required vs delivered top-of-search clicks', 'Run has the requirement; engine does not compare', f'In focus: {_deal:.0f} of {_need:.0f} a day ({_deal / _need * 100:.0f}%)', 'The traffic-gap test; sizes the push.'],
    ['Reachability: delivered ÷ impression share', 'Not computed', 'Most Cal King requirements are above what the whole auction holds', 'Stops a push chasing clicks that do not exist; grade on share instead.'],
    ['Budget hours out', 'Not read', '“Bamboo Cooling Sheets” at 108% of budget', 'Separates price from budget as the cause of low share.'],
    ['The advertised child’s stock against the push', 'Stock rules only', 'Queen White 22 days; Full/Twin White at 0', 'A push on a hero that runs out rents the position; rank losses on Queen/Full/Twin follow the stock-outs.'],
    ['Rank response per 100 top-of-search clicks', 'Not computed', 'Bamboo|Cooling 36 → 26 in the deal', 'Tells the push how much volume a position costs; replaces a fixed step.'],
    ['Rank tracking on every in-focus term', 'Tracker covers most', '3 in-focus terms without a rank read (2 untracked, 1 not returned)', 'An untracked term cannot be graded.'],
    ['TACoS and organic sales trend', 'Daily TACoS available', '13.9% before the deal → 19.4% in it; guard 21.7%', 'The payoff of rank is organic sales; read TACoS falling after the deal as the test that rank moved.'],
    ['Brand click and purchase share (SQP)', 'Field present, empty', '—', 'The rank driver itself: our share of the term’s purchases vs the ranks above us.'],
]
table(['Metric', 'Today', 'Value now', 'Why it matters'], rows, widths=[4, 3.6, 4, 5.4], size=7.5)

# ---------------------------------------------------------------- D10 recommendations
section('What to do')
for b in [
    '**In the deal (to 09-28):** push every in-focus exact ranking campaign on top of search daily by the delivery gap (Part A), inside the loss stop (3× contribution per top-of-search order) and the envelope; stop new steps and hold if weekly TACoS passes 21.7%. Read impression share and clicks every morning; share rising with price = keep stepping; share flat after a step = halve the next step and check budget hours out and the child the ad serves.',
    '**Cooling:** “bamboo cooling sheets” to its loss stop with the base carrying the raise past the 900% wall, re-pointed off Queen White; the moving king-size cooling terms keep stepping while share rises. No extra bid on the broad “cooling sheets” until the offer answers the cooling query.',
    '**California King:** push toward the top slot and grade on share and 7-day rank; the requirement is above what the auction holds on most terms. Investigate “bamboo sheets cal king” (23 → 37) the same day.',
    '**Out of focus:** hold the hero terms flat (the flagship included); freeze the Queen, Full and Twin rank losses and open them as re-ranking cases once stock lands; run the +10% test on the 19 converting out-of-focus campaigns for two weeks.',
    '**New formats:** build the SB headline and SBV on the two focus terms, SD defence, and the conquest row on a child with stock, at readable budgets.',
    '**From 09-29:** terms at their requirement — evaluate cost per click, probing down 3–5% a step; terms short of it — step down to break-even over two writes four days apart, unless the 7-day rank reached the target range (then hold one week and read). No term keeps climbing after the deal.',
]:
    bullet(b)
