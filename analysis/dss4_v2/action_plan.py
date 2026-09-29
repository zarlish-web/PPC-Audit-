"""Deliverable 3 — the action plan: one row per change, each with the data behind it and the reason, in execution order.
Read from the same v2/out/*.json as the document and the review workbook. Output: v2/deliver/DSS4_Action_Plan_20260929.xlsx"""
import json, re, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import load as L

O = L.OUT
J = lambda n: json.load(open(O + n))
D, R, INV, LT, KA, F, AR, P, CF = J('decisions.json'), J('ranking.json'), J('inventory.json'), J('ltsf.json'), J('kw_actions.json'), J('financials.json'), J('audit_review.json'), J('product.json'), J('conflicts.json')
PATH = L.ROOT + 'v2/deliver/DSS4_Action_Plan_20260929.xlsx'
OPP = {o['keyword']: o for o in R['opportunities']}
INVR = {r['sku']: r for r in INV['rows']}
LTR = {r['sku']: r for r in LT['rows']}
SKU = {s['sku']: s for s in P['skus']}
TBC = collections.defaultdict(dict)
for f in ('targets_by_campaign_W30.json', 'targets_by_campaign_W90only.json'):
    for r in json.load(open(L.RAW + 'sb_ppc/' + f)):
        TBC[str(r['_campaign_id'])].setdefault((r.get('Name') or '').lower(), r)
ARC = collections.defaultdict(list)
for a in AR:
    ARC[a['campaign_id']].append(f"{a['kind']} {a['current']}→{a['suggested']}: {a['verdict']}")

wb = Workbook()
wb.remove(wb.active)
HEAD, HF = PatternFill('solid', start_color='1F3864'), Font(color='FFFFFF', bold=True)
STEPF = {1: 'E2EFDA', 2: 'DDEBF7', 3: 'FFF2CC', 4: 'FCE4D6', 5: 'EDEDED'}


def sheet(name, cols, rows, widths, note=None, colour_col=None):
    ws = wb.create_sheet(name[:31])
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = Font(italic=True)
        ws.cell(1, 1).alignment = Alignment(wrap_text=False)
        r0 = 3
    for j, c in enumerate(cols, 1):
        x = ws.cell(r0, j, c)
        x.fill, x.font = HEAD, HF
        x.alignment = Alignment(wrap_text=True, vertical='top')
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, json.dumps(v) if isinstance(v, (list, dict)) else v)
        if colour_col is not None and isinstance(row[colour_col], int) and row[colour_col] in STEPF:
            ws.cell(i, colour_col + 1).fill = PatternFill('solid', start_color=STEPF[row[colour_col]])
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 4)
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(cols))}{r0 + max(1, len(rows))}"
    return ws


pct = lambda x: None if x is None else round(x * 100, 1)

# ================================================================= 1. Action list: one row per change
STEP_OF = {'Ranking – push': 2, 'Ranking – maintain': 2, 'Duplicate of a ranking owner': 3, 'Discovery / LTSF clearance': 4, 'Colour exact (LTSF)': 4}
rows = []
n = 0


def add(step, d, level, entity, change, frm, to, why, ids=('', ''), tos_price=None):
    global n
    n += 1
    m = d['now']
    ch = d['action']['child_to'] or d['child']
    rows.append([n, step, d['campaign_id'], d['name'], d['role'], level, entity, change, frm, to, tos_price, ids[0], ids[1],
                 round(m['clicks_day'] * 30), round(m['orders_day'] * 30), round(m['spend_day'] * 30, 2), round(m['sales_day'] * 30, 2), m['acos'], d['be_acos'] and round(d['be_acos'] * 100, 1),
                 m['tos_cvr90'], m['clicks90'], m['orders90'], d['budget'], d['utilization'] and round(d['utilization'], 1), d['tos_is'],
                 ', '.join(d['owned_terms'][:3]), (OPP.get(d['owned_terms'][0]) or {}).get('rank_now') if d['owned_terms'] else None,
                 (OPP.get(d['owned_terms'][0]) or {}).get('target') if d['owned_terms'] else None, ch, d['checks'].get('stock_child_days'), (LTR.get(ch) or {}).get('rank'),
                 d['checks'].get('d_clicks_day'), d['checks'].get('d_orders_day'), d['checks'].get('d_spend_day'), d['checks'].get('d_profit_day'), why,
                 ' | '.join(ARC.get(d['campaign_id'], []))[:300]])


for d in D['decisions']:
    a = d['action']
    if d['role'] == 'Other product (out of scope)':
        continue
    step = STEP_OF.get(d['role'], 5 if not d['role'].startswith('Colour') else 4)
    why = ' | '.join(d['why'])
    ags = sorted({k.split('|')[0] for k in (r['Id'] for r in TBC.get(d['campaign_id'], {}).values())})
    ag = ags[0] if len(ags) == 1 else ''
    if a['state']:
        add(3 if a['state'] == 'PAUSED' else step, d, 'Campaign', 'Campaign', 'State', d['status'], a['state'], why)
    for sku in a['ads_add']:
        add(step, d, 'Product ad', sku, 'Add product ad (enabled)', '', sku, why, (ag, ''))
    for sku in a['ads_pause']:
        add(step, d, 'Product ad', sku, 'Pause product ad', 'enabled', 'paused', why, (ag, ''))
    for sku in a.get('ads_add_paused') or []:
        add(step, d, 'Product ad', sku, 'Add back-up product ad (PAUSED)', '', sku + ' (paused)', a.get('switch_rule') or '', (ag, ''))
    if a['budget_to']:
        add(step, d, 'Campaign', 'Daily budget', 'Budget', d['budget'], a['budget_to'], why)
    if a['strategy_to'] and a['strategy_to'] != d['strategy']:
        add(step, d, 'Campaign', 'Bidding strategy', 'Bidding strategy', d['strategy'], a['strategy_to'], why)
    for k, lab in (('tos', 'Top of search %'), ('ros', 'Rest of search %'), ('pp', 'Product pages %')):
        if a[k + '_to'] is not None and a[k + '_to'] != d['mods'].get(k):
            add(step, d, 'Placement', lab, 'Placement modifier', d['mods'].get(k), a[k + '_to'], d['checks'].get('placement') or why)
    for b in a['bid_changes']:
        if b.get('term') == '(all targets)':
            for nm, r in TBC.get(d['campaign_id'], {}).items():
                if r.get('current_bid'):
                    add(5, d, 'Target', nm, 'Bid (−15% step)', r['current_bid'], round(r['current_bid'] * 0.85, 2), b.get('reason', ''), tuple((r['Id'].split('|') + [''])[:2]))
        else:
            r = TBC.get(d['campaign_id'], {}).get(b['term'])
            add(step, d, 'Target', b['term'], 'Base bid', b.get('bid_from'), b.get('bid_to'), why, tuple((r['Id'].split('|') + [''])[:2]) if r else ('', ''), b.get('eff_tos_to'))
    for ng in a['negatives_add']:
        add(5, d, 'Negative', ng['term'], 'Add campaign negative exact', '', ng['term'], why)
rows.sort(key=lambda r: (r[1], r[2], r[5]))
for i, r in enumerate(rows, 1):
    r[0] = i
ACOLS = ['#', 'Step', 'Campaign ID', 'Campaign', 'Role', 'Level', 'Entity (target / SKU / setting)', 'Change', 'From', 'To', 'Top-of-search price $ (base × 10)', 'Ad group ID', 'Keyword/target ID',
         'Clicks 30d', 'Orders 30d', 'Spend 30d $', 'Sales 30d $', 'ACoS 30d %', 'Break-even ACoS %', 'TOS CVR 90d %', 'Clicks 90d', 'Orders 90d', 'Budget now $',
         'Budget used %', 'TOS impr. share %', 'Ranking terms owned', 'Rank now', 'Rank target', 'Child (after)', 'Child days of stock', 'Child LTSF rank',
         'Δ clicks/day', 'Δ orders/day', 'Δ spend/day', 'Δ profit/day', 'Why (data behind the change)', 'Audit said → verdict']
sheet('Action list', ACOLS, rows, [6, 6, 17, 48, 22, 11, 30, 26, 14, 20, 12, 17, 17] + [9] * 22 + [110, 70],
      note='Every change, one row each, in execution order (Step 1 stock · 2 ranking · 3 pauses/duplicates · 4 LTSF · 5 non-ranking steps, TOS setup on other exact campaigns, negatives). '
           'Data columns are the campaign\'s current numbers (Command Center 30/90 days, Sellerboard budget/utilisation/TOS share).', colour_col=1)
ACTION_ROWS = len(rows)

# ================================================================= 2. Step 1 — stock and reorder
rows = []
for r in INV['rows']:
    s = SKU[r['sku']]
    role = 'Ranking child' if r['sku'] in {v['chosen'] for v in R['preferred'].values()} else ('Back-up' if any(r['sku'] == b['sku'] for v in R['preferred'].values() for b in v.get('backups') or []) else 'LTSF')
    rows.append([1, r['sku'], s['size'], role, r['sellable'], r['in_transfer'], r['inbound_shipped'], r['inbound_plan_only'], r['stock_total'], r['units_day_30'], r['plan_pace_sep_equiv'],
                 r['stockout_now_pace'] or 'after 31 Jan', r['stockout_date'] or 'after 31 Jan', r['units_short_to_31jan'],
                 ('REORDER ' + f"{r['units_short_to_31jan']:,}" + ' units') if (role == 'Ranking child' and r['units_short_to_31jan'] > 0) else ('—' if r['units_short_to_31jan'] == 0 else 'no reorder: LTSF/aged stock clears')])
sheet('Stock & reorder', ['Step', 'Child', 'Size', 'Role', 'Sellable', 'In FC transfer', 'Inbound shipped', 'Inbound plan only', 'Stock counted', 'Units/day (30d)', 'Plan pace/day',
                          'Stock-out today\'s pace', 'Stock-out under plan', 'Short to 31 Jan', 'Action'], rows, [6, 30, 7, 13] + [11] * 10 + [34],
      note='Sellerboard stock 29 Sep; last year\'s seasonality applied (Oct ×1.25, Nov ×1.87, Dec ×2.49, Jan ×1.54). Confirm Queen Black in Manage FBA Inventory (Data Dive showed 2,328 on 23 Sep).', colour_col=0)

# ================================================================= 3. Back-up switch triggers
rows = []
for b in INV.get('backups', []):
    pr = INVR.get(b['preferred'], {})
    rows.append([b['size'], b['preferred'], pr.get('sellable'), pr.get('plan_pace_sep_equiv'), f"#{b['order']}", b['backup'], b['level'], b['backup_stock'], b['switch_date'] or 'not before 31 Jan',
                 b['ranking_orders_moved'], b['backup_stockout'] or 'after 31 Jan',
                 f"When {b['preferred']} < 14 days of stock or out of stock: enable {b['backup']} product ad in every {b['size']} ranking campaign, pause {b['preferred']}" +
                 (' and drop the push to maintain level' if b['level'].startswith('maintain') else '') + f"; switch back at ≥30 days of {b['preferred']}"])
sheet('Back-up switch', ['Size', 'Preferred child', 'Sellable now', 'Plan pace/day', 'Back-up #', 'Back-up child', 'Can carry', 'Back-up stock', 'Projected switch date',
                         'Ranking orders/day moved', 'Back-up lasts to', 'Trigger and action'], rows, [7, 24, 10, 10, 8, 30, 22, 10, 14, 12, 14, 110],
      note='The back-up product ads are loaded PAUSED now (Action list). Check days of cover daily; the switch is manual unless a rule is set in the ads tool.')

# ================================================================= 4. Ranking terms
rows = [[o['keyword'], o.get('tier'), o['size'], o['child'], o['mkt_purchases_day'], o.get('mkt_clicks_day'), pct(o['our_share']), o['rank_now'], o.get('radar_worst_q3'), o['target'],
         o.get('ad_clicks_day_now'), o.get('target_clicks_day'), o['tos_cpc'], o.get('target_cost_day'), D['owners'].get(o['keyword']) or 'held by discovery (no new campaign)',
         pct(o['paid_cvr']), o['acos30'], o['be_acos'] and round(o['be_acos'] * 100, 1)] for o in R['opportunities'] if o.get('tier')]
sheet('Ranking terms', ['Term', 'Tier', 'Size', 'Ranking child', 'Market purchases/day', 'Market clicks/day', 'Our purchase share %', 'Rank now', 'Worst rank Q3', 'Target',
                        'Paid clicks/day now', 'Target clicks/day', 'TOS CPC $', 'Target cost/day $', 'Owner campaign', 'Paid CVR %', 'ACoS 30d %', 'Break-even %'], rows,
      [28, 9, 7, 24] + [11] * 10 + [22, 9, 9, 9])

# ================================================================= 5. Keyword actions (every term with an action)
rows = [[k['keyword'], k['action'], k['why'], k['cls'], k['size'], k['colour'], k['clicks90'], k['orders90'], k['spend90'], k['acos90'], k['sqp_volume_q3'],
         pct(k['our_purchase_share_q3']), k['rank_radar'], k['owner']] for k in KA if not k['action'].startswith('No change')]
sheet('Keyword actions', ['Search term', 'Action', 'Why', 'Class', 'Size', 'Colour', 'Clicks 90d', 'Orders 90d', 'Spend 90d $', 'ACoS 90d %', 'SQP searches Q3', 'Our share Q3 %',
                          'Rank', 'Owner campaign'], rows, [30, 34, 80, 18, 7, 8] + [10] * 7 + [18],
      note='Negative exact terms go into the discovery (auto/broad/phrase) campaigns. Harvest candidates are added only to an existing campaign of the same size/objective — no new campaigns.')

# ================================================================= 6. LTSF
rows = [[r['rank'], r['sku'], r['size'], r['total_ais'], r['ais_units'], r['oldest_bucket'], r['sellable_now'], r['est_aged_left'], r['avoid_per_unit'], r['pba_unit30'],
         sum(1 for d in D['decisions'] if (d['action']['child_to'] or d['child']) == r['sku'] and d['plan']['spend_day']),
         round(sum(d['plan']['orders_day'] for d in D['decisions'] if (d['action']['child_to'] or d['child']) == r['sku']), 2),
         (INVR.get(r['sku']) or {}).get('stockout_date')] for r in LT['rows']]
sheet('LTSF', ['LTSF # (col V)', 'SKU', 'Size', 'Total AIS Charge $', 'AIS units', 'Oldest bucket (15 Jul)', 'Sellable now', 'Aged left (est.)', 'Storage/AIS avoided per unit $',
               'Profit before ads/unit $', 'Campaigns advertising it (plan)', 'Plan ad orders/day', 'Sells out (plan)'], rows, [8, 30, 7] + [12] * 10,
      note='From the S4-LTSF sheet (verified: $1,632.92 / 3,121 units). LTSF campaign product-ad fixes are in the Action list (Step 4).')

# ================================================================= 7. Owner decisions & monitoring
comp = F['components']
rows = [
    ['Ranking investment', f"Full push adds ${comp['Ranking – push']['d_spend_day'] * 30:,.0f}/month ad spend, direct profit ${comp['Ranking – push']['d_profit_day'] * 30:,.0f}/month",
     'Keep all push terms, or drop terms in the Ranking terms tab (their campaigns then use maintain settings)'],
    ['Spend-share definition', f"Ranking {F['share']['ranking_of_all_spend'] * 100:.0f}% of all spend; {F['share']['ranking_of_keyword_growth_spend'] * 100:.0f}% with LTSF clearance and brand defence ring-fenced",
     f"80% of ALL spend needs non-ranking ≤ ${F['share']['nonranking_cap_for_80pct']:.0f}/day (cuts LTSF clearance)"],
    ['Reorder timing', f"Queen Black out {INVR['SATIN-QUEEN-BLACK']['stockout_date']}, Full Black out {INVR['SATIN-FULL-BLACK']['stockout_date']} (plan pace)", 'Back-ups carry Queen to ~10 Dec, Full to ~20 Nov at maintain level'],
    ['Step size on non-ranking', '−15% per step, re-read after 7 clean days, stop if organic rank slips ≥3', 'A setting, not fixed by the data'],
    ['Back-up switch threshold', '14 days of stock (switch), 30 days (switch back)', 'A setting, not fixed by the data'],
]
sheet('Owner decisions', ['Decision', 'The numbers', 'Options / note'], rows, [28, 90, 80])
rows = [['Daily', 'Ranking clicks vs target (Ranking terms tab) — budget spent but short of target → raise the top-of-search price 10%'],
        ['Daily', 'Top-of-search impression share and placement split on ranking campaigns'],
        ['Daily', 'Organic rank on each push term (Rank Radar)'],
        ['Daily', 'Days of stock on each preferred child — under 14 days → back-up switch'],
        ['Weekly', 'Non-ranking step: re-read ACoS vs break-even after 7 clean days; stop if the term\'s rank slipped ≥3'],
        ['Weekly', 'LTSF: aged units left per SKU; when Queen Grey is ~2 weeks from selling out, move its campaigns to Queen Stone Grey (LTSF #2)'],
        ['Weekly', 'TACoS and net profit vs the plan (Summary tab of the review workbook)']]
sheet('Monitoring', ['When', 'Check'], rows, [10, 140])

# ================================================================= 8. Summary (first tab)
ws = wb.create_sheet('Summary', 0)
cnt = collections.Counter(r[7] for r in wb['Action list'].iter_rows(min_row=4, values_only=True) if r[0])
lines = [('Decolure Satin Sheets 4-Piece — action plan (29 Sep 2026)', ''),
         ('Changes in this plan', f"{ACTION_ROWS:,} rows across {len({r[2] for r in wb['Action list'].iter_rows(min_row=4, values_only=True) if r[0]})} campaigns"),
         *[(f"  {k}", v) for k, v in cnt.most_common()],
         ('Ranking share of spend', f"now {D['summary']['now_ranking_spend_share'] * 100:.0f}% → plan {F['share']['ranking_of_all_spend'] * 100:.0f}% (ring-fenced {F['share']['ranking_of_keyword_growth_spend'] * 100:.0f}%)"),
         ('P&L (30-day)', f"now: sales ${F['now']['sales']:,}, ads ${F['now']['ads']:,}, net ${F['now']['net']:,} · plan: sales ${F['plan']['sales']:,}, ads ${F['plan']['ads']:,}, net ${F['plan']['net']:,}"),
         ('Ranking children', ', '.join(f"{z}: {v['chosen']}" for z, v in R['preferred'].items())),
         ('Back-ups', '; '.join(f"{z}: " + ' → '.join(b['sku'] for b in v.get('backups') or []) for z, v in R['preferred'].items())),
         ('Reorder (to 31 Jan)', ', '.join(f"{s}: {INVR[s]['units_short_to_31jan']:,}" for s in ('SATIN-QUEEN-BLACK', 'SATIN-FULL-BLACK', 'SATIN-TWIN-BLACK-NEW'))),
         ('Execution order', 'Step 1 stock → 2 ranking campaigns → 3 pauses/duplicates → 4 LTSF → 5 other exact TOS set-up, −15% steps, negatives'),
         ('Same data as', 'DSS4_Audit_Review_20260929.docx and DSS4_Audit_Review_Workbook_20260929.xlsx (39 automated cross-checks pass)')]
for i, (a, b) in enumerate(lines, 1):
    ws.cell(i, 1, a).font = Font(bold=(i == 1 or not a.startswith('  ')))
    ws.cell(i, 2, b)
ws.column_dimensions['A'].width = 36
ws.column_dimensions['B'].width = 150
wb.save(PATH)
json.dump(dict(path=PATH, action_rows=ACTION_ROWS), open(O + 'action_plan_meta.json', 'w'))
print(PATH, ACTION_ROWS, dict(cnt.most_common()))
