import pandas as pd, json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
D='/tmp/claude-0/-home-user-PPC-Audit-/6199518d-f662-55e7-a2b4-84935779d092/scratchpad/'
R='/home/user/PPC-Audit-/research/bedding-feasibility/'
CH=R+'charts/'
F='Arial'
H=Font(name=F,bold=True,color='FFFFFF',size=10); HF=PatternFill('solid',fgColor='1F3A5F')
T=Font(name=F,bold=True,size=14); ST=Font(name=F,italic=True,size=9,color='52514E')
N=Font(name=F,size=10); BLUE=Font(name=F,size=10,color='0000FF'); GREEN=Font(name=F,size=10,color='008000'); BOLD=Font(name=F,size=10,bold=True)
YEL=PatternFill('solid',fgColor='FFFF00'); thin=Side(style='thin',color='D0D0D0'); BRD=Border(top=thin,bottom=thin,left=thin,right=thin)
USD='$#,##0.00;($#,##0.00);-'; USD0='$#,##0;($#,##0);-'; PCT='0.0%;(0.0%);-'; NUM='#,##0;(#,##0);-'
wb=Workbook(); wb.remove(wb.active)

def sheet(name,title,sub):
    ws=wb.create_sheet(name[:31]); ws['A1']=title; ws['A1'].font=T; ws['A2']=sub; ws['A2'].font=ST; return ws
def header(ws,row,cols,start=1):
    for i,c in enumerate(cols):
        cell=ws.cell(row=row,column=start+i,value=c); cell.font=H; cell.fill=HF; cell.alignment=Alignment(wrap_text=True,vertical='center'); cell.border=BRD
def widths(ws,w):
    for i,x in enumerate(w): ws.column_dimensions[get_column_letter(i+1)].width=x
def df_to(ws,df,row=4,fmt=None,maxrows=None):
    if maxrows: df=df.head(maxrows)
    header(ws,row,list(df.columns))
    for r_i,rec in enumerate(df.itertuples(index=False)):
        for c_i,v in enumerate(rec):
            if isinstance(v,float) and pd.isna(v): v=None
            c=ws.cell(row=row+1+r_i,column=c_i+1,value=v); c.font=N; c.border=BRD
            if fmt and df.columns[c_i] in fmt: c.number_format=fmt[df.columns[c_i]]
    ws.freeze_panes=ws.cell(row=row+1,column=1)
    return row+1+len(df)
def rows_to(ws,rows,row,cols,fonts=None,fmts=None):
    header(ws,row,cols)
    for i,r in enumerate(rows):
        for j,v in enumerate(r):
            c=ws.cell(row=row+1+i,column=j+1,value=v); c.font=N; c.border=BRD; c.alignment=Alignment(wrap_text=True,vertical='top')
            if fmts and j in fmts: c.number_format=fmts[j]
    return row+1+len(rows)

# ---------------- 1 Data inventory ----------------
ws=sheet('01_Data_Inventory','Data inventory and quality audit','Research date 26 Oct 2026 · primary window YTD 2026 (1 Jan – 26 Oct) · labels: ACTUAL / CALC / EST / ASSUMPTION')
inv=[
['Sellerboard P&L (Ecotero)','US, CA','2023-10 → 2026-10-26 (37 months)','Account-month, child ASIN, parent','Sales, units, orders, refunds, promos, SP/SB/SBV/SD, referral, FBA, storage, LTSF, COGS, gross/net profit, sessions, unit session %','Complete','COGS single all-in figure; parents partially listed (residual rows); ~$91.6k US fees only at account level','Account history, product P&L, refunds, YoY'],
['Command Center PPC','US, CA','Daily 2025 →; search terms from 2025-12-29','Account, product, syntax group, search term, placement, match type','Impressions, clicks, CTR, CPC, spend, orders, sales, CVR, ACoS, TACoS','Complete','7-day attribution, last 7 days settling; match-type coverage partial; off-Amazon CVR withheld','Ad benchmarks, CAC, launch curves'],
['Command Center unit economics','US, CA','Snapshot 2026-10-26 (28–30 day units)','SKU, size, product','Price, referral, FBA fee, landed cost, refund rate, contribution, break-even ACoS','Mostly complete','COGS outliers (e.g. 59 US Queen Light Blue $60.89); missing costs on 39 CA, 32 US SKUs; CA FBA fees high','Factory-to-customer cost model'],
['Command Center aged inventory','US, CA','Age snapshot 2026-10-23','Product, SKU, age bracket','Aged units, projected surcharge, 6-month exposure','Complete (projection)','Projections not bills; CAD FX stale (2025-08-20)','Overstock lessons'],
['Command Center competitors + Keepa','US (CA unmapped)','Roster 2026-10-25; Keepa 2025-07 →','Rival ASIN','Price, reviews, rating, BSR, sales band, coupons, deals, listing assets','Partial','Sales bands ≈1.4× badge method; no CA roster','Competitor leaderboard'],
['Brand Analytics SQP','US, CA','Weekly from 2025-01-05','Query × week','Search volume, impressions, clicks, cart adds, purchases, our share, median price','Sample','Only queries our ASINs appear on','Seasonality, demand trend'],
['Data Dive niches','US (27), CA (2), UK (1 knee pillow, not bedding)','Research dates 2025-06 → 2026-10','Keyword, competitor ASIN','Search volume, suggested bids, rival price/reviews/est. revenue, ranking juice','Partial','Estimates; sibling ASINs repeat parent sales (upper bounds)','Demand sizing, review barrier, CPC'],
['Google Drive — SLQS audit','US','Jul–Aug 2026','Product, SKU','Quilt P&L, landed vs unit cost, refunds by colour, deals','Complete for scope','Single product','New-brand failure case study'],
['External research (fees, duties, markets, reviews)','US, CA, UK','Read 2026-10-26','Market, regulation','Fee tables, tariffs, market sizes, complaint themes','Partial','Official and Amazon pages blocked; search extracts of cited pages','Fees, landed cost, compliance, UK'],
['NOT AVAILABLE','—','—','—','Factory cost sheets, freight quotes, return-reason codes, review text, UK sales/ads, CA competitor roster','Unavailable','—','Open inputs (see 18_Assumptions)']]
rows_to(ws,inv,4,['Data source','Marketplace','Coverage','Granularity','Metrics available','Data quality','Limitations','Research use']); widths(ws,[30,16,26,26,48,16,46,30])

# ---------------- 2 Marketplace comparison ----------------
ws=sheet('02_Marketplace_Compare','Marketplace comparison','US, Canada and UK assessed separately · EST unless marked')
mk=[['Bed linen market','$8.6–10.9bn; 28–34% online','~$1.37bn; ~24% online','~$3.5bn (range $1.8–5.4bn)','EST (Mordor, FMI, ECDB, Grand View)'],
['Demand signal','"cotton sheets" 97k + "100% cotton sheets" 71k searches/month','~8% of US search volume (SQP, ACTUAL)','No internal data','Data Dive; SQP'],
['Price tiers','$13–30 / $40–70 / $80–150','CAD 25–45 / 60–95 / 100+','£6–20 / £11–34 / £40+','Roster, Data Dive, UK snapshots'],
['Main competitors','Bedsure, Pure Bamboo, CGK, California Design Den, LANE LINEN, MooMee','Decolure (ours), Amazon Basics, CGK, unbranded hotel-style','Utopia Bedding, Bedsure, Amazon Basics, Catherine Lansfield, Dreamscene','Roster; research_market.md'],
['CPC','Account $1.79 (ACTUAL); cotton terms $2.21–2.64 suggested','Account CAD 1.93 (ACTUAL)','No data','Command Center; Data Dive'],
['Review barrier (median)','Cotton sets 2,453; pre-washed 492','Bamboo 522; satin 1,406','Leaders 100k+','Data Dive'],
['Duty, cotton set','16.7% incl. 10% Sec. 301','17%','0% (DCTS); 20% import VAT recoverable','research_duties_compliance.md'],
['P1-type contribution before ads','$19.84 (29.8%)','CAD 25.96 (28.4%)','£8.23 (21.1%) duvet set','12_Fee_Contribution'],
['Test capital','$72k / $130k / $200k','+$35–60k incremental','$60–90k','13_Launch_Budget'],
['Difficulty for new brand','High overall; medium in pre-washed cotton','Medium–low','High','Judgement'],
['Recommendation','LAUNCH FIRST','LAUNCH LATER (months 4–6)','TEST SELECTIVELY LATER (months 9–12)','Decision memo']]
rows_to(ws,mk,4,['Dimension','Amazon US','Amazon Canada','Amazon UK','Source']); widths(ws,[30,48,40,44,34])

# ---------------- 3 Category analysis ----------------
ws=sheet('03_Category_Analysis','Category and sub-category analysis (Data Dive niches)','Estimates at each niche research date; revenue totals are upper bounds')
ns=pd.read_csv(D+'data/niche_summary.csv'); ne=pd.read_csv(D+'data/niche_summary_extra.csv')
cols=['label','marketplace','segment','research_date','total_sv_all_mkl','top_kw','top_kw_sv','n_competitors','median_price','reviews_p50','reviews_p75','pct_competitors_lt500_reviews','dd_total_monthly_revenue','dd_top3_revenue_share']
a=ns[cols].copy()
r=df_to(ws,a,4,fmt={'median_price':USD,'dd_total_monthly_revenue':USD0,'pct_competitors_lt500_reviews':PCT,'dd_top3_revenue_share':PCT,'total_sv_all_mkl':NUM,'top_kw_sv':NUM})
ws.cell(row=r+1,column=1,value='Additional niches (competitor pull 26 Oct 2026)').font=BOLD
df_to(ws,ne,r+2)
widths(ws,[38,10,26,12,14,26,10,10,10,10,10,12,14,12])

# ---------------- 4 Opportunity matrix ----------------
ws=sheet('04_Opportunity_Matrix','Product opportunity matrix (risk-adjusted)','Scores 1–5 (5 = best for us). Edit blue scores or weights; totals recalculate.')
ws['A4']='Weights'; ws['A4'].font=BOLD
wts=[('Demand',0.20),('Profit',0.20),('Competition',0.20),('Differentiation',0.15),('Operations',0.15),('Risk',0.10)]
for i,(n,w) in enumerate(wts):
    ws.cell(row=4,column=3+i,value=n).font=BOLD; c=ws.cell(row=5,column=3+i,value=w); c.font=BLUE; c.number_format=PCT; c.fill=YEL
ws['I5']='=SUM(C5:H5)'; ws['I5'].number_format=PCT; ws['I4']='Check = 100%'
opp=[['Pre-washed cotton percale set','US',4,4,4,4,5,4,'492 median reviews; 73% < 500; 29.8% break-even; 16.7% duty'],
['Cotton waffle duvet cover set','US',4,3,4,4,4,3,'504k "duvet cover"; median 131 reviews; Bedsure 90%'],
['Cotton percale set','CA',2,4,4,4,4,4,'~8% of US demand; review barrier 522'],
['Cotton sateen set','US',2,4,3,3,5,4,'"sateen sheets" 6k searches'],
['Cotton fitted sheet (standalone)','US',3,2,4,3,5,4,'Fitted median 501 reviews; $17–39'],
['Solid cotton duvet cover','US',5,3,2,2,5,3,'Median 3,194 reviews'],
['Viscose-from-bamboo set','US',5,5,2,2,3,2,'CPC $6.43; cannibalises Decolure; origin risk'],
['Cooling sheet set','US',4,4,2,3,3,3,'CooLinens $41/order'],
['Organic (GOTS) cotton set','US',3,4,2,4,3,3,'$63.99 median; 2,453 reviews'],
['Quilt set','US',4,2,3,3,3,2,'SLQS lost ~$20k/month'],
['Linen set','US',2,4,3,3,2,3,'Flax imported'],
['Polyester satin set','US',4,2,2,2,4,3,'$26 median; our satin ~0% net'],
['Cotton percale duvet set','UK',3,2,2,3,4,3,'£8.23 contribution after VAT'],
['Silk pillowcase','US',4,4,2,2,1,3,'Silk not core'],
['Waterproof mattress protector','US',4,3,1,2,3,3,'Median 6,073 reviews'],
['Comforter','US',3,3,3,2,2,2,'Bulky FBA'],
['Microfiber generic set','US/UK',5,1,1,1,3,2,'CGK 443k reviews'],
['Latex mattress topper','US',1,3,3,2,1,2,'Not textile; oversize']]
header(ws,7,['Opportunity','Market','Demand','Profit','Competition','Differentiation','Operations','Risk','Weighted score','Rank','Key evidence'])
for i,o in enumerate(opp):
    rr=8+i
    for j,v in enumerate(o[:8]):
        c=ws.cell(row=rr,column=1+j,value=v); c.border=BRD; c.font=BLUE if j>=2 else N
    ws.cell(row=rr,column=9,value=f'=SUMPRODUCT(C{rr}:H{rr},$C$5:$H$5)').number_format='0.00'
    ws.cell(row=rr,column=10,value=f'=RANK(I{rr},$I$8:$I${7+len(opp)})')
    ws.cell(row=rr,column=11,value=o[8]).font=N
ws.conditional_formatting.add(f'I8:I{7+len(opp)}',CellIsRule(operator='greaterThanOrEqual',formula=['3.6'],fill=PatternFill('solid',fgColor='CFE8CF')))
widths(ws,[34,9,9,8,12,14,11,7,14,7,60])

# ---------------- 5 Competitor leaderboard ----------------
ws=sheet('05_Competitor_Leaderboard','Competitor leaderboard by segment','Estimated units/revenue from sales-rank brackets (as of 25 Oct 2026) — ESTIMATES, high side (~1.4× badge method)')
lb=pd.read_csv(D+'data/leaderboard.csv'); r=df_to(ws,lb,4)
lc=pd.read_csv(D+'data/leaderboard_concentration.csv'); ws.cell(row=r+1,column=1,value='Concentration').font=BOLD; df_to(ws,lc,r+2)
widths(ws,[22]+[16]*20)

# ---------------- 6 Competitor product intelligence ----------------
ws=sheet('06_Competitor_Products','Competitor-level product intelligence','Command Center roster (25 Oct 2026) + Data Dive niche competitors; Keepa monthly price/coupon/review history YTD 2026')
cp=pd.read_csv(D+'data/competitors.csv'); r=df_to(ws,cp,4)
ws2=sheet('06b_Keepa_History','Keepa price, coupon and review history (monthly)','Command Center keepa_history, 2025-07 → 2026-10-26; main window 2026')
df_to(ws2,pd.read_csv(D+'data/keepa_samples.csv'),4)
ws3=sheet('06c_Niche_Competitors','Data Dive niche competitors','Estimates; research dates per niche')
x=pd.concat([pd.read_csv(D+'data/niche_competitors.csv'),pd.read_csv(D+'data/niche_competitors_extra.csv')],ignore_index=True,sort=False)
df_to(ws3,x,4)

# ---------------- 7 Keyword demand ----------------
ws=sheet('07_Keyword_Demand','Keyword and search-demand analysis','Data Dive search volume (top 150 per niche) and Brand Analytics SQP monthly market data')
k=pd.read_csv(D+'data/niche_keywords.csv'); r=df_to(ws,k[['niche_label','marketplace','research_date','keyword','search_volume','relevancy','suggested_bid_median_usd','suggested_bid_min_usd','suggested_bid_max_usd']],4,fmt={'search_volume':NUM,'suggested_bid_median_usd':USD})
ws2=sheet('07b_SQP_Market','Search Query Performance — market volume and our share (monthly)','SQP is a sample; read ratios and shares. US and CA. 2025 context, 2026 primary.')
df_to(ws2,pd.read_csv(D+'data/sqp_market.csv'),4)

# ---------------- 8 Reviews and complaints ----------------
ws=sheet('08_Review_Complaints','Review and customer-complaint analysis','Qualitative themes (Amazon blocked review text) + ACTUAL refund rates from Sellerboard YTD 2026')
rv=pd.read_csv(D+'data/review_themes.csv'); r=df_to(ws,rv,4)
ref=[['Bamboo sheet sets','64.8% light / 35.2% dark',0.075,0.062],['Satin sheet sets','23.3% / 76.7%',0.075,0.050],['Cooling sheet sets','45.4% / 54.6%',0.068,0.050],['Satin fitted sheet','27.9% / 72.1%',0.067,0.053],['Quilt sets','70.4% / 29.6%',0.163,0.059]]
ws.cell(row=r+1,column=1,value='Refund rate by colour tone, US, Jan–Oct 2026 (ACTUAL)').font=BOLD
header(ws,r+2,['Family','Unit mix light / dark','Refund % light','Refund % dark/mid','Gap (pts)'])
for i,x in enumerate(ref):
    rr=r+3+i
    for j,v in enumerate(x): c=ws.cell(row=rr,column=1+j,value=v); c.font=N; c.border=BRD; c.number_format=PCT if j>=2 else 'General'
    ws.cell(row=rr,column=5,value=f'=C{rr}-D{rr}').number_format=PCT

# ---------------- 9 Account history ----------------
ws=sheet('09_Account_History','Existing account historical performance (Ecotero bedding)','Sellerboard monthly P&L, Oct 2023 → 26 Oct 2026 (Oct 2026 partial); USD')
m=pd.read_csv(D+'data/sb_monthly.csv'); r=df_to(ws,m,4,fmt={c:USD0 for c in ['Sales','Promo','Ads','SP','SB','SBV','SD','AmazonFees','Referral','FBAFulfil','Storage','LTSF','COGS','GrossProfit','NetProfit','EstPayout']})
ws2=sheet('09b_YoY','Year-on-year windows','Sellerboard; YTD = 1 Jan – 26 Oct in both years'); df_to(ws2,pd.read_csv(D+'data/sb_yoy.csv'),4)
ws3=sheet('09c_Parents_YTD2026','Parent-level P&L YTD 2026','Sellerboard parent rows (RESIDUAL row = children not grouped under a listed parent)'); df_to(ws3,pd.read_csv(D+'data/sb_parent_ytd2026.csv'),4)
ws4=sheet('09d_Aged_Inventory','Aged inventory (181+ days)','Command Center ltsf_portfolio, age snapshot 23 Oct 2026 — projections'); df_to(ws4,pd.read_csv(D+'data/ltsf_snapshot.csv'),4)

# ---------------- 10 Ad benchmarks ----------------
ws=sheet('10_Ad_Benchmarks','Advertising benchmarks','Command Center YTD 2026 (US USD, CA CAD); 7-day attribution')
t=pd.read_csv(D+'data/ppc_product_table_ytd.csv'); t=t[t.spend>0].copy()
cols=['marketplace','currency','product_id','product_name','impressions','clicks','ctr_pct','orders','cvr_pct','spend','ad_sales','acos_pct','cpc','units','tacos_pct','spend_share_pct']
r=df_to(ws,t[cols],4,fmt={'spend':USD0,'ad_sales':USD0,'cpc':USD,'impressions':NUM,'clicks':NUM,'orders':NUM})
ws.cell(row=4,column=len(cols)+1,value='Ad spend per order').font=H; ws.cell(row=4,column=len(cols)+1).fill=HF
for i in range(len(t)): rr=5+i; ws.cell(row=rr,column=len(cols)+1,value=f'=IFERROR(J{rr}/H{rr},0)').number_format=USD
ws2=sheet('10b_MatchType_Objective','Spend by match type, ad product and objective (YTD 2026)','Command Center ppc_performance_split; partial coverage of keyword-targeted spend'); df_to(ws2,pd.read_csv(D+'data/perf_split_matchtype_ytd.csv'),4)
ws3=sheet('10c_Windows','PPC 30/90/180-day and YTD windows','Command Center ppc_summary per product and account'); df_to(ws3,pd.read_csv(D+'data/ppc_windows.csv'),4)
ws4=sheet('10d_Launch_Curves','Weekly launch curves','Command Center ppc_daily aggregated weekly; products 73, 69, 70, 48 and mature 59'); df_to(ws4,pd.read_csv(D+'data/launch_curve.csv'),4)
lv=sheet('10e_Spend_per_Orders','Spend required to generate 10 / 100 / 1,000 orders','CPC and conversion are inputs (blue); cost per order = CPC ÷ CVR')
header(lv,4,['Scenario','CPC ($)','Conversion','Cost per order','10 orders','100 orders','1,000 orders','ACoS at $66.49'])
for i,(n,c,v) in enumerate([('Low',1.80,0.09),('Base',2.40,0.08),('High',3.00,0.07)]):
    rr=5+i; lv.cell(row=rr,column=1,value=n).font=N
    a=lv.cell(row=rr,column=2,value=c); a.font=BLUE; a.number_format=USD
    b=lv.cell(row=rr,column=3,value=v); b.font=BLUE; b.number_format=PCT
    lv.cell(row=rr,column=4,value=f'=B{rr}/C{rr}').number_format=USD
    for j,q in enumerate([10,100,1000]): lv.cell(row=rr,column=5+j,value=f'=D{rr}*{q}').number_format=USD0
    lv.cell(row=rr,column=8,value=f"=D{rr}/'12_Fee_Contribution'!B5").number_format=PCT
widths(lv,[12,10,12,14,12,12,14,14])

# ---------------- 11 Landed cost & 12 Fee/contribution (formulas) ----------------
inp=json.load(open(D+'model/inputs.json')); P=inp['P']; FX=inp['FX']
ws=sheet('11_Landed_Cost','Factory cost and landed-cost model','Blue = input (ASSUMPTION unless noted). FOB from vendor guide pending mill cost sheets; uplift anchored to SLQS audit (40–44%).')
ws['A4']='FX (EST, 23–26 Oct 2026)'; ws['A4'].font=BOLD
for i,(k,v) in enumerate([('USD/CAD',FX['USD_CAD']),('GBP/USD',FX['GBP_USD']),('USD/PKR',FX['USD_PKR'])]):
    ws.cell(row=5+i,column=1,value=k); c=ws.cell(row=5+i,column=2,value=v); c.font=BLUE
hdr=['Product','Marketplace','Currency','Size','Mix','FOB (USD)','FOB (local)','Freight+duty uplift (x)','Landed (local)','Duty rate (info)','Notes']
header(ws,10,hdr); rr=11; landed_rows={}
duty={'US':'16.7% cotton (6.7% MFN + 10% Sec.301)','CA':'17% MFN','UK':'0% DCTS'}
for pid,p in P.items():
    first=rr
    for s,mx in p['mix'].items():
        ws.cell(row=rr,column=1,value=p['name']); ws.cell(row=rr,column=2,value=p['mkt']); ws.cell(row=rr,column=3,value=p['cur']); ws.cell(row=rr,column=4,value=s)
        c=ws.cell(row=rr,column=5,value=mx); c.font=BLUE; c.number_format=PCT
        c=ws.cell(row=rr,column=6,value=p['fob_usd'][s]); c.font=BLUE; c.number_format=USD; c.fill=YEL
        conv={'USD':f'=F{rr}','CAD':f'=F{rr}*$B$5','GBP':f'=F{rr}/$B$6'}[p['cur']]
        ws.cell(row=rr,column=7,value=conv).number_format=USD
        c=ws.cell(row=rr,column=8,value=p['landed_mult']); c.font=BLUE; c.number_format='0.00"x"'
        ws.cell(row=rr,column=9,value=f'=G{rr}*H{rr}').number_format=USD
        ws.cell(row=rr,column=10,value=duty[p['mkt']]); ws.cell(row=rr,column=11,value='FOB: Karachi vendor guide 250TC percale $12.40–14.80 (search extract)' if 'percale' in pid else 'FOB: no public benchmark')
        rr+=1
    landed_rows[pid]=(first,rr-1); rr+=1
widths(ws,[52,11,9,12,8,11,11,14,13,30,50])

ws=sheet('12_Fee_Contribution','Amazon fee and contribution model (per unit, blended size mix)','Formulas reference 11_Landed_Cost. Contribution = price/(1+VAT)×(1−refund) − referral×(1−refund) − FBA − landed − storage − unsellable − promo')
ids=list(P.keys())
header(ws,4,['Line item']+[P[i]['name'][:40]+f" ({P[i]['mkt']}, {P[i]['cur']})" for i in ids]+['Type / source'])
lines=[('Blended price',None,'ASSUMPTION (target price)'),('VAT rate',None,'UK 20%; US/CA 0% (sales tax outside price)'),('Refund rate',None,'ACTUAL proxy: our sheet families 5.4–7.5%'),('Referral rate',None,'EST 15% Home'),('FBA fee (blended)',None,'ACTUAL proxy / EST'),('Storage + inbound placement',None,'CALC from US $0.78/cu ft Jan–Sep, $2.40 Oct–Dec'),('Unsellable share of refunds',None,'ASSUMPTION'),('Promotions allowance (% price)',None,'ASSUMPTION'),
('Landed cost (blended)','landed','From 11_Landed_Cost tab'),('Net revenue after VAT & refunds','f_net','CALC'),('Referral fee (net of refunds)','f_ref','CALC'),('Unsellable returns cost','f_uns','CALC'),('Promotions','f_promo','CALC'),
('CONTRIBUTION BEFORE ADS','f_contrib','CALC'),('Contribution margin = break-even ACoS','f_be','CALC'),('Break-even ROAS','f_roas','CALC'),('Target ACoS for 10% margin','f_t10','CALC'),('Target ACoS for 15% margin','f_t15','CALC'),
('Launch CPC',None,'CALC from internal + Data Dive'),('Launch conversion',None,'CALC'),('Launch ad spend per order','f_cpo','CALC'),('Contribution after launch-phase ads','f_after','CALC'),('Mature TACoS',None,'ASSUMPTION 12%'),('Contribution at mature TACoS','f_mature','CALC'),('Units to repay $27k one-time + $35k ads (pre-ads contribution)','f_units','CALC')]
rowmap={}
for i,(n,k,src) in enumerate(lines):
    rowmap[n]=5+i; ws.cell(row=5+i,column=1,value=n).font=BOLD if n.isupper() else N; ws.cell(row=5+i,column=2+len(ids),value=src).font=ST
for j,pid in enumerate(ids):
    col=get_column_letter(2+j); p=P[pid]; a,b=landed_rows[pid]
    price=sum(p['mix'][s]*p['price'][s] for s in p['mix']); fba=sum(p['mix'][s]*p['fba'][s] for s in p['mix'])
    vals={'Blended price':price,'VAT rate':p['vat'],'Refund rate':p['refund'],'Referral rate':p['referral'],'FBA fee (blended)':fba,'Storage + inbound placement':p['storage_inbound'],'Unsellable share of refunds':p['unsellable_share'],'Promotions allowance (% price)':p['promo'],'Launch CPC':p['cpc'],'Launch conversion':p['cvr'],'Mature TACoS':0.12}
    for n,v in vals.items():
        c=ws[f'{col}{rowmap[n]}']; c.value=v; c.font=BLUE; c.number_format=PCT if ('rate' in n.lower() or '%' in n or 'share' in n.lower() or 'conversion' in n.lower() or 'TACoS' in n) else USD
    R_=lambda n: f'{col}{rowmap[n]}'
    f={'Landed cost (blended)':f"=SUMPRODUCT('11_Landed_Cost'!E{a}:E{b},'11_Landed_Cost'!I{a}:I{b})",
       'Net revenue after VAT & refunds':f'={R_("Blended price")}/(1+{R_("VAT rate")})*(1-{R_("Refund rate")})',
       'Referral fee (net of refunds)':f'={R_("Blended price")}*{R_("Referral rate")}*(1-{R_("Refund rate")})',
       'Unsellable returns cost':f'={R_("Refund rate")}*{R_("Unsellable share of refunds")}*{R_("Landed cost (blended)")}',
       'Promotions':f'={R_("Blended price")}*{R_("Promotions allowance (% price)")}',
       'CONTRIBUTION BEFORE ADS':f'={R_("Net revenue after VAT & refunds")}-{R_("Referral fee (net of refunds)")}-{R_("FBA fee (blended)")}-{R_("Landed cost (blended)")}-{R_("Storage + inbound placement")}-{R_("Unsellable returns cost")}-{R_("Promotions")}',
       'Contribution margin = break-even ACoS':f'={R_("CONTRIBUTION BEFORE ADS")}/{R_("Blended price")}',
       'Break-even ROAS':f'=IFERROR(1/{R_("Contribution margin = break-even ACoS")},0)',
       'Target ACoS for 10% margin':f'={R_("Contribution margin = break-even ACoS")}-0.10','Target ACoS for 15% margin':f'={R_("Contribution margin = break-even ACoS")}-0.15',
       'Launch ad spend per order':f'={R_("Launch CPC")}/{R_("Launch conversion")}',
       'Contribution after launch-phase ads':f'={R_("CONTRIBUTION BEFORE ADS")}-{R_("Launch ad spend per order")}',
       'Contribution at mature TACoS':f'={R_("CONTRIBUTION BEFORE ADS")}-{R_("Blended price")}*{R_("Mature TACoS")}',
       'Units to repay $27k one-time + $35k ads (pre-ads contribution)':f'=IFERROR(62000/{R_("CONTRIBUTION BEFORE ADS")},0)'}
    for n,v in f.items():
        c=ws[R_(n)]; c.value=v; c.number_format=PCT if ('ACoS' in n) else ('0.00' if 'ROAS' in n else (NUM if 'Units' in n else USD))
        if n.isupper(): c.font=BOLD
widths(ws,[46,24,24,24,24,46])
# sensitivity table (values from model)
s=pd.read_csv(D+'model/out_sensitivity_p1.csv'); r0=6+len(lines)+1
ws.cell(row=r0,column=1,value='Sensitivity of P1 contribution before ads (model.py)').font=BOLD; df_to(ws,s,r0+1)
ws.freeze_panes=None

# ---------------- 13 Launch budget ----------------
ws=sheet('13_Launch_Budget','Launch-budget scenarios (one product family, US)','Blue = inputs. Working capital vs expenses separated. Amazon fees are inside unit economics.')
header(ws,4,['Item','Low','Base','High','Type','Notes'])
items=[('Initial units',1500,2500,4000,'Input','12 SKUs'),('Landed cost per unit',None,None,None,'Link','from 12_Fee_Contribution'),('Initial inventory, landed',None,None,None,'Working capital','units × landed'),
('Samples and lab testing (incl. OEKO-TEX)',3000,6000,9000,'One-time',''),('Branding, packaging design, trademarks',3000,5000,7000,'One-time','USPTO $350/class ×2'),('Photography, video, A+ content',3000,6000,10000,'One-time',''),('Listing preparation and localisation',1000,1500,2000,'One-time',''),('Importer of record, bond, broker, 3PL set-up',2500,4500,6000,'One-time','EO 14411 risk'),('Vine and launch coupons',2000,4000,6000,'Launch expense','Vine $200 for 11–30 units'),('PPC, first 90 days',20000,35000,45000,'Launch expense',''),
('Subtotal',None,None,None,'Calc',''),('Contingency %',0.10,0.15,0.20,'Input',''),('Contingency',None,None,None,'Reserve',''),('CAPITAL TO TEST',None,None,None,'Calc',''),('Of which working capital',None,None,None,'Calc',''),('Of which one-time + launch expense',None,None,None,'Calc',''),('First reorder (units × landed), funded from sales + reserve',None,None,None,'Working capital','base reorder month 4')]
for i,(n,lo,ba,hi,ty,no) in enumerate(items):
    rr=5+i; ws.cell(row=rr,column=1,value=n).font=BOLD if n.isupper() else N; ws.cell(row=rr,column=5,value=ty); ws.cell(row=rr,column=6,value=no)
    for j,v in enumerate([lo,ba,hi]):
        col=get_column_letter(2+j); c=ws[f'{col}{rr}']
        if v is not None: c.value=v; c.font=BLUE; c.number_format=PCT if n=='Contingency %' else (NUM if n=='Initial units' else USD0)
for col in 'BCD':
    ws[f'{col}6']="='12_Fee_Contribution'!B13"; ws[f'{col}6'].font=GREEN; ws[f'{col}6'].number_format=USD
    ws[f'{col}7']=f'={col}5*{col}6'; ws[f'{col}7'].number_format=USD0
    ws[f'{col}15']=f'=SUM({col}7:{col}14)'; ws[f'{col}15'].number_format=USD0
    ws[f'{col}17']=f'={col}15*{col}16'; ws[f'{col}17'].number_format=USD0
    ws[f'{col}18']=f'={col}15+{col}17'; ws[f'{col}18'].number_format=USD0; ws[f'{col}18'].font=BOLD
    ws[f'{col}19']=f'={col}7'; ws[f'{col}19'].number_format=USD0
    ws[f'{col}20']=f'=SUM({col}8:{col}14)'; ws[f'{col}20'].number_format=USD0
    ws[f'{col}21']=f'={col}5*{col}6'; ws[f'{col}21'].number_format=USD0
ws['A23']='Capital to scale (months 6–18, if gates pass): about $200–350k — inventory cover for ~1,500–2,000 units/month ($130–170k), Canada ($35–60k), UK test ($60–90k), peak-season ads.'; ws['A23'].font=ST
widths(ws,[52,14,14,14,16,30])

# ---------------- 14 Profit scenarios (formulas) ----------------
ws=sheet('14_Profit_Scenarios','Product-level profitability scenarios — 12 months from first sale (P1 US + P2 from month 5)','Blue rows are inputs per month; all other rows are formulas referencing 12_Fee_Contribution. Peak cash comes from model.py (inventory timing).')
from model_import import SCEN
rr=4
for name,s in SCEN.items():
    ws.cell(row=rr,column=1,value=f'{name} scenario').font=Font(name=F,bold=True,size=12)
    header(ws,rr+1,['Metric']+[f'M{m}' for m in range(1,13)]+['12-month total'])
    labels=['Sheet sets per day','Price discount','Conversion','CPC','Organic share of units','Duvet sets as % of sheet units (from M5)',
            'Sheet units','Duvet units','Revenue','Paid units','Ad spend','ACoS (ad spend / revenue)','Contribution after ads','Cumulative contribution']
    for i,l in enumerate(labels): ws.cell(row=rr+2+i,column=1,value=l).font=N
    for m in range(12):
        c=get_column_letter(2+m); base=rr+2
        vals=[s['ramp'][m],s['price_disc'][m],s['cvr'][m],s['cpc'],s['organic'][m],(0.25 if (s['p2'] and m>=4) else 0)]
        for i,v in enumerate(vals):
            cell=ws[f'{c}{base+i}']; cell.value=v; cell.font=BLUE; cell.number_format=PCT if i in (1,2,4,5) else (USD if i==3 else NUM)
        ws[f'{c}{base+6}']=f'={c}{base}*30'; ws[f'{c}{base+7}']=f'=ROUND({c}{base+6}*{c}{base+5},0)'
        ws[f'{c}{base+8}']=f"={c}{base+6}*'12_Fee_Contribution'!$B$5*(1-{c}{base+1})+{c}{base+7}*'12_Fee_Contribution'!$C$5"
        ws[f'{c}{base+9}']=f'=({c}{base+6}+{c}{base+7})*(1-{c}{base+4})'
        ws[f'{c}{base+10}']=f'={c}{base+9}*{c}{base+3}/{c}{base+2}'
        ws[f'{c}{base+11}']=f'=IFERROR({c}{base+10}/{c}{base+8},0)'
        ws[f'{c}{base+12}']=f"={c}{base+6}*('12_Fee_Contribution'!$B$18-'12_Fee_Contribution'!$B$5*{c}{base+1}*(1-'12_Fee_Contribution'!$B$8))+{c}{base+7}*'12_Fee_Contribution'!$C$18-{c}{base+10}"
        prev=f'{get_column_letter(1+m)}{base+13}' if m>0 else '0'
        ws[f'{c}{base+13}']=f'={prev}+{c}{base+12}' if m>0 else f'={c}{base+12}'
        for i in (6,7,9): ws[f'{c}{base+i}'].number_format=NUM
        for i in (8,10,12,13): ws[f'{c}{base+i}'].number_format=USD0
        ws[f'{c}{base+11}'].number_format=PCT
    tc='N'
    for i in (6,7,8,9,10,12):
        ws[f'{tc}{rr+2+i}']=f'=SUM(B{rr+2+i}:M{rr+2+i})'; ws[f'{tc}{rr+2+i}'].number_format=USD0 if i in (8,10,12) else NUM; ws[f'{tc}{rr+2+i}'].font=BOLD
    ws[f'{tc}{rr+2+11}']=f'=IFERROR(N{rr+2+10}/N{rr+2+8},0)'; ws[f'{tc}{rr+2+11}'].number_format=PCT
    rr+=2+len(labels)+2
summ=pd.read_csv(D+'model/out_scenarios_summary.csv')
ws.cell(row=rr,column=1,value='Model summary incl. inventory timing, reorders and peak cash (model.py; values)').font=BOLD
df_to(ws,summ,rr+1,fmt={'revenue_12m':USD0,'ad_spend_12m':USD0,'contribution_12m':USD0,'one_time_costs':USD0,'initial_inventory_landed':USD0,'peak_cash_need':USD0,'tacos_12m':PCT,'contingency_pct':PCT})
ws.freeze_panes='B4'; widths(ws,[40]+[11]*12+[14])
img=XLImage(CH+'07_scenario_contribution.png'); img.width,img.height=720,400; ws.add_image(img,f'P4')

# ---------------- 15 Logistics & compliance ----------------
ws=sheet('15_Logistics_Compliance','Country-specific logistics, duties and compliance','EST — search extracts of official pages read 26 Oct 2026; confirm with licensed broker')
lg=[['Cotton sheet set duty','6.7% MFN + 10% Sec. 301 = 16.7% (+0.3464% MPF, 0.125% HMF)','17% MFN (GPT excludes ch. 63)','0% DCTS Enhanced (12% MFN)','HTS 6302.31.90 / CBSA ch.63 / UK Trade Tariff'],
['Viscose or polyester set duty','11.4% + 10% = 21.4%','17% (unconfirmed)','0%','6302.32.20'],
['Cotton with embroidery/piping','20.9% + 10% = 30.9%','17%','0%','6302.31.50 — avoid trim'],
['Poly comforter / quilt','12.8% + 10% = 22.8%','14%','0% (unconfirmed)','9404.40.90'],
['Origin rule','Where fabric is woven','Where fabric is woven','DCTS origin rules; weave in-house','19 CFR 102.21 (unverified)'],
['Import tax','Sales tax collected by Amazon (marketplace facilitator)','GST 5% / HST, recoverable','Import VAT 20%, postponed VAT accounting','—'],
['Importer of record','Non-resident IOR possible; EO 14411 tightening → US LLC or third-party IOR','Non-resident importer + own CARM security','GB EORI + UK VAT registration','—'],
['Fibre & origin labelling','FTC Textile Act; "Made in Pakistan"; RN','Bilingual EN/FR fibre content; CA ID number','Textile Labelling Regs 2012','—'],
['Bamboo claims','"Rayon/viscose made from bamboo" (FTC $5.5M penalties 2022)','Viscose/rayon disclosure','"Viscose" — bamboo not a fibre name','—'],
['Flammability','Sheets: none; mattress pads/toppers 16 CFR 1632','Textile Flammability Regs','Fire regs cover mattresses/pillows/fillings, not sheets','—'],
['Chemicals','Prop 65; state PFAS bans (CA, ME, MN, CO, VT)','CCPSA','UK REACH','—'],
['Ocean transit (ASSUMPTION)','25–35 days West Coast; 35–45 East Coast','35–45 days','25–35 days','Quotes required'],
['Production lead time (ASSUMPTION)','30–45 days after sample approval','same','same','Mill to confirm'],
['Cash cycle (ASSUMPTION)','140–200 days','similar','similar','—'],
['Amazon prep','Prep & labelling services ended 1 Jan 2026 — FNSKU at factory','—','—','research_fees_ads.md']]
rows_to(ws,lg,4,['Topic','US','Canada','UK','Source / note']); widths(ws,[30,52,40,44,34])

# ---------------- 16 Roadmap ----------------
ws=sheet('16_Launch_Roadmap','Launch roadmap','Dates assume decision 2 Nov 2026; anything after week 6 depends on samples, mill lead time, Amazon approvals and freight')
rm=[['1','Phase 1 Validation','Brief; FOB by size; 4-colour samples; trademark (24, 20); US LLC/IOR','Cost sheet requested'],['2','Phase 1','Lab tests booked; freight/broker/bond/3PL quotes; packaging brief','Quotes'],['3','Phase 1','Samples measured and weighed; FBA fee previews; listing copy and keyword map','Real FBA fee'],
['4','Phase 1','Lab tests running; photography plan','—'],['5','Phase 1','Lab results; final FOB; rebuild model','Contribution ≥ 25%?'],['6','GATE 1','Go/no-go on production; PO 2,500 + 300 air units','PO or stop'],
['7–8','Phase 2 Launch prep','Production; FNSKU labels at mill; Brand Registry','Brand Registry'],['9','Phase 2','Photo/video/A+; listings built','Listings ready'],['10','Phase 2','Air-ship 300 units; Vine enrolment','Inbound created'],
['11–12','Phase 2','300 units live; Vine; auto/broad at $150–250/day','Conversion read on 500+ clicks'],['12–13','Phase 2','Inspection; ocean booking; ISF','Main batch ships'],['18–20','Phase 2','Main stock live; 90-day $35k ad budget starts','—'],
['Month 4 of selling','GATE 2','Conversion ≥ 8%; refunds ≤ 7%; cost/order ≤ $25; contribution after ads ≥ 0','Scale or stop'],['Months 4–6','Phase 3 Optimisation','Search-term pruning; price test; image refresh; reorder','—'],
['Months 5–9','Phase 4 Expansion','Waffle duvet set; Canada listing; Prime Day 2027','—'],['Month 9','GATE 3','US TACoS ≤ 22%; organic ≥ 35% → decide Canada scale and UK test','—'],['Months 9–12','Phase 4','BFCM 2027 stock (reorder by early Aug 2027)','—']]
rows_to(ws,rm,4,['Week / month','Phase','Actions','Output / gate']); widths(ws,[18,22,80,30])

# ---------------- 17 Risk register ----------------
ws=sheet('17_Risk_Register','Risk register','Likelihood/impact H/M/L (judgement); score = L×I with H=3, M=2, L=1')
rk=[['New-brand conversion < 7%','M','H','Tested-fit claim, Vine, median price, small range','Conversion on first 1,500 clicks'],['Cost per order > $30','H','H','Broad/phrase discovery, long-tail exact, no blanket TOS premium','Weekly cost/order'],['Overstock / aged surcharges','M','H','2,500-unit first order; 4 colours; 3PL buffer','Days of cover > 150'],['Sec. 301 tariff persists/rises','M','M','Price for 16.7%; keep entries open','CIT ruling'],['IOR rules (EO 14411)','M','M','US LLC / third-party IOR','CBP guidance'],['Shrinkage/pilling returns','M','H','AATCC 135 / ASTM D3512 per batch','Refunds > 7%'],['Colour-mismatch returns','M','M','Calibrated photos; colour tolerance','Refunds by colour'],['Non-Pakistani fabric origin','L','H','Weave in-house; records','Supplier audit'],['Incumbent price/ad war','M','M','Compete on spec; margin floor','Price index'],['Cannibalisation of Decolure','L','M','Cotton only; negatives','Search-term overlap'],['Claims enforcement','L','M','No cooling/bamboo/organic claims without proof','Suppressions'],['Freight disruption','M','M','Two forwarders; buffer','Transit time'],['Cash cycle longer than plan','M','M','Air starter batch; staged POs','Peak cash vs plan'],['UK complexity for thin margin','M','M','No UK before Gate 3','UK gate']]
header(ws,4,['Risk','Likelihood','Impact','Score','Mitigation','Early warning metric'])
for i,x in enumerate(rk):
    rr=5+i
    for j,v in enumerate(x[:3]): ws.cell(row=rr,column=1+j,value=v).font=N
    ws.cell(row=rr,column=4,value=f'=(IF(B{rr}="H",3,IF(B{rr}="M",2,1)))*(IF(C{rr}="H",3,IF(C{rr}="M",2,1)))')
    ws.cell(row=rr,column=5,value=x[3]); ws.cell(row=rr,column=6,value=x[4])
dv=DataValidation(type='list',formula1='"H,M,L"'); ws.add_data_validation(dv); dv.add(f'B5:C{4+len(rk)}')
ws.conditional_formatting.add(f'D5:D{4+len(rk)}',CellIsRule(operator='greaterThanOrEqual',formula=['6'],fill=PatternFill('solid',fgColor='F4CCCC')))
widths(ws,[34,11,9,8,54,30])

# ---------------- 18 Assumptions & evidence ----------------
ws=sheet('18_Assumptions_Evidence','Assumptions and evidence register','Replace ASSUMPTION rows first; every external figure carries a source and date')
ae=[['Factory FOB percale set F/Q/K','$13.00 / $14.00 / $15.50','ASSUMPTION','Karachi vendor guide 250TC $12.40–14.80 (search extract 26 Oct 2026)','Mill cost sheet'],['FOB waffle duvet Q/K','$17.50 / $19.50','ASSUMPTION','No benchmark','Mill cost sheet'],['US uplift','×1.40 cotton / ×1.45 viscose','ASSUMPTION anchored to ACTUAL','SLQS audit landed 40–44% above unit cost','Forwarder + broker quotes'],['CA / UK uplift','×1.38 / ×1.18','ASSUMPTION','Duty 17% / 0% + freight','Quotes'],
['US duty cotton / viscose','16.7% / 21.4%','EST','HTS mirrors + Sec. 301 (unisco.com, htshub.com; bakerdonelson.com)','Broker'],['Referral','15%','EST','Amazon fee pages via search extracts','Seller Central'],['FBA sheet set','$8.80–9.20','ACTUAL proxy','Our bamboo 4-pc avg $9.06','Fee preview'],['Refund rate','6.5% US / 7.5% CA / 5% UK','ACTUAL proxy','Sellerboard YTD 2026','Live'],
['Launch CPC / CVR','$2.40 / 8%','CALC','Command Center launches; Data Dive bids','First 1,500 clicks'],['Organic share M12','30% / 45% / 55%','ASSUMPTION','Mature account ~57%','Business reports'],['FX','USD/CAD 1.3993; GBP/USD 1.3367; USD/PKR 276.82','EST','Bank of Canada / SBP 23–26 Oct 2026','Rate at PO'],['Rival sales','Brackets / Data Dive','EST','Roster 25 Oct 2026; Data Dive','—'],['Market sizes','US $8.6–10.9bn; CA ~$1.37bn; UK ~$3.5bn','EST','Mordor, FMI, Grand View, ECDB','—']]
rows_to(ws,ae,4,['Input','Value used','Type','Source / date','Replace with']); widths(ws,[30,40,26,60,26])

# ---------------- 19 Shortlist ----------------
ws=sheet('19_Final_Shortlist','Final product shortlist','Ranked by risk-adjusted attractiveness')
sl=[['1','Pre-washed 100% cotton percale 4-pc set','US','Launch first','F/Q/K × White, Stone, Slate Grey, Navy (12 SKUs)','$59.99 / $64.99 / $69.99',"='12_Fee_Contribution'!B13","='12_Fee_Contribution'!B18","='12_Fee_Contribution'!B19",'$35k (20–45k)','2,500 units','Medium'],
['2','100% cotton waffle duvet cover set','US','Month 4–5 after Gate 2','Q/K × White, Oat, Charcoal (6 SKUs)','$74.99 / $79.99',"='12_Fee_Contribution'!C13","='12_Fee_Contribution'!C18","='12_Fee_Contribution'!C19",'$10–15k','600–900 units','Medium–low'],
['3','Pre-washed cotton percale set','CA','Months 4–6, same SKUs','Best 3 US colours × 3 sizes','CAD 84.99 / 89.99 / 94.99',"='12_Fee_Contribution'!D13","='12_Fee_Contribution'!D18","='12_Fee_Contribution'!D19",'CAD 10–15k','500–800 units','Medium'],
['4','Cotton percale duvet set + fitted sheet (UK sizes)','UK','Test months 9–12 after Gate 3','Double/King/Super King × 3 colours','£34.99 / 39.99 / 44.99',"='12_Fee_Contribution'!E13","='12_Fee_Contribution'!E18","='12_Fee_Contribution'!E19",'£10–15k','600–1,000 units','Low']]
header(ws,4,['Rank','Product','Market','When','Sizes × colours','Target price','Landed cost','Contribution before ads','Break-even ACoS','Launch ads (90 d)','Initial inventory','Confidence'])
for i,x in enumerate(sl):
    for j,v in enumerate(x):
        c=ws.cell(row=5+i,column=1+j,value=v); c.font=GREEN if str(v).startswith('=') else N; c.border=BRD
        if j in (6,7): c.number_format=USD
        if j==8: c.number_format=PCT
ws['A11']='Avoid / postpone: bamboo viscose (cannibalises Decolure; CPC $6.43), polyester satin (~0% net), microfiber (price war), quilts & comforters (bulky; SLQS loss), toppers (non-textile), mattress protectors (6,073 median reviews), linen (flax imported), silk pillowcases (not core), cooling (CooLinens $41/order).'; ws['A11'].font=ST
widths(ws,[6,42,8,24,40,24,12,14,12,14,14,12])

# ---------------- 20 Go/no-go dashboard ----------------
ws=sheet('20_GoNoGo_Dashboard','Go / no-go decision dashboard','Enter actuals in the yellow column; status updates automatically')
gates=[['Gate 1 — production (week 6)','Factory FOB blended (USD)','<=',16.00,None],['Gate 1','Shrinkage after 5 washes','<=',0.03,None],['Gate 1','Pilling grade (ASTM D3512)','>=',3.5,None],['Gate 1','Landed cost per unit (USD)','<=',22.00,None],['Gate 1','Contribution before ads (% price)','>=',0.25,None],['Gate 1','FBA fee per set (USD)','<=',10.50,None],
['Gate 2 — month 4 of selling','Conversion (on ≥ 3,000 clicks)','>=',0.08,None],['Gate 2','Refund rate','<=',0.07,None],['Gate 2','Average rating','>=',4.3,None],['Gate 2','Ad spend per order (USD)','<=',25.00,None],['Gate 2','Contribution after ads per unit, 4 weeks (USD)','>=',0.00,None],['Gate 2','Organic share of units','>=',0.25,None],
['Gate 3 — month 9','US TACoS','<=',0.22,None],['Gate 3','Organic share of units','>=',0.35,None],['Gate 3','Months of positive contribution in a row','>=',3,None]]
header(ws,4,['Gate','Metric','Rule','Threshold','Actual (enter)','Status'])
pct_rows=('Shrinkage','Contribution before ads','Conversion','Refund','TACoS','Organic')
for i,(g,mn,op,th,ac) in enumerate(gates):
    rr=5+i; ws.cell(row=rr,column=1,value=g); ws.cell(row=rr,column=2,value=mn); ws.cell(row=rr,column=3,value=op)
    c=ws.cell(row=rr,column=4,value=th); c.font=BLUE; c.number_format=PCT if mn.startswith(pct_rows) else ('0.0' if 'grade' in mn or 'rating' in mn or 'Months' in mn else USD)
    a=ws.cell(row=rr,column=5); a.fill=YEL; a.number_format=c.number_format
    cond=f'E{rr}<=D{rr}' if op=='<=' else f'E{rr}>=D{rr}'
    ws.cell(row=rr,column=6,value=f'=IF(E{rr}="","PENDING",IF({cond},"PASS","FAIL"))')
n=4+len(gates)
ws.conditional_formatting.add(f'F5:F{n}',FormulaRule(formula=['F5="PASS"'],fill=PatternFill('solid',fgColor='CFE8CF')))
ws.conditional_formatting.add(f'F5:F{n}',FormulaRule(formula=['F5="FAIL"'],fill=PatternFill('solid',fgColor='F4CCCC')))
ws.cell(row=n+2,column=1,value='Gate 1 decision').font=BOLD; ws.cell(row=n+2,column=2,value='=IF(COUNTIF(F5:F10,"FAIL")>0,"NO-GO / REWORK",IF(COUNTIF(F5:F10,"PENDING")>0,"PENDING","GO"))')
ws.cell(row=n+3,column=1,value='Gate 2 decision').font=BOLD; ws.cell(row=n+3,column=2,value='=IF(COUNTIF(F11:F16,"FAIL")>0,"STOP / REWORK",IF(COUNTIF(F11:F16,"PENDING")>0,"PENDING","SCALE"))')
ws.cell(row=n+4,column=1,value='Gate 3 decision').font=BOLD; ws.cell(row=n+4,column=2,value='=IF(COUNTIF(F17:F19,"FAIL")>0,"HOLD EXPANSION",IF(COUNTIF(F17:F19,"PENDING")>0,"PENDING","EXPAND CA / TEST UK"))')
ws.cell(row=n+6,column=1,value='Key model outputs').font=BOLD
for i,(lbl,f_,fmt) in enumerate([('P1 contribution before ads',"='12_Fee_Contribution'!B18",USD),('P1 break-even ACoS',"='12_Fee_Contribution'!B19",PCT),('Capital to test (base)',"='13_Launch_Budget'!C18",USD0),('Capital to test (low / high)',"=TEXT('13_Launch_Budget'!B18,\"$#,##0\")&\" / \"&TEXT('13_Launch_Budget'!D18,\"$#,##0\")",'General')]):
    ws.cell(row=n+7+i,column=1,value=lbl); c=ws.cell(row=n+7+i,column=2,value=f_); c.font=GREEN; c.number_format=fmt
widths(ws,[30,44,8,12,16,12])

# ---------------- Charts sheet ----------------
ws=sheet('Charts','Visual evidence','PNG charts generated by charts.py from the same data')
row=4
for f in sorted(os.listdir(CH)):
    img=XLImage(CH+f); img.width,img.height=760,int(760*img.height/img.width); ws.add_image(img,f'A{row}'); row+=int(img.height/20)+3

# README sheet first
ws=wb.create_sheet('README',0); ws['A1']='Bedding Market Entry Feasibility — research workbook'; ws['A1'].font=T
info=['Prepared 26 Oct 2026 · primary window YTD 2026 (1 Jan – 26 Oct) · marketplaces US, CA, UK','Colour code: blue = input/assumption · black = formula · green = link to another sheet · yellow fill = key input to replace',
'Labels: ACTUAL (our accounts) · CALC (calculated) · EST (third-party estimate) · ASSUMPTION (awaiting factory/forwarder quotes)','Core model tabs: 11_Landed_Cost → 12_Fee_Contribution → 13_Launch_Budget / 14_Profit_Scenarios → 19_Final_Shortlist / 20_GoNoGo_Dashboard',
'Data tabs come from Sellerboard, Ecotero Command Center, Data Dive, Keepa and Brand Analytics extracts pulled on 26 Oct 2026','Competitor sales and market sizes are estimates; do not treat them as observed facts','The executive report and decision memo are in the Claude doc "Bedding Market Entry Feasibility — Amazon US, Canada & UK" and in report.md in this folder']
for i,t in enumerate(info): ws.cell(row=3+i,column=1,value=t).font=N
ws.column_dimensions['A'].width=140
for w_ in wb.worksheets:
    for row_ in w_.iter_rows():
        for c in row_:
            if c.font is None or c.font.name!=F: c.font=Font(name=F,size=c.font.size if c.font else 10,bold=c.font.bold if c.font else False,italic=c.font.italic if c.font else False,color=c.font.color if c.font else None)
out=R+'Bedding_Market_Entry_Feasibility_Workbook.xlsx'; wb.save(out); print(out, len(wb.sheetnames), wb.sheetnames)
