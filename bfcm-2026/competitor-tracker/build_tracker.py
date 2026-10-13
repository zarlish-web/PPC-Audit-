"""BFCM 2026 Competitor Tracker builder (v1 for Thu Oct 15).

Reads the scratchpad data pulled from the Command Center (primary) and Data Dive:
  cc/{pid}_{US|CA}.json   rosters + our product profile/economics
  ours/{pid}_{US|CA}.json our BFCM 2025 / pre / Dec windows + daily + events
  dd_out/asins.json       Data Dive listing details per ASIN
  pt_summary.json         Command Center daily placement scores, summarised per window
Writes the xlsx plus keepa_request.csv.
"""
import json, glob, os, sys, csv, statistics, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.comments import Comment

_here = os.path.dirname(os.path.abspath(__file__))
SP = os.path.join(_here, "data") if os.path.isdir(os.path.join(_here, "data", "cc")) else _here
OUT = sys.argv[1]
OUT_CSV = sys.argv[2]

def jl(p, default=None):
    try:
        return json.load(open(p))
    except Exception:
        return default

# ---------------------------------------------------------------- products
PRODUCTS = [
    # pid, family, short, hero/role (from the BFCM 2026 plan, product family roles)
    (59, "Bamboo", "DC B4", "Hero: rank and profit"),
    (30, "Bamboo", "DC B6", "Hero: rank and profit"),
    (32, "Bamboo", "SL B6", "Core"),
    (36, "Bamboo", "SL B4", "Core"),
    (73, "Bamboo", "DC Bamboo Fitted", "Excluded (ranking variations OOS to Dec 3)"),
    (40, "Bamboo", "SL Bamboo Pillowcases", "Other"),
    (37, "Satin", "DC S4", "Clearance with margin"),
    (31, "Satin", "DC S6", "Clearance with margin"),
    (39, "Satin", "SSS4", "Clearance with margin"),
    (74, "Satin", "SL Satin 6pc", "Clearance with margin"),
    (38, "Satin", "DC Satin Fitted", "Opportunistic"),
    (55, "Silk", "DC Silk Pillowcase", "Gifting"),
    (56, "Silk", "SL Silk Mask Contoured", "Gifting"),
    (57, "Silk", "SL Silk Mask Velcro", "Gifting"),
    (48, "Cooling", "DC Cooling 4pc", "Off-season (aged variations only)"),
    (45, "Cooling", "SL Cooling 4pc", "Off-season (aged variations only)"),
    (76, "Cooling", "SL Cooling 6pc", "Off-season (aged variations only)"),
    (69, "Cooling", "CooLinens Cooling Sheets", "Launch"),
    (70, "Cooling", "CooLinens Cooling Pillowcases", "Launch"),
    (43, "Cooling", "SL Cooling Pillowcases", "Off-season"),
    (42, "Cooling", "SL Cooling Comforter", "Off-season (lost $2.4K at BFCM 2025)"),
    (44, "Cooling", "SL Cooling Fitted", "Off-season"),
    (46, "Other", "SL Quilt Set", "Liquidation"),
    (53, "Other", "DC Linen Curtains", "Liquidation"),
    (49, "Other", "DC Beach Towel", "Other"),
    (58, "Other", "DC Welcome Sign", "Other"),
    (33, "Other", "DC Chair Cover", "Other"),
]
PINFO = {p[0]: p for p in PRODUCTS}

def econ(e):
    e = e or {}
    b = e.get("blended") if isinstance(e.get("blended"), dict) else e
    g = lambda *ks: next((b.get(k) for k in ks if b.get(k) is not None), None)
    units = g("units_28d", "units28d", "units")
    if units is None and isinstance(e.get("by_size"), list):
        us = [s.get("units_28d") for s in e["by_size"] if s.get("units_28d") is not None]
        units = sum(us) if us else None
    return {
        "price": g("blended_price", "price"),
        "contribution": g("contribution_per_unit", "contribution"),
        "be_acos": g("break_even_acos", "be_acos"),
        "units_28d": units,
    }

cc = {}
for f in glob.glob(os.path.join(SP, "cc", "*.json")):
    d = jl(f)
    if d:
        cc[(int(d["product_id"]), d["marketplace"])] = d
ours = {}
for f in glob.glob(os.path.join(SP, "ours", "*.json")):
    d = jl(f)
    if d:
        ours[(int(d["product_id"]), d["marketplace"])] = d
dd = (jl(os.path.join(SP, "dd_out", "asins.json"), {}) or {}).get("asins", {})
ddb = (jl(os.path.join(SP, "dd_out", "asins.json"), {}) or {}).get("benchmarks", {})
pt = jl(os.path.join(SP, "pt_summary.json"), {}) or {}

# ---------------------------------------------------------------- styles
FONT = "Arial"
def F(**k):
    k.setdefault("size", 9)
    return Font(name=FONT, **k)
HDR = PatternFill("solid", fgColor="1F3864")
FILL = {k: PatternFill("solid", fgColor=v) for k, v in {
    "id": "D9E1F2", "cc": "E2EFDA", "dd": "DDEBF7", "vis": "EDE7F6", "kp": "FCE4D6", "str": "FFF2CC",
    "calc": "EDEDED", "flag": "F8CBAD", "ours": "FFF8E1"}.items()}
INPUT = PatternFill("solid", fgColor="FFFF99")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BLUE = "0000FF"
GREEN = "008000"

def header(ws, row, cols, height=60):
    for i, c in enumerate(cols, 1):
        h = ws.cell(row=row, column=i, value=c[1])
        h.font = F(bold=True, color="FFFFFF")
        h.fill = HDR
        h.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        h.border = BOX
        ws.column_dimensions[get_column_letter(i)].width = c[2]
        if len(c) > 4 and c[4]:
            h.comment = Comment(c[4], "Tracker")
    ws.row_dimensions[row].height = height

def group_row(ws, row, cols, names):
    spans, prev, start = [], None, 1
    for i, c in enumerate(cols, 1):
        g = c[3]
        if g != prev:
            if prev is not None:
                spans.append((prev, start, i - 1))
            prev, start = g, i
    spans.append((prev, start, len(cols)))
    for g, a, b in spans:
        if b > a:
            ws.merge_cells(start_row=row, start_column=a, end_row=row, end_column=b)
        cell = ws.cell(row=row, column=a, value=names.get(g, ""))
        cell.font = F(bold=True, color="1F3864")
        cell.fill = FILL.get(g, FILL["calc"])
        cell.alignment = Alignment(horizontal="center", vertical="center")

def pct_or_none(a, b):
    return None if (a is None or b in (None, 0)) else a / b - 1

wb = Workbook()

# ================================================================= README
rd = wb.active
rd.title = "README"
rd.sheet_view.showGridLines = False
today = "2026-10-13"
R = [
    ("BFCM 2026 Competitor Tracker (v1)", dict(bold=True, size=15)),
    (f"Owner: Shayan Rana  ·  Backup: Zarlish  ·  Due Thu Oct 15, 2026 (ClickUp BFCM-2026 › Competitor BFCM tracker)  ·  Built {today} from the Command Center (primary) and Data Dive", dict(italic=True)),
    ("", None),
    ("What it is for", dict(bold=True, size=11)),
    ("1. Lift forecast (Erik Oct 8 / Oct 13): tune our BFCM lift per product on what competitors did, not only our own history - deal type, depth and BSR change last BFCM, and how hard they push at events.", {}),
    ("2. Pricing and STR: where competitors set List Price today, how far over their price, and who already sits below our planned BFCM deal price (plan risk P2 'Competitors undercut on heroes' - report only, no regular-price cut without Joseph).", {}),
    ("3. Event watch: one log for Nov 20 - Dec 1 competitor checks.", {}),
    ("", None),
    ("Tabs", dict(bold=True, size=11)),
    ("Dashboard - one row per our product: our profile (SKUs, sizes, colours, price, margin), the competitor set, price position, List Price/STR read, our BFCM 2025 lift, Prime Day 2026 visibility lift (ours vs theirs) and Keepa completion.", {}),
    ("Competitors - every mapped competitor for every product (Command Center map), with current price / List Price / BSR / weekly units, listing details (variations, age, ads) from Data Dive, ASINsight traffic, Prime Day 2026 visibility, and the yellow BFCM 2025 + STR columns to fill from Keepa.", {}),
    ("Our BFCM 2025 - our own result per product: BFCM week vs the 30 days before and vs December, from the Command Center PPC/TACoS data, plus the deals and events the Command Center logged.", {}),
    ("Our Daily 2025 - day-by-day Oct 15 - Dec 15, 2025 per product (ad spend, ad sales, orders) to see the shape of the lift.", {}),
    ("Visibility History - Command Center daily placement scores summarised per window, per ASIN (organic and advertising). Competitor history starts Jan-Apr 2026, ours goes back to Oct 2025.", {}),
    ("Lift Inputs - per product, every input the lift forecast needs in one place, with the competitor benchmarks rolling up live as Keepa columns are filled.", {}),
    ("Keepa Request - the exact ASIN list and file spec to pull from Keepa (also saved as keepa_request.csv for the retrieval agent).", {}),
    ("Data Gaps - what is missing, why, and exactly what to request.", {}),
    ("Event Watch - log for the run-up and Nov 20 - Dec 1.", {}),
    ("", None),
    ("Colour legend", dict(bold=True, size=11)),
    ("Green header group = Command Center data (filled).  Blue = Data Dive (filled).  Purple = Command Center visibility history (filled).  Grey = formula.  YELLOW CELLS WITH BLUE TEXT = to fill (Keepa / Amazon page).", {}),
    ("", None),
    ("Definitions", dict(bold=True, size=11)),
    ("BFCM 2025 window = Nov 20 - Dec 1, 2025 (Black Friday Nov 28, Cyber Monday Dec 1). Pre-event baseline = Oct 21 - Nov 19, 2025 (after PBDD Oct 7-8). Prime Day 2026 = Jun 23-26, 2026 vs Jun 1-20 baseline.", {}),
    ("List Price % over = List Price / price - 1. Joseph's STR study: Amazon shows the List Price reliably when it is 12% or less over the price AND the item actually sold at it recently; 30-50% over is usually ignored.", {}),
    ("Discount % = 1 - BFCM low price / pre-event typical price.  BSR lift (x) = pre-event avg BSR / BFCM-week avg BSR (2.0x = BSR halved ~ roughly double the sales rate).", {}),
    ("Visibility score = Command Center ASINsight placement score (organic or advertising). It is a score, not traffic or sales; compare an ASIN with itself across windows.", {}),
    ("Est weekly units = Command Center snapshot estimate (Oct 12, 2026). Units 30d (DD) = Data Dive estimate. Both are estimates; use them for ranking competitors and market share, not as exact sales.", {}),
    ("Our BFCM 2025 total sales are derived from the Command Center: ad spend / TACoS. Marked 'derived' wherever used.", {}),
]
for i, (t, f) in enumerate(R, 1):
    c = rd.cell(row=i, column=1, value=t)
    if f is not None:
        c.font = F(**f) if f else F(size=10)
    c.alignment = Alignment(wrap_text=True, vertical="top")
rd.column_dimensions["A"].width = 160

# ================================================================= Competitors
cs = wb.create_sheet("Competitors")
C = [
    # key, header, width, group, comment
    ("family", "Family", 8, "id"),
    ("product", "Our product", 20, "id"),
    ("pid", "Prod ID", 5, "id"),
    ("role", "BFCM role (plan)", 14, "id"),
    ("cid", "Comp ID", 6, "id"),
    ("brand", "Brand", 16, "id"),
    ("parent", "Parent ASIN", 12, "id"),
    ("child", "Tracked child ASIN", 12, "id", "The child ASIN the Command Center snapshots for this competitor (its best seller / mapped child)."),
    ("title", "Title", 34, "id"),
    ("tier", "Mapping tier", 10, "id", "Command Center mapping tier for this rival against this product: Aspirational / Beatable / Poor."),
    ("price", "Price now ($)", 8, "cc"),
    ("lp", "List Price now ($)", 8, "cc", "List Price the Command Center snapshot holds (seen_on date). A List Price above the price is what Amazon can show as the strikethrough."),
    ("lp_pct", "List Price % over price", 8, "calc"),
    ("lp_band", "List Price band", 11, "calc", "<=12%: inside the band where STR shows reliably (if sold at it). 12-25%: risky. >25%: usually ignored by Amazon."),
    ("size", "Size (from title)", 8, "calc", "Size read from the Data Dive / Command Center title. Blank = not stated; the gap then uses our Queen price (or our blended price) and says so."),
    ("our_same", "Our price, same size ($)", 8, "calc"),
    ("gap", "Their price vs ours (same size)", 8, "calc"),
    ("basis", "Comparison basis", 10, "calc"),
    ("bsr", "BSR now", 8, "cc"),
    ("wk", "Est weekly units", 8, "cc"),
    ("rev", "Est monthly revenue ($)", 10, "calc"),
    ("rating", "Rating", 6, "cc"),
    ("reviews", "Reviews", 8, "cc"),
    ("seen", "Snapshot date", 9, "cc"),
    ("var", "Variations (DD)", 7, "dd", "Number of variations in the competitor's family (Data Dive)."),
    ("listed", "Listed (DD)", 9, "dd"),
    ("age", "Listing age (yrs)", 6, "calc"),
    ("dd30", "Units 30d (DD est.)", 8, "dd"),
    ("ful", "Fulfilment", 6, "dd"),
    ("sellers", "Active sellers", 6, "dd"),
    ("country", "Seller country", 6, "dd"),
    ("adkw", "Advertised KW %", 7, "dd", "Share of the niche's keywords this ASIN advertises on (Data Dive)."),
    ("tos", "TOS ad SV %", 7, "dd", "Share of niche search volume where this ASIN holds a top-of-search ad (Data Dive)."),
    ("p1", "Page-1 KWs", 6, "dd"),
    ("asi_tr", "ASINsight 7d traffic", 10, "cc", "Impression-like 7-day figure from the competitor's ASINsight export; not clicks or sales. Blank = unmeasured."),
    ("asi_kw", "ASINsight keywords", 8, "cc"),
    ("vis_from", "Visibility history from", 9, "vis"),
    ("pd_pre_org", "Organic score Jun 1-20", 7, "vis"),
    ("pd_org", "Organic score PD Jun 23-26", 7, "vis"),
    ("pd_org_lift", "PD organic lift", 7, "calc"),
    ("pd_pre_ad", "Ad score Jun 1-20", 7, "vis"),
    ("pd_ad", "Ad score PD", 7, "vis"),
    ("pd_ad_lift", "PD ad lift", 7, "calc"),
    ("q2_org", "Organic score Q2 (Apr 1-Jun 20)", 7, "vis"),
    ("l30_org", "Organic score Aug 18-Sep 16", 7, "vis"),
    ("trend", "Organic trend (last30 vs Q2)", 7, "calc"),
    ("str_show", "STR showing now? (Y/N)", 7, "str"),
    ("str_label", "STR label", 9, "str"),
    ("lp_sold", "Sold at List Price in last 90d? (Y/N)", 8, "str"),
    ("deal_now", "Deal / coupon live now", 10, "str"),
    ("bf_type", "BFCM 2025 deal type", 10, "kp"),
    ("bf_start", "Deal start", 8, "kp"),
    ("bf_end", "Deal end", 8, "kp"),
    ("bf_pre", "Pre-event typical price ($)", 8, "kp"),
    ("bf_low", "BFCM 2025 low price ($)", 8, "kp"),
    ("bf_disc", "BFCM 2025 discount %", 7, "calc"),
    ("bf_coupon", "Coupon on top (%)", 7, "kp"),
    ("bsr_pre", "BSR pre-event avg", 8, "kp"),
    ("bsr_bf", "BSR BFCM week avg", 8, "kp"),
    ("bsr_lift", "BSR lift (x)", 6, "calc"),
    ("bsr_dec", "BSR Dec 2-31 avg", 8, "kp"),
    ("bf_oos", "OOS / lost Buy Box in BFCM? (Y/N)", 7, "kp"),
    ("pd25", "PBDD 2025 deal price ($)", 8, "kp"),
    ("p_low90", "Lowest price last 90d ($)", 8, "kp"),
    ("below_ours", "Price below our price?", 7, "calc"),
    ("undercut", "Below our 2026 deal price?", 7, "calc"),
    ("notes", "Notes", 24, "kp"),
]
K = {c[0]: i + 1 for i, c in enumerate(C)}
L = {k: get_column_letter(v) for k, v in K.items()}
group_row(cs, 1, C, {"id": "COMPETITOR (Command Center map)", "cc": "CURRENT MARKET - Command Center snapshot", "calc": "",
                     "dd": "LISTING DETAIL - Data Dive", "vis": "VISIBILITY HISTORY - Command Center (2026)",
                     "str": "STR NOW - fill (Amazon page / Keepa)", "kp": "BFCM 2025 + PRICE HISTORY - fill from Keepa"})
header(cs, 2, C)
cs.freeze_panes = "G3"

def our_best_child(pid, mkt="US"):
    d = cc.get((pid, mkt))
    return (d or {}).get("our", {}).get("best_child_asin")

import re
our_size_price, our_blend = {}, {}
for (pid_, mkt_), d_ in cc.items():
    if mkt_ != "US":
        continue
    e_ = d_.get("our", {}).get("economics") or {}
    bs_ = e_.get("by_size") or (e_.get("blended") or {}).get("by_size") or []
    m_ = {}
    for x_ in bs_:
        if x_.get("price") is not None:
            nm = {"Cal King": "California King"}.get(x_.get("size"), x_.get("size"))
            m_.setdefault(nm, x_["price"])
    our_size_price[pid_] = m_
    our_blend[pid_] = econ(e_)["price"]
r = 3
FIRST = r
comp_rows = []
for pid, fam, short, role in PRODUCTS:
    d = cc.get((pid, "US"))
    if not d:
        continue
    o = d.get("our", {})
    pname = f"{o.get('brand') or ''} {o.get('name') or short}".strip()
    comps = d.get("competitors", [])
    tier_rank = {"Aspirational": 0, "Beatable": 1, "Poor": 2}
    comps = sorted(comps, key=lambda x: (tier_rank.get(x.get("tier"), 3), -(x.get("est_weekly_sales") or 0)))
    for cp in comps:
        ch = cp.get("child_asin")
        ddr = dd.get(ch) or dd.get(cp.get("parent_asin")) or {}
        v = pt.get(ch) or {}
        vals = {
            "family": fam, "product": pname, "pid": pid, "role": role, "cid": cp.get("id"),
            "brand": (cp.get("brand") or "").strip(), "parent": cp.get("parent_asin"), "child": ch,
            "title": (ddr.get("title") or cp.get("name") or "").strip()[:90], "tier": cp.get("tier"),
            "price": cp.get("price"), "lp": cp.get("list_price"), "bsr": cp.get("bsr"), "wk": cp.get("est_weekly_sales"),
            "rating": cp.get("rating"), "reviews": cp.get("reviews"), "seen": cp.get("seen_on"),
            "var": ddr.get("numberOfVariations"), "listed": ddr.get("listingCreationDate"), "dd30": ddr.get("sales"),
            "ful": ddr.get("fulfillment"), "sellers": ddr.get("numberOfActiveSellers"), "country": ddr.get("sellerCountry"),
            "adkw": ddr.get("advertisedKwsPercent"), "tos": ddr.get("tosSvAdsPercent"), "p1": ddr.get("kwRankedOnP1"),
            "asi_tr": cp.get("asinsight_traffic_7d"), "asi_kw": cp.get("asinsight_keywords") or None,
            "vis_from": v.get("first"), "pd_pre_org": v.get("pd26_pre_org"), "pd_org": v.get("pd26_org"),
            "pd_pre_ad": v.get("pd26_pre_ad"), "pd_ad": v.get("pd26_ad"), "q2_org": v.get("q2_org"), "l30_org": v.get("last30_org"),
        }
        ttl = f"{ddr.get('title') or ''} {cp.get('name') or ''}".lower()
        m = re.search(r"\b(california king|cal king|king|queen|full|twin xl|twin|standard)\b", ttl)
        sz = {"cal king": "California King", "twin xl": "Twin"}.get(m.group(1), m.group(1).title()) if m else None
        bys = our_size_price.get(pid, {})
        if sz and bys.get(sz) is not None:
            vals["our_same"], vals["basis"] = bys[sz], f"Same size ({sz})"
        elif bys.get("Queen") is not None:
            vals["our_same"], vals["basis"] = bys["Queen"], "Size unknown - vs our Queen"
        elif our_blend.get(pid) is not None:
            vals["our_same"], vals["basis"] = our_blend[pid], "vs our blended price"
        vals["size"] = sz
        comp_rows.append(vals)
        for key, hdr, w, grp, *rest in C:
            cell = cs.cell(row=r, column=K[key])
            cell.border = BOX
            cell.font = F()
            if grp in ("kp", "str"):
                cell.fill = INPUT
                cell.font = F(color=BLUE)
            elif grp == "calc" and key not in ("size", "our_same", "basis"):
                cell.fill = FILL["calc"]
            elif key in vals:
                val = vals[key]
                if key == "listed" and isinstance(val, str) and len(val) >= 10:
                    try:
                        val = dt.date.fromisoformat(val[:10])
                    except ValueError:
                        pass
                if key == "seen" and isinstance(val, str) and len(val) >= 10:
                    try:
                        val = dt.date.fromisoformat(val[:10])
                    except ValueError:
                        pass
                cell.value = val
        # formulas
        P, LPc = f"{L['price']}{r}", f"{L['lp']}{r}"
        cs[f"{L['lp_pct']}{r}"] = f'=IF(AND(ISNUMBER({LPc}),ISNUMBER({P}),{P}>0),{LPc}/{P}-1,"")'
        cs[f"{L['lp_band']}{r}"] = (f'=IF({L["lp_pct"]}{r}="","No List Price",IF({L["lp_pct"]}{r}<=0,"At/below price",'
                                    f'IF({L["lp_pct"]}{r}<=0.12,"<=12% (STR zone)",IF({L["lp_pct"]}{r}<=0.25,"12-25%",">25% (likely ignored)"))))')
        cs[f"{L['rev']}{r}"] = f'=IF(AND(ISNUMBER({P}),ISNUMBER({L["wk"]}{r})),{P}*{L["wk"]}{r}*52/12,"")'
        cs[f"{L['age']}{r}"] = f'=IF(ISNUMBER({L["listed"]}{r}),(DATE(2026,10,13)-{L["listed"]}{r})/365,"")'
        cs[f"{L['pd_org_lift']}{r}"] = f'=IF(AND(ISNUMBER({L["pd_org"]}{r}),ISNUMBER({L["pd_pre_org"]}{r}),{L["pd_pre_org"]}{r}>=5),{L["pd_org"]}{r}/{L["pd_pre_org"]}{r},"")'
        cs[f"{L['pd_ad_lift']}{r}"] = f'=IF(AND(ISNUMBER({L["pd_ad"]}{r}),ISNUMBER({L["pd_pre_ad"]}{r}),{L["pd_pre_ad"]}{r}>=5),{L["pd_ad"]}{r}/{L["pd_pre_ad"]}{r},"")'
        cs[f"{L['trend']}{r}"] = f'=IF(AND(ISNUMBER({L["l30_org"]}{r}),ISNUMBER({L["q2_org"]}{r}),{L["q2_org"]}{r}>0),{L["l30_org"]}{r}/{L["q2_org"]}{r}-1,"")'
        cs[f"{L['bf_disc']}{r}"] = f'=IF(AND(ISNUMBER({L["bf_pre"]}{r}),ISNUMBER({L["bf_low"]}{r}),{L["bf_pre"]}{r}>0),1-{L["bf_low"]}{r}/{L["bf_pre"]}{r},"")'
        cs[f"{L['bsr_lift']}{r}"] = f'=IF(AND(ISNUMBER({L["bsr_pre"]}{r}),ISNUMBER({L["bsr_bf"]}{r}),{L["bsr_bf"]}{r}>0),{L["bsr_pre"]}{r}/{L["bsr_bf"]}{r},"")'
        cs[f"{L['gap']}{r}"] = f'=IF(AND(ISNUMBER({P}),ISNUMBER({L["our_same"]}{r}),N({L["our_same"]}{r})>0),{P}/{L["our_same"]}{r}-1,"")'
        cs[f"{L['below_ours']}{r}"] = f'=IF(ISNUMBER({L["gap"]}{r}),IF({L["gap"]}{r}<0,"YES","no"),"")'
        cs[f"{L['undercut']}{r}"] = f'=IFERROR(IF(AND(ISNUMBER({P}),ISNUMBER(INDEX(\'Lift Inputs\'!$H:$H,MATCH({L["pid"]}{r},\'Lift Inputs\'!$C:$C,0)))),IF({P}<INDEX(\'Lift Inputs\'!$H:$H,MATCH({L["pid"]}{r},\'Lift Inputs\'!$C:$C,0)),"YES","no"),""),"")'
        # formats
        fm = {"price": "$#,##0.00", "lp": "$#,##0.00", "lp_pct": "0.0%", "bsr": "#,##0", "wk": "#,##0", "rev": "$#,##0",
              "reviews": "#,##0", "seen": "mmm d", "listed": "yyyy-mm-dd", "age": "0.0", "dd30": "#,##0", "adkw": "0%", "tos": "0%",
              "asi_tr": "#,##0", "asi_kw": "#,##0", "pd_org_lift": '0.00"x"', "pd_ad_lift": '0.00"x"', "trend": "+0%;-0%;0%",
              "bf_start": "mmm d", "bf_end": "mmm d", "bf_pre": "$#,##0.00", "bf_low": "$#,##0.00", "bf_disc": "0.0%",
              "bf_coupon": "0%", "bsr_pre": "#,##0", "bsr_bf": "#,##0", "bsr_lift": '0.0"x"', "bsr_dec": "#,##0",
              "pd25": "$#,##0.00", "p_low90": "$#,##0.00", "our_same": "$#,##0.00", "gap": "+0%;-0%;0%",
              "pd_pre_org": "0.0", "pd_org": "0.0", "pd_pre_ad": "0.0", "pd_ad": "0.0", "q2_org": "0.0", "l30_org": "0.0"}
        for k2, nf in fm.items():
            cs.cell(row=r, column=K[k2]).number_format = nf
        r += 1
LAST = r - 1
rng = lambda k: f"Competitors!${L[k]}${FIRST}:${L[k]}${LAST}"

for dv_list, keys in (('"Y,N"', ("str_show", "lp_sold", "bf_oos")),
                      ('"List Price,Typical price,Was price,None"', ("str_label",)),
                      ('"Lightning Deal,Best Deal,Prime Exclusive Discount,Coupon only,Price cut only,None,Not live in 2025,Unknown"', ("bf_type",)),
                      ('"None,Coupon,Lightning Deal,Best Deal,Prime Exclusive Discount,Price cut"', ("deal_now",))):
    dv = DataValidation(type="list", formula1=dv_list, allow_blank=True)
    cs.add_data_validation(dv)
    for k in keys:
        dv.add(f"{L[k]}{FIRST}:{L[k]}{LAST}")
for k in ("below_ours", "undercut"):
    cs.conditional_formatting.add(f"{L[k]}{FIRST}:{L[k]}{LAST}",
                                  FormulaRule(formula=[f'{L[k]}{FIRST}="YES"'], fill=FILL["flag"], font=Font(name=FONT, bold=True, color="C00000")))
cs.conditional_formatting.add(f"{L['lp_band']}{FIRST}:{L['lp_band']}{LAST}",
                              FormulaRule(formula=[f'{L["lp_band"]}{FIRST}="<=12% (STR zone)"'], fill=PatternFill("solid", fgColor="C6EFCE")))
cs.auto_filter.ref = f"A2:{get_column_letter(len(C))}{LAST}"

# ================================================================= Dashboard
db = wb.create_sheet("Dashboard", 1)
D = [
    ("family", "Family", 8, "id"), ("product", "Our product", 22, "id"), ("pid", "Prod ID", 5, "id"), ("role", "BFCM role (plan)", 16, "id"),
    ("stage", "Stage", 10, "ours"), ("parent", "Our parent ASIN", 11, "ours"), ("skus", "Our SKUs (variations)", 7, "ours"),
    ("sizes", "Sizes", 5, "ours"), ("colours", "Colours", 6, "ours"), ("price", "Our blended price ($)", 8, "ours"),
    ("contrib", "Contribution / unit ($)", 8, "ours"), ("beacos", "Break-even ACoS", 7, "ours"), ("u28", "Our units 28d", 7, "ours"),
    ("ncomp", "Competitors mapped", 7, "calc"), ("nasp", "Aspirational", 6, "calc"), ("nbeat", "Beatable", 6, "calc"), ("npoor", "Poor", 5, "calc"),
    ("medp", "Comp median price ($)", 8, "calc"), ("vsmed", "Our price vs comp median", 7, "calc"),
    ("nbelow", "Comps priced below us (same size)", 6, "calc"), ("nlp", "Comps with List Price set", 7, "calc"),
    ("nlp12", "of which <=12% over (STR zone)", 7, "calc"), ("avglp", "Comp avg List Price % over", 7, "calc"),
    ("compwk", "Comp est weekly units (sum)", 9, "calc"), ("ourwk", "Our weekly units (28d/4)", 8, "calc"), ("share", "Our share of tracked units", 7, "calc"),
    ("medvar", "Comp median variations (DD)", 7, "calc"),
    ("bf_lift", "Our BFCM 2025 lift (daily sales vs 30d pre)", 8, "calc"), ("bf_tacos", "Our BFCM 2025 TACoS", 7, "calc"),
    ("our_pd", "Our PD 2026 organic vis lift", 7, "vis"), ("comp_pd", "Comp median PD 2026 organic vis lift", 7, "calc"),
    ("kp_done", "Keepa rows filled", 6, "calc"), ("comp_disc", "Comp avg BFCM 2025 discount", 7, "calc"), ("comp_bsr", "Comp avg BFCM 2025 BSR lift", 7, "calc"),
]
DK = {c[0]: i + 1 for i, c in enumerate(D)}
DL = {k: get_column_letter(v) for k, v in DK.items()}
group_row(db, 1, D, {"id": "PRODUCT", "ours": "OUR PRODUCT - Command Center", "calc": "COMPETITOR SET + OUR HISTORY (formulas)", "vis": ""})
header(db, 2, D)
db.freeze_panes = "C3"
dr = 3
D_FIRST = dr
for pid, fam, short, role in PRODUCTS:
    d = cc.get((pid, "US"))
    if not d:
        continue
    o = d.get("our", {})
    e = econ(o.get("economics"))
    pname = f"{o.get('brand') or ''} {o.get('name') or short}".strip()
    bc = o.get("best_child_asin")
    v = pt.get(bc) or {}
    our_pd = (v["pd26_org"] / v["pd26_pre_org"]) if v.get("pd26_org") is not None and (v.get("pd26_pre_org") or 0) >= 5 else None
    vals = {"family": fam, "product": pname, "pid": pid, "role": role, "stage": o.get("stage"), "parent": o.get("parent_asin"),
            "skus": o.get("sku_count"), "sizes": len(o.get("sizes") or []) or None, "colours": o.get("colour_count"),
            "price": e["price"], "contrib": e["contribution"], "beacos": (e["be_acos"] / 100 if e["be_acos"] is not None and abs(e["be_acos"]) > 1.5 else e["be_acos"]),
            "u28": e["units_28d"], "our_pd": our_pd}
    for key, hdr, w, grp, *rest in D:
        c = db.cell(row=dr, column=DK[key]); c.border = BOX; c.font = F()
        if key in vals:
            c.value = vals[key]
            if grp in ("ours", "vis"):
                c.fill = FILL["ours"]
        else:
            c.fill = FILL["calc"]
    P = f"$C{dr}"
    f = {
        "ncomp": f"=COUNTIF({rng('pid')},{P})",
        "nasp": f'=COUNTIFS({rng("pid")},{P},{rng("tier")},"Aspirational")',
        "nbeat": f'=COUNTIFS({rng("pid")},{P},{rng("tier")},"Beatable")',
        "npoor": f'=COUNTIFS({rng("pid")},{P},{rng("tier")},"Poor")',
        "vsmed": f'=IF(AND(ISNUMBER({DL["price"]}{dr}),ISNUMBER({DL["medp"]}{dr}),{DL["medp"]}{dr}>0),{DL["price"]}{dr}/{DL["medp"]}{dr}-1,"")',
        "nbelow": f'=COUNTIFS({rng("pid")},{P},{rng("gap")},"<0")',
        "nlp": f'=COUNTIFS({rng("pid")},{P},{rng("lp_pct")},">0")',
        "nlp12": f'=COUNTIFS({rng("pid")},{P},{rng("lp_pct")},">0",{rng("lp_pct")},"<=0.12")',
        "avglp": f'=IFERROR(AVERAGEIFS({rng("lp_pct")},{rng("pid")},{P},{rng("lp_pct")},">0"),"")',
        "compwk": f"=SUMIF({rng('pid')},{P},{rng('wk')})",
        "ourwk": f'=IF(ISNUMBER({DL["u28"]}{dr}),{DL["u28"]}{dr}/4,"")',
        "share": f'=IF(AND(ISNUMBER({DL["ourwk"]}{dr}),{DL["compwk"]}{dr}>0),{DL["ourwk"]}{dr}/({DL["ourwk"]}{dr}+{DL["compwk"]}{dr}),"")',
        "bf_lift": f"=IFERROR(IF(N(INDEX('Our BFCM 2025'!$S:$S,MATCH({P},'Our BFCM 2025'!$C:$C,0)))=0,\"\",INDEX('Our BFCM 2025'!$S:$S,MATCH({P},'Our BFCM 2025'!$C:$C,0))),\"\")",
        "bf_tacos": f"=IFERROR(IF(N(INDEX('Our BFCM 2025'!$H:$H,MATCH({P},'Our BFCM 2025'!$C:$C,0)))=0,\"\",INDEX('Our BFCM 2025'!$H:$H,MATCH({P},'Our BFCM 2025'!$C:$C,0))),\"\")",
        "kp_done": f'=COUNTIFS({rng("pid")},{P},{rng("bf_type")},"<>")',
        "comp_disc": f'=IFERROR(AVERAGEIFS({rng("bf_disc")},{rng("pid")},{P}),"")',
        "comp_bsr": f'=IFERROR(AVERAGEIFS({rng("bsr_lift")},{rng("pid")},{P}),"")',
    }
    for k2, fx in f.items():
        db[f"{DL[k2]}{dr}"] = fx
    # median price / median variations are written as values (MEDIAN(IF()) needs array entry); documented in header comment
    pr = [x["price"] for x in comp_rows if x["pid"] == pid and isinstance(x.get("price"), (int, float))]
    db[f"{DL['medp']}{dr}"] = statistics.median(pr) if pr else None
    cpd_ = [x["pd_org"] / x["pd_pre_org"] for x in comp_rows if x["pid"] == pid and x.get("pd_org") is not None and (x.get("pd_pre_org") or 0) >= 5]
    db[f"{DL['comp_pd']}{dr}"] = statistics.median(cpd_) if cpd_ else None
    vr = [x["var"] for x in comp_rows if x["pid"] == pid and isinstance(x.get("var"), (int, float))]
    db[f"{DL['medvar']}{dr}"] = statistics.median(vr) if vr else None
    fm = {"price": "$#,##0.00", "contrib": "$#,##0.00", "beacos": "0.0%", "u28": "#,##0", "medp": "$#,##0.00", "vsmed": "+0%;-0%;0%",
          "avglp": "0.0%", "compwk": "#,##0", "ourwk": "#,##0", "share": "0.0%", "bf_lift": '0.00"x"', "bf_tacos": "0.0%",
          "our_pd": '0.00"x"', "comp_pd": '0.00"x"', "comp_disc": "0.0%", "comp_bsr": '0.0"x"'}
    for k2, nf in fm.items():
        db.cell(row=dr, column=DK[k2]).number_format = nf
    dr += 1
D_LAST = dr - 1
db.cell(row=2, column=DK["medp"]).comment = Comment("Median of the Competitors tab 'Price now' for this product, computed at build time (Oct 13). Re-run the build to refresh.", "Tracker")
db.cell(row=2, column=DK["medvar"]).comment = Comment("Median Data Dive variation count of mapped competitors, computed at build time.", "Tracker")
db.cell(row=2, column=DK["bf_lift"]).comment = Comment("From 'Our BFCM 2025': derived total sales per day in BFCM week (Nov 20-Dec 1, 2025) / per day in the 30 days before (Oct 21-Nov 19). Total sales = ad spend / TACoS.", "Tracker")
db.cell(row=D_LAST + 2, column=1, value="Medians are build-time values; every other grey cell is a live formula over the Competitors and Our BFCM 2025 tabs. Our PD 2026 organic visibility lift is for our best-selling child ASIN.").font = F(italic=True)

# ================================================================= Our BFCM 2025
ob = wb.create_sheet("Our BFCM 2025")
O = [("family", "Family", 8), ("product", "Our product", 22), ("pid", "Prod ID", 5), ("mkt", "Mkt", 4),
     ("bf_spend", "BFCM wk ad spend ($)", 9), ("bf_adsales", "BFCM wk ad sales ($)", 9), ("bf_acos", "BFCM wk ACoS", 7), ("bf_tacos", "BFCM wk TACoS", 7),
     ("bf_total", "BFCM wk total sales ($, derived)", 10), ("bf_orders", "BFCM wk ad orders", 7), ("bf_cvr", "BFCM wk CVR", 6),
     ("pre_spend", "Pre-30d ad spend ($)", 9), ("pre_tacos", "Pre-30d TACoS", 7), ("pre_total", "Pre-30d total sales ($, derived)", 10),
     ("dec_tacos", "Dec TACoS", 7), ("dec_total", "Dec 2-31 total sales ($, derived)", 10),
     ("bf_day", "BFCM wk total sales / day", 9), ("pre_day", "Pre-30d total sales / day", 9), ("lift", "BFCM lift (x) vs pre-30d", 7),
     ("dec_day", "Dec total sales / day", 9), ("dec_lift", "Dec vs pre-30d (x)", 7), ("prior12", "BFCM wk ad sales vs prior 12d (x)", 7),
     ("events", "Deals / events logged by the Command Center (Oct 15 - Dec 15, 2025)", 60)]
OK = {c[0]: i + 1 for i, c in enumerate(O)}
OL = {k: get_column_letter(v) for k, v in OK.items()}
header(ob, 1, [(a, b, c, "x") for a, b, c in O])
ob.freeze_panes = "D2"

def W(d, *keys):
    for k in keys:
        if isinstance(d, dict) and d.get(k) is not None:
            return d[k]
    return None
def tile(w, name):
    if not isinstance(w, dict):
        return None
    v = W(w, name, name.replace("_", ""))
    if isinstance(v, dict):
        v = v.get("value")
    return v

orow = 2
for pid, fam, short, role in PRODUCTS:
    for mkt in ("US", "CA"):
        d = ours.get((pid, mkt))
        if not d:
            continue
        o = (cc.get((pid, mkt)) or cc.get((pid, "US")) or {}).get("our", {})
        pname = f"{o.get('brand') or ''} {o.get('name') or short}".strip()
        w1, w2, w3 = d.get("w1_bfcm") or {}, d.get("w2_pre30") or {}, d.get("w3_dec") or {}
        def tot(w):
            s, t = tile(w, "spend"), tile(w, "tacos")
            if W(w, "total_sales", "total_sales_derived") is not None:
                return W(w, "total_sales", "total_sales_derived")
            return s / (t / 100) if (s and t) else None
        ev = d.get("events") or []
        evs = "; ".join(
            " ".join(str(x) for x in [e.get("type") or e.get("kind"), e.get("start") or e.get("date") or e.get("from"), ("to " + str(e.get("end") or e.get("to"))) if (e.get("end") or e.get("to")) else "", e.get("description") or e.get("name") or e.get("label") or ""] if x)
            if isinstance(e, dict) else str(e) for e in ev)[:500]
        prior = w1.get("prior") or {}
        for w_ in (w1, w2, w3):
            if isinstance(w_, dict) and not tile(w_, "spend"):
                w_["tacos"] = None
        vals = {"family": fam, "product": pname, "pid": pid, "mkt": mkt,
                "bf_spend": tile(w1, "spend"), "bf_adsales": tile(w1, "sales") or tile(w1, "ad_sales"),
                "bf_acos": (tile(w1, "acos") or 0) / 100 if tile(w1, "acos") is not None else None,
                "bf_tacos": (tile(w1, "tacos") or 0) / 100 if tile(w1, "tacos") is not None else None,
                "bf_total": tot(w1), "bf_orders": tile(w1, "orders"),
                "bf_cvr": (tile(w1, "cvr") / 100) if tile(w1, "cvr") is not None else None,
                "pre_spend": tile(w2, "spend"), "pre_tacos": (tile(w2, "tacos") / 100) if tile(w2, "tacos") is not None else None,
                "pre_total": tot(w2), "dec_tacos": (tile(w3, "tacos") / 100) if tile(w3, "tacos") is not None else None,
                "dec_total": tot(w3), "events": evs or None}
        psales = tile(prior, "sales") or tile(prior, "ad_sales")
        for key, hdr, w in O:
            c = ob.cell(row=orow, column=OK[key]); c.border = BOX; c.font = F()
            if key in vals:
                c.value = vals[key]
        ob[f"{OL['bf_day']}{orow}"] = f'=IF(ISNUMBER({OL["bf_total"]}{orow}),{OL["bf_total"]}{orow}/12,"")'
        ob[f"{OL['pre_day']}{orow}"] = f'=IF(ISNUMBER({OL["pre_total"]}{orow}),{OL["pre_total"]}{orow}/30,"")'
        ob[f"{OL['lift']}{orow}"] = f'=IF(AND(ISNUMBER({OL["bf_day"]}{orow}),ISNUMBER({OL["pre_day"]}{orow}),N({OL["pre_day"]}{orow})>0),{OL["bf_day"]}{orow}/{OL["pre_day"]}{orow},"")'
        ob[f"{OL['dec_day']}{orow}"] = f'=IF(ISNUMBER({OL["dec_total"]}{orow}),{OL["dec_total"]}{orow}/30,"")'
        ob[f"{OL['dec_lift']}{orow}"] = f'=IF(AND(ISNUMBER({OL["dec_day"]}{orow}),ISNUMBER({OL["pre_day"]}{orow}),N({OL["pre_day"]}{orow})>0),{OL["dec_day"]}{orow}/{OL["pre_day"]}{orow},"")'
        if psales and vals["bf_adsales"]:
            ob[f"{OL['prior12']}{orow}"] = vals["bf_adsales"] / psales
        for k2 in ("bf_day", "pre_day", "lift", "dec_day", "dec_lift", "prior12"):
            ob[f"{OL[k2]}{orow}"].fill = FILL["calc"]
        fm = {"bf_spend": "$#,##0", "bf_adsales": "$#,##0", "bf_acos": "0.0%", "bf_tacos": "0.0%", "bf_total": "$#,##0", "bf_orders": "#,##0",
              "bf_cvr": "0.0%", "pre_spend": "$#,##0", "pre_tacos": "0.0%", "pre_total": "$#,##0", "dec_tacos": "0.0%", "dec_total": "$#,##0",
              "bf_day": "$#,##0", "pre_day": "$#,##0", "lift": '0.00"x"', "dec_day": "$#,##0", "dec_lift": '0.00"x"', "prior12": '0.00"x"'}
        for k2, nf in fm.items():
            ob[f"{OL[k2]}{orow}"].number_format = nf
        ob[f"{OL['events']}{orow}"].alignment = Alignment(wrap_text=True, vertical="top")
        orow += 1
ob.cell(row=orow + 1, column=1, value="Source: Command Center PPC summary (ppc-daily-product-day). Total sales are derived as ad spend / TACoS. Sales carry 7-day attribution. Windows: BFCM wk Nov 20-Dec 1, 2025; pre-30d Oct 21-Nov 19, 2025; Dec 2-31, 2025.").font = F(italic=True)
# Dashboard looks up column S (lift) and H (TACoS) by product id in column C - keep these positions
assert OL["lift"] == "S" and OL["bf_tacos"] == "H" and OL["pid"] == "C", (OL["lift"], OL["bf_tacos"])

# ================================================================= Our Daily 2025
od = wb.create_sheet("Our Daily 2025")
header(od, 1, [("a", "Product", 24, "x"), ("b", "Prod ID", 6, "x"), ("c", "Mkt", 5, "x"), ("d", "Date", 10, "x"),
               ("e", "Ad spend ($)", 9, "x"), ("f", "Ad sales ($)", 9, "x"), ("g", "Ad orders", 7, "x"), ("h", "Units", 7, "x"), ("i", "Window", 10, "x")], height=30)
rr = 2
for pid, fam, short, role in PRODUCTS:
    for mkt in ("US", "CA"):
        d = ours.get((pid, mkt))
        if not d:
            continue
        o = (cc.get((pid, "US")) or {}).get("our", {})
        pname = f"{o.get('brand') or ''} {o.get('name') or short}".strip()
        for x in d.get("daily") or []:
            date = x.get("date")
            win = ("BFCM wk" if "2025-11-20" <= date <= "2025-12-01" else "Pre-30d" if "2025-10-21" <= date <= "2025-11-19" else "")
            vals = [pname, pid, mkt, dt.date.fromisoformat(date[:10]) if date else None, x.get("spend"), x.get("sales"), x.get("orders"), x.get("units"), win]
            for i, val in enumerate(vals, 1):
                c = od.cell(row=rr, column=i, value=val); c.font = F()
            od.cell(row=rr, column=4).number_format = "yyyy-mm-dd"
            od.cell(row=rr, column=5).number_format = "$#,##0"
            od.cell(row=rr, column=6).number_format = "$#,##0"
            rr += 1
od.freeze_panes = "A2"
od.auto_filter.ref = f"A1:I{max(rr - 1, 1)}"

# ================================================================= Visibility History
vh = wb.create_sheet("Visibility History")
VH = [("asin", "ASIN", 12), ("side", "Ours / competitor", 9), ("brand", "Brand", 16), ("product", "Mapped to (our product)", 26),
      ("first", "History from", 10), ("last", "History to", 10),
      ("bfp_o", "Org Oct21-Nov19 2025", 8), ("bf_o", "Org BFCM wk 2025", 8), ("bf_ol", "BFCM 2025 org lift", 7),
      ("bfp_a", "Ad Oct21-Nov19 2025", 8), ("bf_a", "Ad BFCM wk 2025", 8), ("bf_al", "BFCM 2025 ad lift", 7),
      ("pdp_o", "Org Jun 1-20 2026", 8), ("pd_o", "Org PD Jun 23-26", 8), ("pd_ol", "PD org lift", 7),
      ("pdp_a", "Ad Jun 1-20 2026", 8), ("pd_a", "Ad PD", 8), ("pd_al", "PD ad lift", 7),
      ("q2_o", "Org Q2", 8), ("l30_o", "Org Aug18-Sep16", 8), ("tr", "Org trend", 7)]
header(vh, 1, [(a, b, c, "x") for a, b, c in VH])
vh.freeze_panes = "B2"
asin_meta = {}
for x in comp_rows:
    asin_meta.setdefault(x["child"], ("competitor", x["brand"], set()))[2].add(x["product"])
for (pid, mkt), d in cc.items():
    if mkt != "US":
        continue
    o = d.get("our", {})
    if o.get("best_child_asin"):
        asin_meta[o["best_child_asin"]] = ("OURS", o.get("brand"), {f"{o.get('brand')} {o.get('name')}"})
vr_ = 2
for a, v in sorted(pt.items(), key=lambda kv: (asin_meta.get(kv[0], ("z",))[0] != "OURS", kv[0])):
    side, brand, prods = asin_meta.get(a, ("competitor", None, set()))
    vals = [a, side, brand, "; ".join(sorted(prods))[:120], v.get("first"), v.get("last"),
            v.get("bf25_pre_org"), v.get("bf25_org"), None, v.get("bf25_pre_ad"), v.get("bf25_ad"), None,
            v.get("pd26_pre_org"), v.get("pd26_org"), None, v.get("pd26_pre_ad"), v.get("pd26_ad"), None,
            v.get("q2_org"), v.get("last30_org"), None]
    for i, val in enumerate(vals, 1):
        c = vh.cell(row=vr_, column=i, value=val); c.font = F(bold=(side == "OURS")); c.border = BOX
    for lift_col, a_col, b_col, nf in (("I", "G", "H", '0.00"x"'), ("L", "J", "K", '0.00"x"'), ("O", "M", "N", '0.00"x"'), ("R", "P", "Q", '0.00"x"')):
        vh[f"{lift_col}{vr_}"] = f'=IF(AND(ISNUMBER({b_col}{vr_}),ISNUMBER({a_col}{vr_}),N({a_col}{vr_})>=5),{b_col}{vr_}/{a_col}{vr_},"")'
        vh[f"{lift_col}{vr_}"].number_format = nf
        vh[f"{lift_col}{vr_}"].fill = FILL["calc"]
    vh[f"U{vr_}"] = f'=IF(AND(ISNUMBER(T{vr_}),ISNUMBER(S{vr_}),N(S{vr_})>0),T{vr_}/S{vr_}-1,"")'
    vh[f"U{vr_}"].number_format = "+0%;-0%;0%"
    vh[f"U{vr_}"].fill = FILL["calc"]
    if side == "OURS":
        for i in range(1, 7):
            vh.cell(row=vr_, column=i).fill = FILL["ours"]
    vr_ += 1
vh.auto_filter.ref = f"A1:U{max(vr_ - 1, 1)}"
vh.cell(row=vr_ + 1, column=1, value="Source: Command Center traffic_position_trends (ASINsight daily placement scores; feed retired Sep 16, 2026). Scores, not traffic. Competitor history starts Jan-Apr 2026, so BFCM 2025 columns are filled for our ASINs only.").font = F(italic=True)

# ================================================================= Lift Inputs
li = wb.create_sheet("Lift Inputs")
LI = [("family", "Family", 8), ("product", "Our product", 22), ("pid", "Prod ID", 5), ("role", "BFCM role (plan)", 16),
      ("deal26", "Planned 2026 deal type (fill)", 12), ("dealdates", "2026 deal dates (fill)", 12), ("deal_pct", "2026 deal discount % (fill)", 8),
      ("dealprice", "2026 hero deal price $ (fill)", 9), ("u28", "Our units 28d", 7),
      ("base_day", "Base units / day (28d)", 7), ("our25", "Our BFCM 2025 lift (x)", 7), ("our25_tacos", "Our BFCM 2025 TACoS", 7),
      ("our_pd", "Our PD26 organic vis lift", 7), ("comp_pd", "Comp median PD26 organic vis lift", 7),
      ("comp_bsr", "Comp avg BFCM25 BSR lift (Keepa)", 7), ("comp_disc", "Comp avg BFCM25 discount (Keepa)", 7),
      ("our_bsr", "Our BFCM25 BSR lift (Keepa, fill)", 7), ("ratio", "Comp BSR lift / ours", 7),
      ("oos25", "Were we OOS / restricted in BFCM 2025? (fill)", 10), ("notes", "Notes for Joseph's forecast", 40)]
LK = {c[0]: i + 1 for i, c in enumerate(LI)}
LLc = {k: get_column_letter(v) for k, v in LK.items()}
header(li, 1, [(a, b, c, "x") for a, b, c in LI])
assert LLc["dealprice"] == "H" and LLc["pid"] == "C"
li.freeze_panes = "D2"
lr = 2
for pid, fam, short, role in PRODUCTS:
    if not cc.get((pid, "US")):
        continue
    o = cc[(pid, "US")]["our"]
    pname = f"{o.get('brand') or ''} {o.get('name') or short}".strip()
    vals = {"family": fam, "product": pname, "pid": pid, "role": role}
    for key, hdr, w in LI:
        c = li.cell(row=lr, column=LK[key]); c.border = BOX; c.font = F()
        if key in vals:
            c.value = vals[key]
        elif key in ("deal26", "dealdates", "deal_pct", "dealprice", "our_bsr", "oos25", "notes"):
            c.fill = INPUT; c.font = F(color=BLUE)
        else:
            c.fill = FILL["calc"]
    m = f"MATCH($C{lr},Dashboard!$C:$C,0)"
    li[f"{LLc['u28']}{lr}"] = f"=IFERROR(INDEX(Dashboard!${DL['u28']}:${DL['u28']},{m}),\"\")"
    li[f"{LLc['base_day']}{lr}"] = f'=IF(ISNUMBER({LLc["u28"]}{lr}),{LLc["u28"]}{lr}/28,"")'
    li[f"{LLc['our25']}{lr}"] = f"=IFERROR(INDEX(Dashboard!${DL['bf_lift']}:${DL['bf_lift']},{m}),\"\")"
    li[f"{LLc['our25_tacos']}{lr}"] = f"=IFERROR(INDEX(Dashboard!${DL['bf_tacos']}:${DL['bf_tacos']},{m}),\"\")"
    li[f"{LLc['our_pd']}{lr}"] = f"=IFERROR(INDEX(Dashboard!${DL['our_pd']}:${DL['our_pd']},{m}),\"\")"
    li[f"{LLc['comp_pd']}{lr}"] = f"=IFERROR(INDEX(Dashboard!${DL['comp_pd']}:${DL['comp_pd']},{m}),\"\")"
    li[f"{LLc['comp_bsr']}{lr}"] = f"=IFERROR(INDEX(Dashboard!${DL['comp_bsr']}:${DL['comp_bsr']},{m}),\"\")"
    li[f"{LLc['comp_disc']}{lr}"] = f"=IFERROR(INDEX(Dashboard!${DL['comp_disc']}:${DL['comp_disc']},{m}),\"\")"
    li[f"{LLc['ratio']}{lr}"] = f'=IF(AND(ISNUMBER({LLc["comp_bsr"]}{lr}),ISNUMBER({LLc["our_bsr"]}{lr}),N({LLc["our_bsr"]}{lr})>0),{LLc["comp_bsr"]}{lr}/{LLc["our_bsr"]}{lr},"")'
    fm = {"deal_pct": "0%", "dealprice": "$#,##0.00", "u28": "#,##0", "base_day": "0.0", "our25": '0.00"x"', "our25_tacos": "0.0%",
          "our_pd": '0.00"x"', "comp_pd": '0.00"x"', "comp_bsr": '0.0"x"', "comp_disc": "0.0%", "our_bsr": '0.0"x"', "ratio": '0.00"x"'}
    for k2, nf in fm.items():
        li[f"{LLc[k2]}{lr}"].number_format = nf
    lr += 1
li.cell(row=lr + 1, column=1, value="How to read it: 'Comp BSR lift / ours' > 1 means competitors gained more rank in BFCM 2025 than we did - our own 2025 lift understates what the product can do (e.g. we were OOS or had no deal). Joseph applies it to the forecast; this tab only assembles the inputs.").font = F(italic=True)
li.cell(row=lr + 2, column=1, value="Yellow = fill. Planned 2026 deal type / dates / price come from Umair's deal coverage list and Joseph's D1 board. The hero deal price drives the 'Below our 2026 deal price?' flag on the Competitors tab.").font = F(italic=True)

# ================================================================= Keepa Request
kr = wb.create_sheet("Keepa Request")
kr.sheet_view.showGridLines = False
spec = [
    ("What to pull from Keepa (send this tab or keepa_request.csv to the retrieval agent)", dict(bold=True, size=12)),
    ("FILE 1 - Product Viewer export (one CSV, all ASINs below, domain Amazon.com): Keepa > Data > Product Viewer > paste the ASIN list > Export (CSV, all columns). Needed columns: ASIN, Title, Brand, Parent ASIN, Variation count / Variation ASINs, Buy Box: Current, Buy Box: 90 days avg, Buy Box: Lowest (90 days), List Price: Current, Sales Rank: Current, Sales Rank: 30/90 days avg, Bought in past month, Lightning Deals: Current, Coupon (one-time / S&S), Reviews: Rating / Review Count, Buy Box Seller, Count of offers.", {}),
    ("FILE 2 - Price + rank history, one CSV per ASIN (or one long CSV with an ASIN column), daily rows Sep 1, 2025 - Jan 5, 2026 AND Jun 1 - Jul 5, 2026 AND Sep 1 - Oct 13, 2026. Columns: date, buy_box_price, new_price, amazon_price, list_price, lightning_deal_price, coupon (amount or %), sales_rank (main category), sales_rank_subcategory, out_of_stock (Y/N). From the Keepa graph: 'Export' / Download data, or Keepa API product?history=1&days=420&offers=20&rating=1&buybox=1&stats=420.", {}),
    ("FILE 3 (optional, only if quick) - Keepa Deals history for Nov 20 - Dec 1, 2025 for the same ASINs (deal type and deal price as Keepa logged them).", {}),
    ("Format: CSV, comma separated, dates as YYYY-MM-DD, prices in USD as numbers (no $). Name files keepa_viewer.csv and keepa_history_<ASIN>.csv (or keepa_history_all.csv). I ingest them straight into the yellow columns.", {}),
    ("Priority: P1 = Aspirational/Beatable competitors on hero and core products + our own child ASINs (needed for the Thu Oct 15 v1). P2 = the rest of the mapped set. Canada: same pull on Amazon.ca for our CA ASINs once the CA competitor set is mapped (see Data Gaps).", {}),
]
for i, (t, f) in enumerate(spec, 1):
    c = kr.cell(row=i, column=1, value=t); c.font = F(**f) if f else F(size=10)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    kr.merge_cells(start_row=i, start_column=1, end_row=i, end_column=8)
    kr.row_dimensions[i].height = 18 if i == 1 else 44
hr = len(spec) + 2
KR = [("p", "Priority", 7), ("role", "Ours / competitor", 10), ("asin", "ASIN (pull this)", 12), ("parent", "Parent ASIN", 12), ("brand", "Brand", 18),
      ("tier", "Tier", 10), ("prod", "Our product(s)", 40), ("fam", "Family", 8)]
for i, (k_, h_, w_) in enumerate(KR, 1):
    c = kr.cell(row=hr, column=i, value=h_); c.font = F(bold=True, color="FFFFFF"); c.fill = HDR; c.border = BOX
    kr.column_dimensions[get_column_letter(i)].width = w_
hero_core = {59, 30, 32, 36, 37, 31, 39, 38, 48, 45, 69}
req = {}
for (pid, mkt), d in cc.items():
    if mkt != "US":
        continue
    o = d["our"]
    if o.get("best_child_asin"):
        req.setdefault(o["best_child_asin"], {"p": "P1", "role": "OURS", "parent": o.get("parent_asin"), "brand": o.get("brand"), "tier": "", "prod": set(), "fam": PINFO.get(pid, (0, ""))[1]})["prod"].add(f"{o.get('brand')} {o.get('name')}")
for x in comp_rows:
    if not x["child"]:
        continue
    p = "P1" if (x["pid"] in hero_core and x["tier"] in ("Aspirational", "Beatable")) else "P2"
    e = req.setdefault(x["child"], {"p": p, "role": "competitor", "parent": x["parent"], "brand": x["brand"], "tier": x["tier"], "prod": set(), "fam": x["family"]})
    e["prod"].add(x["product"])
    if p == "P1":
        e["p"] = "P1"
rows_req = sorted(req.items(), key=lambda kv: (kv[1]["p"], kv[1]["role"] != "OURS", kv[1]["fam"], kv[1]["brand"] or ""))
with open(OUT_CSV, "w", newline="") as fh:
    w = csv.writer(fh)
    w.writerow(["priority", "role", "asin", "parent_asin", "brand", "tier", "our_products", "family", "marketplace"])
    for i, (a, e) in enumerate(rows_req, hr + 1):
        vals = [e["p"], e["role"], a, e["parent"], e["brand"], e["tier"], "; ".join(sorted(e["prod"])), e["fam"]]
        for j, val in enumerate(vals, 1):
            c = kr.cell(row=i, column=j, value=val); c.font = F(bold=(e["role"] == "OURS")); c.border = BOX
        w.writerow(vals + ["US"])
kr.freeze_panes = kr.cell(row=hr + 1, column=1)
kr.auto_filter.ref = f"A{hr}:H{hr + len(rows_req)}"
N_REQ = len(rows_req)
N_P1 = sum(1 for _, e in rows_req if e["p"] == "P1")

# ================================================================= Data Gaps
dg = wb.create_sheet("Data Gaps")
dg.sheet_view.showGridLines = False
GAPS = [
    ("Gap", "Why it matters", "Where it can come from", "What exactly to request", "Owner"),
    ("Competitors' BFCM 2025 deal type, dates, prices, discount", "Core of the lift forecast and the deal-depth benchmark", "Keepa (not in Command Center)", "Keepa Request tab, File 2 (+ File 3) for the P1 ASINs", "Shayan Rana / retrieval agent"),
    ("Competitors' BSR before / during / after BFCM 2025", "BSR lift per competitor = the multiple Erik asked to tune our forecast against", "Keepa sales-rank history", "Keepa File 2, sales_rank column, Sep 1 2025 - Jan 5 2026", "Shayan Rana / retrieval agent"),
    ("Is the strikethrough actually SHOWING now, and its label", "Command Center holds the List Price, not whether Amazon displays it", "Amazon page check, or Joseph's Command Center STR Panel (Data Hub > Competitors > STR Panel, 300 competitor ASINs daily)", "Ask Joseph for an export of the STR Panel for these ASINs; else check P1 ASINs on the page", "Shayan Rana / Joseph"),
    ("Did the competitor sell at its List Price recently", "Joseph's study: STR shows when List Price <=12% over AND sold at it", "Keepa price history (List Price vs Buy Box over 90 days)", "Keepa File 2, Sep 1 - Oct 13 2026 window", "Shayan Rana"),
    ("Competitor history for BFCM 2025 in Command Center", "Command Center competitor placement history starts Jan-Apr 2026 (feed retired Sep 16, 2026)", "Not recoverable from Command Center", "Covered by Keepa instead; Prime Day 2026 visibility is used as the event proxy", "-"),
    ("Canada competitor set", "CA files hold no competitor snapshots ('no snapshot in CA'), so CA has no competitor prices", "Command Center mapping for CA, or Keepa on Amazon.ca", "Ask Joseph to map CA rivals for B4, B6, S4, S6, SF, SSS4; then Keepa .ca pull", "Joseph / Shayan Rana"),
    ("Variation count + launch date for competitors missing from Data Dive", "Variation breadth and listing age shape how much a deal lifts the family", "Keepa Product Viewer (variation count, listed since)", "Keepa File 1 for rows with blank 'Variations (DD)'", "Retrieval agent"),
    ("Last year's deal prices and dates for OUR products", "Needed to read our own 2025 lift fairly (deal vs no deal, depth)", "Seller Central Promotions history / Umair; Keepa on our ASINs", "Umair: BFCM 2025 deal list per parent (type, dates, deal price); Keepa File 2 on our ASINs", "Umair / Shayan Rana"),
    ("Our stock position during BFCM 2025", "OOS variations suppress the lift we saw - Erik asked to adjust for it", "Inventory history (Haroon / David), Keepa OOS flags on our ASINs", "Days OOS per hero variation Nov 15 - Dec 5, 2025", "Haroon / David"),
    ("Planned 2026 deal type, dates, price per product", "Drives the undercut flag and the forecast scenario", "Umair's deal coverage list (due Oct 13) and Joseph's D1 board", "Fill the yellow cells on Lift Inputs", "Umair / Shayan Rana"),
    ("ASINsight traffic for unmeasured competitors", "Many rivals carry no ASINsight export (traffic blank)", "ASINsight / Data Rova export per competitor", "Export for P1 competitors with blank 'ASINsight 7d traffic'", "Shayan Rana"),
]
for i, row in enumerate(GAPS, 1):
    for j, val in enumerate(row, 1):
        c = dg.cell(row=i, column=j, value=val)
        c.font = F(bold=(i == 1), color=("FFFFFF" if i == 1 else "000000"), size=9)
        if i == 1:
            c.fill = HDR
        c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = BOX
for j, w in enumerate((34, 40, 36, 44, 18), 1):
    dg.column_dimensions[get_column_letter(j)].width = w

# ================================================================= Event Watch
ew = wb.create_sheet("Event Watch")
eh = [("Check date", 10), ("Time (ET)", 7), ("Family", 8), ("Our product", 22), ("Competitor brand", 16), ("ASIN", 12),
      ("Price ($)", 8), ("List Price ($)", 8), ("STR showing (Y/N)", 7), ("Deal live (type)", 12), ("Deal price ($)", 8),
      ("Coupon (%)", 7), ("BSR", 8), ("Our deal price ($)", 8), ("Below our deal price?", 8), ("Posted in #bfcm-2026? (Y/N)", 8), ("Notes", 30)]
header(ew, 1, [("x", h, w, "x") for h, w in eh], height=42)
ew.freeze_panes = "A2"
for rr2 in range(2, 402):
    for cc2 in range(1, len(eh) + 1):
        c = ew.cell(row=rr2, column=cc2); c.border = BOX
        if cc2 not in (14, 15):
            c.fill = INPUT; c.font = F(color=BLUE)
    ew.cell(row=rr2, column=1).number_format = "mmm d"
    ew.cell(row=rr2, column=14).value = f"=IFERROR(INDEX('Lift Inputs'!$H:$H,MATCH(D{rr2},'Lift Inputs'!$B:$B,0)),\"\")"
    ew.cell(row=rr2, column=15).value = (f'=IF(AND(ISNUMBER(N{rr2}),OR(ISNUMBER(G{rr2}),ISNUMBER(K{rr2}))),'
                                         f'IF(MIN(IF(ISNUMBER(G{rr2}),G{rr2},10^9),IF(ISNUMBER(K{rr2}),K{rr2},10^9))<N{rr2},"YES","no"),"")')
    for cc2 in (14, 15):
        ew.cell(row=rr2, column=cc2).fill = FILL["calc"]; ew.cell(row=rr2, column=cc2).font = F()
ex = [dt.date(2026, 11, 20), "10:00", "Bamboo", "Decolure Bamboo Sheets 4 Piece", "EXAMPLE - Pure Bamboo", "B07DZTQ4JC", 109.99, 129.99, "Y", "Best Deal", 87.99, None, 250, None, None, "Y", "Example row - overwrite"]
for i, v in enumerate(ex, 1):
    if i in (14, 15):
        continue
    ew.cell(row=2, column=i, value=v).font = F(color=BLUE, italic=True)
ew.conditional_formatting.add("O2:O401", FormulaRule(formula=['O2="YES"'], fill=FILL["flag"], font=Font(name=FONT, bold=True, color="C00000")))
for lst, rngs in (('"Y,N"', ("I2:I401", "P2:P401")), ('"None,Coupon,Lightning Deal,Best Deal,Prime Exclusive Discount,Price cut"', ("J2:J401",))):
    dv = DataValidation(type="list", formula1=lst, allow_blank=True); ew.add_data_validation(dv)
    for x in rngs:
        dv.add(x)

# ================================================================= Key Insights
ki = wb.create_sheet("Key Insights", 1)
ki.sheet_view.showGridLines = False
hdrs = [("Family", 9), ("Our product", 24), ("Topic", 16), ("Insight (computed from the data on Oct 13, 2026)", 120)]
for j, (h, w) in enumerate(hdrs, 1):
    c = ki.cell(row=1, column=j, value=h); c.font = F(bold=True, color="FFFFFF"); c.fill = HDR; c.border = BOX
    ki.column_dimensions[get_column_letter(j)].width = w
ki.freeze_panes = "A2"

def fmt_money(x):
    return f"${x:,.2f}" if isinstance(x, (int, float)) else "n/a"

def lift_of(pid, mkt="US"):
    d = ours.get((pid, mkt)) or {}
    w1, w2, w3 = d.get("w1_bfcm") or {}, d.get("w2_pre30") or {}, d.get("w3_dec") or {}
    t1, t2, t3 = w1.get("total_sales_derived"), w2.get("total_sales_derived"), w3.get("total_sales_derived")
    if not (t1 and t2):
        return None
    return {"lift": (t1 / 12) / (t2 / 30), "dec": ((t3 / 30) / (t2 / 30)) if t3 else None,
            "tacos": w1.get("tacos"), "pre_tacos": w2.get("tacos"), "spend_x": (w1.get("spend") / 12) / (w2.get("spend") / 30) if w1.get("spend") and w2.get("spend") else None}

ins = []
# portfolio-level
lifts = [(PINFO[p][2], lift_of(p)) for p in PINFO if lift_of(p)]
lifts.sort(key=lambda x: -x[1]["lift"])
ins.append(("All", "Portfolio", "Our BFCM 2025 lift",
            "BFCM week 2025 daily total sales vs the 30 days before (derived from ad spend / TACoS): " +
            ", ".join(f"{n} {v['lift']:.2f}x" for n, v in lifts) +
            ". Ad spend per day rose " + ", ".join(f"{n} {v['spend_x']:.1f}x" for n, v in lifts if v.get("spend_x")) + " - lift came with heavier spend, so read lift and TACoS together."))
allc = [x for x in comp_rows if isinstance(x.get("price"), (int, float))]
withlp = [x for x in allc if isinstance(x.get("lp"), (int, float)) and x["lp"] > x["price"]]
in12 = [x for x in withlp if x["lp"] / x["price"] - 1 <= 0.12]
over25 = [x for x in withlp if x["lp"] / x["price"] - 1 > 0.25]
ins.append(("All", "Portfolio", "List Price / STR landscape",
            f"{len(withlp)} of {len(allc)} mapped competitor rows ({len(withlp)/max(len(allc),1):.0%}) carry a List Price above their price. "
            f"{len(in12)} sit within 12% (the band Joseph's study found Amazon shows), {len(over25)} are more than 25% over (usually ignored). "
            f"Median List Price premium where set: {statistics.median([x['lp']/x['price']-1 for x in withlp]):.0%}. "
            "Whether Amazon actually displays each one still needs the page / STR Panel check (Competitors tab, 'STR showing now?')."))
pdv = [(x["brand"], x["pd_org"] / x["pd_pre_org"]) for x in comp_rows if x.get("pd_org") is not None and (x.get("pd_pre_org") or 0) >= 5]
if pdv:
    ins.append(("All", "Portfolio", "Prime Day 2026 behaviour",
                f"{len(pdv)} competitor rows ({len(set(b for b, _ in pdv))} brands) have Command Center visibility history over Prime Day 2026. Median organic visibility lift Jun 23-26 vs Jun 1-20: "
                f"{statistics.median([v for _, v in pdv]):.2f}x. Competitor history does not reach BFCM 2025 (starts Jan-Apr 2026), so Prime Day is the closest in-house event read; BFCM 2025 needs Keepa."))
for pid, fam, short, role in PRODUCTS:
    d = cc.get((pid, "US"))
    if not d:
        continue
    o = d["our"]
    pname = f"{o.get('brand') or ''} {o.get('name') or short}".strip()
    rows_ = [x for x in comp_rows if x["pid"] == pid]
    top = sorted([x for x in rows_ if isinstance(x.get("wk"), (int, float))], key=lambda x: -x["wk"])[:5]
    e = econ(o.get("economics"))
    if top:
        tot = sum(x["wk"] for x in rows_ if isinstance(x.get("wk"), (int, float)))
        ourwk = (e["units_28d"] or 0) / 4
        parts = []
        for x in top:
            lp = f", List Price {fmt_money(x['lp'])} (+{x['lp']/x['price']-1:.0%})" if isinstance(x.get("lp"), (int, float)) and isinstance(x.get("price"), (int, float)) and x["lp"] > x["price"] else ", no List Price"
            parts.append(f"{x['brand']} ({x['tier'] or 'untiered'}) ~{x['wk']:,}/wk at {fmt_money(x.get('price'))}{lp}")
        ins.append((fam, pname, "Who sells the most",
                    f"Top 5 of {len(rows_)} mapped rivals by est weekly units: " + "; ".join(parts) +
                    f". Tracked rivals ~{tot:,.0f} units/wk; we sell ~{ourwk:,.0f}/wk ({ourwk/(ourwk+tot):.1%} of the tracked set)." if tot else ""))
    lpz = [x for x in rows_ if isinstance(x.get("lp"), (int, float)) and isinstance(x.get("price"), (int, float)) and 0 < x["lp"] / x["price"] - 1 <= 0.12]
    lpn = [x for x in rows_ if isinstance(x.get("lp"), (int, float)) and isinstance(x.get("price"), (int, float)) and x["lp"] > x["price"]]
    if rows_:
        ins.append((fam, pname, "List Price / STR",
                    f"{len(lpn)} of {len(rows_)} rivals carry a List Price above price; {len(lpz)} inside 12%"
                    + (": " + ", ".join(f"{x['brand']} +{x['lp']/x['price']-1:.0%}" for x in lpz[:6]) if lpz else "")
                    + ". These are the ones most likely to show a strikethrough at BFCM - check them on the page first."))
    lf = lift_of(pid)
    if lf:
        ins.append((fam, pname, "Our BFCM 2025",
                    f"BFCM week 2025 total sales {lf['lift']:.2f}x the prior-30-day daily rate; TACoS {lf['tacos']:.1f}% in BFCM week vs {lf['pre_tacos']:.1f}% before; December ran {lf['dec']:.2f}x the pre-event rate." if lf.get("dec") and lf.get("tacos") is not None and lf.get("pre_tacos") is not None else f"BFCM week 2025 total sales {lf['lift']:.2f}x the prior-30-day daily rate."))
    bc = o.get("best_child_asin")
    v = pt.get(bc) or {}
    cpd = [x["pd_org"] / x["pd_pre_org"] for x in rows_ if x.get("pd_org") is not None and (x.get("pd_pre_org") or 0) >= 5]
    if v.get("pd26_pre_org") and cpd:
        ins.append((fam, pname, "Prime Day 2026 visibility",
                    f"Our best child {bc}: organic visibility {v['pd26_org']/v['pd26_pre_org']:.2f}x over Prime Day vs the 3 weeks before; mapped rivals' median {statistics.median(cpd):.2f}x (n={len(cpd)}, rivals with a baseline score of 5+). "
                    + ("Rivals gained more event visibility than we did." if statistics.median(cpd) > v['pd26_org']/v['pd26_pre_org'] else "We held or out-gained rivals on event visibility.")))
    if fam in ("Bamboo", "Satin", "Cooling") and e["price"] and rows_:
        below = [x for x in rows_ if isinstance(x.get("price"), (int, float)) and x.get("our_same") and x["price"] < x["our_same"]]
        ins.append((fam, pname, "Price position",
                    f"{len(below)} of {len(rows_)} rivals are priced below our same-size (or Queen) price today. Our blended price {fmt_money(e['price'])}, contribution {fmt_money(e['contribution'])}/unit, break-even ACoS {e['be_acos'] if e['be_acos'] is None else (e['be_acos'] if e['be_acos'] < 1.5 else e['be_acos']/100):.0%}."
                    if e['be_acos'] is not None else f"{len(below)} of {len(rows_)} rivals are priced below our same-size (or Queen) price today."))
rr3 = 2
for fam, pn, topic, text in ins:
    if not text:
        continue
    for j, val in enumerate((fam, pn, topic, text), 1):
        c = ki.cell(row=rr3, column=j, value=val); c.font = F(size=9, bold=(j == 3)); c.border = BOX
        c.alignment = Alignment(wrap_text=True, vertical="top")
    rr3 += 1
ki.cell(row=rr3 + 1, column=1, value="Computed at build time from the Command Center (rosters, prices, List Price, est weekly units, our PPC/TACoS, visibility) and Data Dive. Estimates are estimates; Keepa fills the BFCM 2025 competitor history.").font = F(italic=True)

wb.save(OUT)
print(json.dumps({"competitor_rows": LAST - FIRST + 1, "products": D_LAST - D_FIRST + 1, "ours_rows": orow - 2,
                  "vis_asins": vr_ - 2, "keepa_asins": N_REQ, "keepa_p1": N_P1, "dd_asins": len(dd)}))
